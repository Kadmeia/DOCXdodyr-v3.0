# -*- coding: utf-8 -*-
"""Инструмент для сквозного тестирования папок с претензиями в DOCXдодыр."""

import argparse
import json
import os
import shutil
import sys
from pathlib import Path

# Добавляем корень репозитория в sys.path
REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))
sys.path.insert(0, str(REPO_ROOT / "app"))

from backend_api import BackendApi
from folder_pipeline import (
    DEFAULT_OUTPUT_DIR_NAME,
    DEFAULT_DECODER_FILENAME,
    FolderAnonymizationPipeline,
    collect_supported_files,
)
from document_restorer import DocumentRestorer, is_valid_decoder_structure
from docx import Document


def get_configured_backend():
    """Создает и настраивает BackendApi для пакетного тестирования."""
    api = BackendApi()
    api.save_docx = True
    api.save_pdf = True
    api.save_decoder = True
    api.emit_audit_sidecars = True
    api.open_output_folder = False
    if not api._pullenti_processor:
        api.init_pullenti()
    return api


def run_anonymize(folder_path: str | Path):
    """Выполняет пакетное обезличивание одной папки."""
    folder = Path(folder_path).resolve()
    if not folder.is_dir():
        print(f"ОШИБКА: Папка не существует: {folder}", file=sys.stderr)
        sys.exit(1)

    print(f"=== Запуск пакетного обезличивания папки: {folder.name} ===")
    api = get_configured_backend()
    pipeline = FolderAnonymizationPipeline(api, output_dir_name=DEFAULT_OUTPUT_DIR_NAME)

    result = pipeline.process(
        folder,
        recursive=True,
        progress_callback=lambda cur, tot, msg: print(f"[{cur+1}/{tot}] {msg}"),
    )

    print("\n=== Результаты обработки ===")
    print(f"Всего файлов найдено: {len(result.files)}")
    print(f"Успешно обработано: {result.processed_count}")
    print(f"Ошибок: {result.error_count}")
    print(f"Всего замен: {result.total_replacements}")
    print(f"Выходная папка: {result.output_dir}")
    if result.decoder_path:
        print(f"Дешифратор: {result.decoder_path}")
    if result.errors:
        print("\nОшибки:")
        for err in result.errors:
            print(f"  - {err.get('file')}: {err.get('error')}")

    return result


def inspect_decoder(decoder_path: str | Path):
    """Выводит сводку по плейсхолдерам из дешифратора."""
    dec_file = Path(decoder_path).resolve()
    if not dec_file.is_file():
        print(f"Дешифратор не найден: {dec_file}")
        return

    with dec_file.open("r", encoding="utf-8") as f:
        mapping = json.load(f)

    print(f"\n=== Содержимое дешифратора ({len(mapping)} записей) ===")
    for token, val in sorted(mapping.items()):
        if isinstance(val, dict):
            orig = val.get("original", "")
            cat = val.get("category", "")
            print(f"  {token:25} -> {orig} [{cat}]")
        else:
            print(f"  {token:25} -> {val}")


def dump_extracted_texts(output_dir: str | Path):
    """Извлекает и печатает текст из всех обезличенных DOCX-файлов."""
    out_dir = Path(output_dir).resolve()
    if not out_dir.is_dir():
        print(f"Выходная папка не найдена: {out_dir}")
        return

    docx_files = sorted(out_dir.glob("**/*.docx"))
    print(f"\n=== Текстовое содержимое обезличенных DOCX ({len(docx_files)} файлов) ===")
    for doc_path in docx_files:
        print(f"\n--- ФАЙЛ: {doc_path.name} ---")
        try:
            doc = Document(str(doc_path))
            for i, p in enumerate(doc.paragraphs):
                text = p.text.strip()
                if text:
                    print(f"P{i+1}: {text}")
            for t_idx, table in enumerate(doc.tables):
                print(f"  [Таблица #{t_idx+1}, строк: {len(table.rows)}]")
                for r_idx, row in enumerate(table.rows):
                    row_vals = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    if any(row_vals):
                        print(f"    R{r_idx+1}: | {' | '.join(row_vals)} |")
        except Exception as exc:
            print(f"  [Ошибка чтения DOCX: {exc}]")


def restore_document(doc_path: str | Path, decoder_path: str | Path, output_doc_path: str | Path = None):
    """Восстанавливает документ по дешифратору."""
    doc_p = Path(doc_path).resolve()
    dec_p = Path(decoder_path).resolve()

    with dec_p.open("r", encoding="utf-8") as f:
        mapping = json.load(f)

    restorer = DocumentRestorer(mapping)
    if output_doc_path:
        out_p = Path(output_doc_path).resolve()
        shutil.copyfile(doc_p, out_p)
        target = out_p
    else:
        target = doc_p

    ok, msg = restorer.restore_docx(str(target))
    print(f"Восстановление {target.name}: ok={ok}, {msg}")
    return ok, target


def cleanup_folder(folder_path: str | Path):
    """Удаляет папку 'Обезличенные документы' и временные файлы."""
    folder = Path(folder_path).resolve()
    out_dir = folder / DEFAULT_OUTPUT_DIR_NAME
    removed_items = []

    if out_dir.exists():
        shutil.rmtree(out_dir, ignore_errors=True)
        removed_items.append(str(out_dir))

    for inc in folder.glob(".docxdodyr-*"):
        if inc.is_dir():
            shutil.rmtree(inc, ignore_errors=True)
            removed_items.append(str(inc))
        else:
            try:
                inc.unlink()
                removed_items.append(str(inc))
            except OSError:
                pass

    for pattern in (
        "*[Дд]ешифратор*.json",
        "*.manifest.json",
        "*.provenance.json",
        "*_хронологическая_справка*",
        "*_cleaned*.*",
        "*_restored*.*",
    ):
        for f in folder.glob(pattern):
            try:
                f.unlink()
                removed_items.append(str(f))
            except OSError:
                pass

    print(f"Очистка папки {folder.name}: удалено {len(removed_items)} объектов.")


def main():
    parser = argparse.ArgumentParser(description="Верификатор папок претензий DOCXдодыр")
    subparsers = parser.add_subparsers(dest="command", required=True)

    # anonymize
    p_anon = subparsers.add_parser("anonymize", help="Обезличить папку")
    p_anon.add_argument("folder", help="Путь к папке")

    # inspect
    p_insp = subparsers.add_parser("inspect", help="Показать дешифратор и тексты")
    p_insp.add_argument("folder", help="Путь к папке")

    # clean
    p_clean = subparsers.add_parser("clean", help="Удалить результаты обезличивания")
    p_clean.add_argument("folder", help="Путь к папке")

    # restore
    p_rest = subparsers.add_parser("restore", help="Восстановить DOCX")
    p_rest.add_argument("docx", help="Путь к DOCX")
    p_rest.add_argument("decoder", help="Путь к дешифратор.json")
    p_rest.add_argument("--out", help="Путь для сохранения")

    args = parser.parse_args()

    if args.command == "anonymize":
        run_anonymize(args.folder)
    elif args.command == "inspect":
        folder = Path(args.folder).resolve()
        out_dir = folder / DEFAULT_OUTPUT_DIR_NAME
        dec_path = out_dir / DEFAULT_DECODER_FILENAME
        inspect_decoder(dec_path)
        dump_extracted_texts(out_dir)
    elif args.command == "clean":
        cleanup_folder(args.folder)
    elif args.command == "restore":
        restore_document(args.docx, args.decoder, args.out)


if __name__ == "__main__":
    main()
