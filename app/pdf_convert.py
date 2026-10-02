# -*- coding: utf-8 -*-
"""Конвертация DOCX/XLSX в PDF и OCR PDF."""

from __future__ import annotations

import importlib.util
import json
import os
import sys
import subprocess
import shutil
import re
import tempfile
from collections import Counter
from pathlib import Path
from typing import Any

import config

import logging
_pdf_logger = logging.getLogger("pdf_convert")

try:
    PHONENUMBERS_AVAILABLE = importlib.util.find_spec("phonenumbers") is not None
    _pdf_logger.debug("Библиотека phonenumbers найдена.")
except (ImportError, ValueError):
    PHONENUMBERS_AVAILABLE = False
    _pdf_logger.info("Библиотека phonenumbers не найдена.")

try:
    from docx2pdf import convert as docx2pdf_convert
    DOCX2PDF_AVAILABLE = True
    _pdf_logger.debug("Библиотека docx2pdf найдена.")
except ImportError:
    DOCX2PDF_AVAILABLE = False
    _pdf_logger.info("Библиотека docx2pdf не найдена.")

def check_libreoffice():
    """Проверяет доступность LibreOffice"""
    possible_paths = [
        "soffice",
        "libreoffice",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for path in possible_paths:
        if shutil.which(path) or (os.path.exists(path) if os.path.isabs(path) else False):
            try:
                result = subprocess.run([path, "--version"],
                                      capture_output=True, text=True, timeout=10)
                if result.returncode == 0:
                    return path
            except Exception:
                continue
    return None

LIBREOFFICE_PATH = check_libreoffice()
LIBREOFFICE_AVAILABLE = LIBREOFFICE_PATH is not None

if LIBREOFFICE_AVAILABLE:
    _pdf_logger.debug(f"LibreOffice найден: {LIBREOFFICE_PATH}")
elif DOCX2PDF_AVAILABLE:
    _pdf_logger.debug("LibreOffice не найден, используем docx2pdf")

PDF_CONVERSION_AVAILABLE = DOCX2PDF_AVAILABLE or LIBREOFFICE_AVAILABLE


def get_pdf_conversion_status() -> dict[str, Any]:
    """Возвращает подробный статус возможностей конвертации документов в PDF.

    Явно отражает ограничения платформы Windows при отсутствии Microsoft Word
    или LibreOffice, при этом подтверждая автономную доступность OCR.
    """
    libreoffice_found = check_libreoffice() is not None
    docx2pdf_found = bool(DOCX2PDF_AVAILABLE)
    conversion_possible = libreoffice_found or docx2pdf_found

    if libreoffice_found:
        message = "Конвертация DOCX/XLSX в PDF доступна через LibreOffice."
        engine = "LibreOffice"
    elif docx2pdf_found:
        message = "Конвертация DOCX/XLSX в PDF доступна через Microsoft Word (docx2pdf)."
        engine = "Microsoft Word (docx2pdf)"
    else:
        message = (
            "Для конвертации DOCX/XLSX в PDF установите LibreOffice "
            "(бесплатно: https://www.libreoffice.org). "
            "Обработка и обезличивание существующих PDF-файлов полностью доступна через встроенный OCR."
        )
        engine = None

    return {
        "available": conversion_possible,
        "conversion_available": conversion_possible,
        "engine": engine,
        "libreoffice_available": libreoffice_found,
        "docx2pdf_available": docx2pdf_found,
        "message": message,
    }

# OCR model instances are intentionally created only on first use and reused
# for a batch of PDFs.  Importing this module never initializes or downloads a
# model.
_OCR_BACKEND_CACHE = {}


def clear_ocr_cache():
    """Освобождает память OCR-моделей при смене языка или сбросе настроек."""
    global _OCR_BACKEND_CACHE
    _OCR_BACKEND_CACHE.clear()


# A text layer shorter than this is still useful (e.g. a one-page certificate
# or a short table).  The previous 200-character cutoff sent such PDFs to OCR,
# which made them fail on installations without an OCR engine.
MIN_TEXT_LENGTH_FOR_FAST_PATH = 20

# Scanner applications often add a tiny searchable watermark to every page of
# an otherwise image-only PDF.  Treating that watermark as a document text
# layer is a privacy failure: the program would anonymise the watermark and
# release the untouched scan.  These markers are removed only for the quality
# decision; the original extracted text is returned unchanged when it passes.
_SCANNER_WATERMARK_RE = re.compile(
    r"^(?:scanned\s+with\s+anyscanner|scanned\s+with\s+camscanner|"
    r"camscanner|anyscanner)$",
    re.IGNORECASE,
)


class PDFTextQualityError(RuntimeError):
    """The PDF has no trustworthy text layer and OCR produced no usable text."""


class PDFOCRCancelled(RuntimeError):
    """OCR was stopped by the caller before the result could be completed."""


def _pdf_page_count(pdf_path: str) -> int | None:
    """Return a cheap page count for quality scoring, or ``None`` on failure."""

    try:
        import fitz

        document = fitz.open(str(pdf_path))
        try:
            return len(document)
        finally:
            document.close()
    except Exception:
        return None


def _pdf_text_layer_profile(pdf_path: str) -> tuple[int, int] | None:
    """Return ``(pages, informative_pages)`` for the native text layer."""

    snapshot = _pdf_native_page_snapshot(pdf_path)
    if snapshot is None:
        return None
    return len(snapshot), sum(1 for item in snapshot if item[3])


def _pdf_native_page_snapshot(
    pdf_path: str,
) -> list[tuple[str, int, int, bool, bool]] | None:
    """Return native text, page size and trust decision for every PDF page."""

    try:
        import fitz

        document = fitz.open(str(pdf_path))
        try:
            snapshot = []
            for page in document:
                raw_text = page.get_text("text")
                lines = [
                    re.sub(r"\s+", " ", line).strip()
                    for line in raw_text.splitlines()
                    if line.strip()
                ]
                content = " ".join(
                    line for line in lines
                    if not _SCANNER_WATERMARK_RE.fullmatch(line)
                )
                rect = page.rect
                # Resource lists may include unused images. Inspect displayed
                # images (including inline images) instead. A long native
                # heading does not prove that the text layer covers the image.
                has_images = bool(page.get_image_info())
                snapshot.append((
                    raw_text.strip(),
                    max(1, int(round(rect.width * 2))),
                    max(1, int(round(rect.height * 2))),
                    sum(char.isalnum() for char in content) >= 20 and not has_images,
                    has_images,
                ))
            return snapshot
        finally:
            document.close()
    except Exception:
        return None


def assess_pdf_text_quality(text: str, *, page_count: int | None = None) -> tuple[bool, str]:
    """Decide whether extracted/OCR text can safely represent the document.

    The gate is deliberately content-oriented rather than tied to a document
    type.  It rejects empty layers, scanner-watermark-only layers and repeated
    boilerplate spread over many pages, while still accepting a genuinely
    short one-page certificate or receipt.
    """

    normalized = re.sub(r"\s+", " ", str(text or "")).strip()
    if not normalized:
        return False, "пустой текст"
    alnum = sum(char.isalnum() for char in normalized)
    if len(normalized) < MIN_TEXT_LENGTH_FOR_FAST_PATH or alnum / max(len(normalized), 1) < 0.25:
        return False, "слишком мало читаемого текста"

    lines = [re.sub(r"\s+", " ", line).strip() for line in str(text).splitlines() if line.strip()]
    informative = [line for line in lines if not _SCANNER_WATERMARK_RE.fullmatch(line)]
    if not informative:
        return False, "найден только водяной знак приложения-сканера"

    pages = max(1, int(page_count or 1))
    informative_text = " ".join(informative)
    informative_alnum = sum(char.isalnum() for char in informative_text)
    counts = Counter(line.casefold() for line in informative)
    unique_alnum = sum(sum(char.isalnum() for char in line) for line in counts)

    # For multi-page files require enough non-boilerplate content across the
    # document.  A repeated header/footer alone must not authorize release.
    if pages > 1:
        if informative_alnum < max(40, pages * 20):
            return False, "недостаточно содержимого для числа страниц"
        if unique_alnum < max(30, pages * 8):
            return False, "текстовый слой состоит преимущественно из повторов"
    return True, "пригодный текстовый слой"

if not PDF_CONVERSION_AVAILABLE:
    print("ВНИМАНИЕ: Ни одна библиотека для конвертации PDF не найдена.")

def convert_docx_to_pdf_libreoffice(docx_path, pdf_path):
    """Конвертирует DOCX в PDF используя LibreOffice"""
    profile_dir = None
    conversion_dir = None
    sibling_temp = None
    try:
        output_dir = str(Path(pdf_path).resolve().parent)
        os.makedirs(output_dir, exist_ok=True)
        profile_dir = tempfile.mkdtemp(prefix="docxdodyr-lo-")
        # LibreOffice always chooses ``<source-stem>.pdf`` and may refuse to
        # overwrite an existing sibling.  Convert in an isolated directory,
        # then atomically replace the requested target.
        conversion_dir = tempfile.mkdtemp(prefix="docxdodyr-pdf-")
        profile_uri = Path(profile_dir).resolve().as_uri()
        cmd = [
            LIBREOFFICE_PATH,
            "--headless",
            f"-env:UserInstallation={profile_uri}",
            "--convert-to", "pdf",
            "--outdir", conversion_dir,
            str(docx_path)
        ]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
        if result.returncode == 0:
            generated_pdf = os.path.join(conversion_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
            if os.path.exists(generated_pdf):
                # The runtime temp directory and the user's output folder can
                # live on different volumes.  Copy to a sibling first, then
                # atomically replace within the destination filesystem.
                fd, sibling_temp = tempfile.mkstemp(
                    prefix=f".{Path(pdf_path).name}.", suffix=".tmp", dir=output_dir
                )
                os.close(fd)
                shutil.copyfile(generated_pdf, sibling_temp)
                os.replace(sibling_temp, pdf_path)
                sibling_temp = None
                return True
            return False
        print(f"LibreOffice error: {result.stderr}")
        return False
    except Exception as e:
        print(f"LibreOffice conversion failed: {e}")
        return False
    finally:
        if sibling_temp:
            Path(sibling_temp).unlink(missing_ok=True)
        if profile_dir:
            shutil.rmtree(profile_dir, ignore_errors=True)
        if conversion_dir:
            shutil.rmtree(conversion_dir, ignore_errors=True)

def convert_docx_to_pdf_word_mac(docx_path, pdf_path):
    """Конвертирует DOCX в PDF используя Microsoft Word на macOS через AppleScript"""
    try:
        script = '''
on run argv
    set srcFile to (POSIX file (item 1 of argv)) as alias
    set tgtPath to item 2 of argv
    tell application "Microsoft Word"
        set wasRunning to running
        if not wasRunning then
            launch
        end if
        set theDoc to missing value
        try
            open srcFile
            set theDoc to active document
            save as theDoc file name tgtPath file format format PDF
            close theDoc saving no
            if not wasRunning then
                quit
            end if
        on error errMsg number errNum
            if theDoc is not missing value then
                try
                    close theDoc saving no
                end try
            end if
            if not wasRunning then
                try
                    quit
                end try
            end if
            error errMsg number errNum
        end try
    end tell
    return "OK"
end run
'''
        src_resolved = str(Path(docx_path).resolve())
        tgt_resolved = str(Path(pdf_path).resolve())
        cmd = ["osascript", "-e", script, src_resolved, tgt_resolved]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=30)

        if result.returncode == 0 and os.path.exists(pdf_path):
            return True
        else:
            print(f"Word on macOS conversion failed (returncode={result.returncode}): {result.stderr.strip()}")
            return False
    except subprocess.TimeoutExpired:
        print("Word on macOS conversion timed out (likely a sandbox/permissions dialog)")
        return False
    except Exception as e:
        print(f"Word on macOS conversion failed: {e}")
        return False

def convert_docx_to_pdf_word(docx_path, pdf_path):
    """Конвертирует DOCX в PDF используя Microsoft Word"""
    word_app = None
    doc = None
    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()
        try:
            word_app = win32com.client.DispatchEx("Word.Application")
            word_app.Visible = False
            word_app.DisplayAlerts = False
            doc = word_app.Documents.Open(str(docx_path))
            doc.SaveAs2(str(pdf_path), FileFormat=17)
            return True
        finally:
            try:
                if doc is not None:
                    doc.Close(SaveChanges=False)
            finally:
                try:
                    if word_app is not None:
                        word_app.Quit()
                finally:
                    pythoncom.CoUninitialize()
    except Exception as e:
        print(f"Word conversion failed: {e}")
        return False

def convert_xlsx_to_pdf_excel(xlsx_path, pdf_path):
    """Конвертирует XLSX в PDF используя Microsoft Excel"""
    excel_app = None
    wb = None
    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()
        try:
            excel_app = win32com.client.DispatchEx("Excel.Application")
            excel_app.Visible = False
            excel_app.DisplayAlerts = False
            wb = excel_app.Workbooks.Open(str(xlsx_path))
            wb.ExportAsFixedFormat(0, str(pdf_path))
            return True
        finally:
            try:
                if wb is not None:
                    wb.Close(SaveChanges=False)
            finally:
                try:
                    if excel_app is not None:
                        excel_app.Quit()
                finally:
                    pythoncom.CoUninitialize()
    except Exception as e:
        print(f"Excel conversion failed: {e}")
        return False


def convert_docx_text_to_pdf(docx_path, pdf_path):
    """Deterministically render a text-first DOCX to PDF without Office.

    PDF-source anonymization creates a plain text DOCX from extraction/OCR.
    This renderer is the safe fallback for that adapter when GUI Office
    automation is unavailable; it embeds a Cyrillic-capable font and paginates
    by measured line width instead of silently omitting a requested PDF.
    """

    pdf = None
    temp_path = None
    try:
        import fitz
        from docx import Document

        windir = os.environ.get("WINDIR") or os.environ.get("SYSTEMROOT") or r"C:\Windows"
        font_candidates = (
            os.path.join(windir, "Fonts", "arial.ttf"),
            os.path.join(windir, "Fonts", "times.ttf"),
            os.path.join(windir, "Fonts", "calibri.ttf"),
            r"C:\Windows\Fonts\arial.ttf",
            r"C:\Windows\Fonts\times.ttf",
            r"C:\Windows\Fonts\calibri.ttf",
            "/System/Library/Fonts/Supplemental/Times New Roman.ttf",
            "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        )
        font_path = next((item for item in font_candidates if os.path.isfile(item)), None)
        if not font_path:
            # Fallback to standard Helvetica if specific font file not found
            font_path = None
        source = Document(str(docx_path))
        paragraphs = [paragraph.text for paragraph in source.paragraphs]
        for table in source.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        pdf = fitz.open()
        page_width, page_height = fitz.paper_size("a4")
        margin = 50.0
        font_size = 11.0
        line_height = 15.0
        font = fitz.Font(fontfile=font_path) if font_path else fitz.Font("helv")
        page = None
        y = margin

        def new_page():
            nonlocal page, y
            page = pdf.new_page(width=page_width, height=page_height)
            y = margin

        def write_line(value):
            nonlocal y
            if page is None or y + line_height > page_height - margin:
                new_page()
            if font_path:
                page.insert_text(
                    (margin, y), value, fontname="DOCXDodyrText",
                    fontfile=font_path, fontsize=font_size, color=(0, 0, 0),
                )
            else:
                page.insert_text(
                    (margin, y), value, fontname="helv",
                    fontsize=font_size, color=(0, 0, 0),
                )
            y += line_height

        for paragraph in paragraphs:
            words = paragraph.split()
            if not words:
                write_line("")
                continue
            line = words[0]
            for word in words[1:]:
                candidate = f"{line} {word}"
                if font.text_length(candidate, fontsize=font_size) <= page_width - 2 * margin:
                    line = candidate
                else:
                    write_line(line)
                    line = word
            write_line(line)
            y += 4.0
        if page is None:
            new_page()
        target = Path(pdf_path).absolute()
        target.parent.mkdir(parents=True, exist_ok=True)
        fd, temp_path = tempfile.mkstemp(prefix=f".{target.name}.", suffix=".tmp", dir=target.parent)
        os.close(fd)
        pdf.save(temp_path, garbage=4, deflate=True)
        pdf.close()
        pdf = None
        os.replace(temp_path, target)
        temp_path = None
        return True
    except Exception as exc:
        print(f"Text DOCX conversion failed: {exc}")
        return False
    finally:
        if pdf is not None:
            pdf.close()
        if temp_path is not None:
            Path(temp_path).unlink(missing_ok=True)

def convert_to_pdf(source_path, pdf_path):
    """Универсальная функция конвертации в PDF (DOCX/DOCM/XLSX/XLS -> PDF)"""
    src = str(source_path).lower()
    is_excel = src.endswith('.xlsx') or src.endswith('.xls') or src.endswith('.xlsm')
    if LIBREOFFICE_AVAILABLE:
        if convert_docx_to_pdf_libreoffice(source_path, pdf_path):
            return True
    if is_excel:
        if convert_xlsx_to_pdf_excel(source_path, pdf_path):
            return True
        # Фоллбэк для систем без Office/LibreOffice: конвертация книги в DOCX и рендер в PDF
        temp_docx = None
        wb = None
        try:
            from openpyxl import load_workbook
            from docx import Document
            wb = load_workbook(str(source_path), data_only=True)
            doc = Document()
            first_sheet = True
            for ws in getattr(wb, "worksheets", []):
                if not first_sheet:
                    doc.add_page_break()
                first_sheet = False
                doc.add_heading(ws.title, level=2)
                rows_data = []
                for row in ws.iter_rows(values_only=True):
                    if any(c is not None and str(c).strip() for c in row):
                        rows_data.append([str(c) if c is not None else "" for c in row])
                if rows_data:
                    max_cols = max(len(r) for r in rows_data)
                    if max_cols > 0:
                        norm_rows = [r + [""] * (max_cols - len(r)) for r in rows_data]
                        table = doc.add_table(rows=len(norm_rows), cols=max_cols)
                        table.style = 'Table Grid'
                        for r_idx, row in enumerate(norm_rows):
                            for c_idx, val in enumerate(row):
                                table.cell(r_idx, c_idx).text = val
            fd, temp_docx_name = tempfile.mkstemp(prefix=".docxdodyr_xlsx_to_pdf_", suffix=".docx")
            os.close(fd)
            temp_docx = Path(temp_docx_name)
            doc.save(str(temp_docx))
            return convert_docx_text_to_pdf(temp_docx, pdf_path)
        except Exception as exc:
            _pdf_logger.warning("excel to pdf fallback failed: %s", exc)
            return False
        finally:
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass
            if temp_docx is not None and temp_docx.exists():
                try:
                    temp_docx.unlink(missing_ok=True)
                except OSError:
                    pass
    if sys.platform == "darwin":
        if convert_docx_to_pdf_word_mac(source_path, pdf_path):
            return True

    if DOCX2PDF_AVAILABLE:
        try:
            docx2pdf_convert(str(source_path), str(pdf_path))
            if Path(pdf_path).is_file():
                return True
        except Exception as exc:
            _pdf_logger.warning("docx2pdf conversion failed: %s", exc)

    if sys.platform == "win32":
        if convert_docx_to_pdf_word(source_path, pdf_path):
            return True
    return convert_docx_text_to_pdf(source_path, pdf_path)

def extract_text_from_pdf(pdf_path: str) -> str:
    """
    Извлекает текст из "текстового" PDF (например, полученного из Word) без OCR.
    Возвращает строку (может быть пустой, если текста нет или он не извлекается).
    """
    pdf_path = str(pdf_path)

    # 1) pypdf/PyPDF2 (предпочтительно для текстовых PDF)
    if getattr(config, "PYPDF_AVAILABLE", False):
        try:
            try:
                from pypdf import PdfReader  # type: ignore
            except Exception:  # noqa: BLE001
                from PyPDF2 import PdfReader  # type: ignore

            reader = PdfReader(pdf_path)
            texts = []
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    texts.append(t)
            return "\n\n".join(texts).strip()
        except Exception:
            # Не роняем весь пайплайн — дадим шанс альтернативному извлечению
            pass

    # 2) PyMuPDF (если установлен) — тоже умеет извлекать текст без OCR
    try:
        import fitz  # PyMuPDF
    except Exception:  # noqa: BLE001
        return ""

    try:
        doc = fitz.open(pdf_path)
        try:
            parts = []
            for page in doc:
                t = page.get_text("text")
                if t:
                    parts.append(t)
            return "\n\n".join(parts).strip()
        finally:
            doc.close()
    except Exception:
        return ""

def pdf_to_text_auto(
    pdf_path: str,
    ocr_lang: str = "rus+eng",
    export_tsv: bool = False,
    export_hocr: bool = False,
) -> tuple[str, str]:
    """
    Автовыбор сценария для PDF:
    - если PDF содержит извлекаемый текст (Word->PDF), используем extract_text_from_pdf()
    - иначе делаем OCR (сканы/картинки) через ocr_pdf_to_text()

    Возвращает (text, method) где method ∈ {"text", "ocr", "mixed"}.
    """
    text, method, _details = pdf_to_text_auto_detailed(
        pdf_path, ocr_lang=ocr_lang, export_tsv=export_tsv,
        export_hocr=export_hocr,
    )
    return text, method


def pdf_to_text_auto_detailed(
    pdf_path: str,
    ocr_lang: str = "rus+eng",
    export_tsv: bool = False,
    export_hocr: bool = False,
    preferred_backend: str | None = None,
    cancel_check=None,
):
    """Like :func:`pdf_to_text_auto`, retaining OCR geometry for review UI.

    The third return value is ``None`` for a usable text layer and an
    ``OCRDocumentResult`` for a scanned or mixed PDF.  Keeping the legacy
    wrapper above avoids breaking existing integrations.
    """
    pdf_path = str(pdf_path)
    if cancel_check and cancel_check():
        raise PDFOCRCancelled("Распознавание PDF отменено пользователем")

    native_pages = _pdf_native_page_snapshot(pdf_path)
    profile = (
        (len(native_pages), sum(1 for item in native_pages if item[3]))
        if native_pages is not None else _pdf_text_layer_profile(pdf_path)
    )
    page_count = profile[0] if profile else _pdf_page_count(pdf_path)
    extracted = extract_text_from_pdf(pdf_path)
    usable, reason = assess_pdf_text_quality(extracted, page_count=page_count)
    ocr_page_numbers = []
    if native_pages:
        ocr_page_numbers = [
            index for index, item in enumerate(native_pages, start=1)
            if not item[3] and item[4]
        ]
    if ocr_page_numbers:
        usable = False
        reason = (
            f"неполный текстовый слой: OCR требуется для "
            f"{len(ocr_page_numbers)} из {len(native_pages)} страниц"
        )
    elif profile and profile[0] > 1 and profile[1] / profile[0] < 0.80:
        usable = False
        reason = (
            f"неполный текстовый слой: содержимое есть только на "
            f"{profile[1]} из {profile[0]} страниц"
        )
    if usable:
        print(f"[PDF] Извлечён текст без OCR (len={len(extracted)})")
        return extracted, "text", None
    print(f"[PDF] Текстовый слой отклонён: {reason}; требуется OCR")

    try:
        details = ocr_pdf_to_result(
            pdf_path, lang=ocr_lang, export_tsv=export_tsv,
            export_hocr=export_hocr,
            preferred_backend=preferred_backend,
            page_numbers=(ocr_page_numbers or None),
            cancel_check=cancel_check,
        )
    except PDFOCRCancelled:
        raise
    except Exception as exc:
        raise PDFTextQualityError(
            f"PDF не имеет пригодного текстового слоя ({reason}), а OCR недоступен: {exc}"
        ) from exc
    if ocr_page_numbers and native_pages is not None:
        from ocr_backend import OCRDocumentResult, OCRPageResult

        recognized = {page.page: page for page in details.pages}
        missing_results = []
        for page_number in ocr_page_numbers:
            page_result = recognized.get(page_number)
            if page_result is None:
                missing_results.append(f"стр.{page_number}: нет данных OCR")
                continue
            raw_text = (page_result.text or "").strip()
            lines = [re.sub(r"\s+", " ", line).strip() for line in raw_text.splitlines() if line.strip()]
            informative = [line for line in lines if not _SCANNER_WATERMARK_RE.fullmatch(line)]
            if raw_text and not informative:
                missing_results.append(f"стр.{page_number}: найден только водяной знак приложения-сканера")
                continue
            alnum = sum(char.isalnum() for char in raw_text)
            if alnum > 0:
                continue
            missing_results.append(f"стр.{page_number}: пустой текст")
        if missing_results:
            raise PDFTextQualityError(
                "OCR не дал пригодного текста для страниц: "
                + "; ".join(missing_results[:5])
                + ". Результат не сохранён."
            )

        combined_pages = []
        for page_number, (page_text, width, height, _informative, _has_images) in enumerate(
            native_pages, start=1
        ):
            if page_number in recognized:
                combined_pages.append(recognized[page_number])
            else:
                combined_pages.append(OCRPageResult(
                    page=page_number,
                    width=width,
                    height=height,
                    regions=(),
                    native_text=page_text,
                    source="text",
                ))
        combined_text = "\n\n".join(page.text for page in combined_pages)
        details = OCRDocumentResult(
            text=combined_text, pages=tuple(combined_pages), backend=details.backend
        )

    ocr_usable, ocr_reason = assess_pdf_text_quality(details.text, page_count=page_count)
    if not ocr_usable:
        raise PDFTextQualityError(
            f"OCR не дал пригодного текста: {ocr_reason}. Результат не сохранён."
        )
    return details.text, ("mixed" if ocr_page_numbers else "ocr"), details

def ocr_pdf_to_text(
    pdf_path: str,
    lang: str = "rus",
    ocr_lang: str | None = None,
    export_tsv: bool = False,
    export_hocr: bool = False,
    progress_cb=None,
    preferred_backend: str | None = None,
    ocr_workers: int | None = None,
    page_numbers: list[int] | tuple[int, ...] | None = None,
    cancel_check=None,
) -> str:
    """Извлекает текст из PDF через OCR (обертка для обратной совместимости)."""
    effective_lang = ocr_lang if ocr_lang is not None else lang
    return ocr_pdf_to_result(
        pdf_path,
        lang=effective_lang,
        export_tsv=export_tsv,
        export_hocr=export_hocr,
        progress_cb=progress_cb,
        preferred_backend=preferred_backend,
        ocr_workers=ocr_workers,
        page_numbers=page_numbers,
        cancel_check=cancel_check,
    ).text


def ocr_pdf_to_result(
    pdf_path: str,
    lang: str = "rus",
    export_tsv: bool = False,
    export_hocr: bool = False,
    progress_cb=None,
    preferred_backend: str | None = None,
    ocr_workers: int | None = None,
    page_numbers: list[int] | tuple[int, ...] | None = None,
    cancel_check=None,
):
    """Return the unified OCR result with page, bbox and confidence.

    ``export_tsv`` and ``export_hocr`` remain accepted for compatibility.  The
    returned :class:`ocr_backend.OCRDocumentResult` uses normalized top-left
    coordinates, making highlights independent of the selected OCR engine.
    """

    del export_tsv, export_hocr
    from dataclasses import replace
    from ocr_backend import (
        OCRBackendError,
        OCRBackendUnavailable,
        OCRDocumentResult,
        OCRPageResult,
        OCRTextRegion,
        create_ocr_backend,
    )
    try:
        import fitz  # PyMuPDF
    except Exception as exc:  # pragma: no cover - dependency guard
        raise RuntimeError(
            "OCR для PDF недоступен: установите PyMuPDF (pymupdf)."
        ) from exc

    try:
        cache_key = (lang, preferred_backend)
        cached = _OCR_BACKEND_CACHE.get(cache_key)
        if cached is None:
            cached = create_ocr_backend(preferred=preferred_backend, language=lang)
            _OCR_BACKEND_CACHE[cache_key] = cached
        backend_info, backend = cached
    except OCRBackendUnavailable as exc:
        raise RuntimeError(str(exc)) from exc

    pdf_path = str(pdf_path)
    print(f"[OCR] {backend_info.label}: {pdf_path} (lang={lang})")
    try:
        document = fitz.open(pdf_path)
    except Exception as exc:
        raise RuntimeError(f"Не удалось открыть PDF для OCR: {exc}") from exc

    pages: list[OCRPageResult] = []
    failed_backends: set[str] = set()
    result_backend_name = backend_info.name
    try:
        # 2x gives small Russian glyphs enough pixels while keeping memory bounded.
        matrix = fitz.Matrix(2, 2)
        total_pages = len(document)
        if page_numbers is None:
            selected_page_numbers = list(range(1, total_pages + 1))
        else:
            selected_page_numbers = [int(value) for value in page_numbers]
            if not selected_page_numbers:
                raise ValueError("page_numbers must not be empty")
            if len(set(selected_page_numbers)) != len(selected_page_numbers):
                raise ValueError("page_numbers must not contain duplicates")
            if any(value < 1 or value > total_pages for value in selected_page_numbers):
                raise ValueError("page_numbers contains a page outside the PDF")
            selected_page_numbers.sort()

        # Apple Vision creates an independent request/handler for every image;
        # pytesseract starts an independent local process per call.  Normal
        # text-heavy pages can therefore be recognized concurrently by these
        # adapters.  Keep the queue bounded: a rendered 2x page may occupy tens
        # of megabytes before PNG compression.  If a page is sparse or the
        # backend fails, discard the speculative result and run the unchanged
        # sequential safety path below, including its second-engine check.
        if ocr_workers is None:
            default_workers = 3 if backend_info.name == "tesseract" else 2
            try:
                ocr_workers = int(os.environ.get(
                    "DOCXDODYR_OCR_WORKERS", str(default_workers)
                ))
            except (TypeError, ValueError):
                ocr_workers = default_workers
        worker_count = max(1, min(
            int(ocr_workers), 4, len(selected_page_numbers) or 1
        ))
        if cancel_check and cancel_check():
            raise PDFOCRCancelled("Распознавание PDF отменено пользователем")
        if backend_info.name in {"applevision", "tesseract"} and worker_count > 1:
            from collections import deque
            from concurrent.futures import ThreadPoolExecutor

            pending = deque()
            parallel_pages: list[OCRPageResult] = []
            parallel_safe = True

            def collect_parallel_page(item) -> None:
                nonlocal parallel_safe
                page_number, width, height, future = item
                try:
                    page_regions = list(future.result() or [])
                except Exception:
                    parallel_safe = False
                    return
                alnum_count = sum(
                    char.isalnum()
                    for region in page_regions
                    for char in str(getattr(region, "text", region) or "")
                )
                if alnum_count < 25:
                    parallel_safe = False
                    return
                normalized = []
                for region in page_regions:
                    if not isinstance(region, OCRTextRegion):
                        region = OCRTextRegion(str(getattr(region, "text", region)).strip())
                    if region.text.strip():
                        normalized.append(
                            replace(region, page=page_number, text=region.text.strip())
                        )
                parallel_pages.append(
                    OCRPageResult(
                        page=page_number,
                        width=width,
                        height=height,
                        regions=tuple(normalized),
                    )
                )

            with ThreadPoolExecutor(max_workers=worker_count) as executor:
                for index in selected_page_numbers:
                    if cancel_check and cancel_check():
                        raise PDFOCRCancelled("Распознавание PDF отменено пользователем")
                    page = document.load_page(index - 1)
                    print(f"[OCR] Обработка страницы {index}/{total_pages}")
                    if progress_cb:
                        try:
                            progress_cb(index, total_pages)
                        except Exception:
                            pass
                    pixmap = page.get_pixmap(matrix=matrix, alpha=False)
                    page_bytes = pixmap.tobytes("png")
                    pending.append((
                        index,
                        int(pixmap.width),
                        int(pixmap.height),
                        executor.submit(backend.recognize_regions, page_bytes),
                    ))
                    if len(pending) >= worker_count:
                        if cancel_check and cancel_check():
                            raise PDFOCRCancelled("Распознавание PDF отменено пользователем")
                        collect_parallel_page(pending.popleft())
                        if not parallel_safe:
                            break
                while pending:
                    if cancel_check and cancel_check():
                        raise PDFOCRCancelled("Распознавание PDF отменено пользователем")
                    collect_parallel_page(pending.popleft())

            if parallel_safe and len(parallel_pages) == len(selected_page_numbers):
                full_text = "\n\n".join(page.text for page in parallel_pages)
                full_text = re.sub(r"(\w)[\-‐‑‒–]\n(\w)", r"\1\2", full_text)
                print(f"[OCR] OCR завершён, длина текста: {len(full_text)} символов")
                return OCRDocumentResult(
                    text=full_text,
                    pages=tuple(parallel_pages),
                    backend=backend_info.name,
                )

        for index in selected_page_numbers:
            if cancel_check and cancel_check():
                raise PDFOCRCancelled("Распознавание PDF отменено пользователем")
            page = document.load_page(index - 1)
            print(f"[OCR] Обработка страницы {index}/{total_pages}")
            if progress_cb:
                try:
                    progress_cb(index, total_pages)
                except Exception:
                    pass
            pixmap = page.get_pixmap(matrix=matrix, alpha=False)
            page_bytes = pixmap.tobytes("png")
            try:
                recognize_regions = getattr(backend, "recognize_regions", None)
                if callable(recognize_regions):
                    regions = recognize_regions(page_bytes)
                else:
                    # Compatibility for third-party/legacy adapters exposing
                    # only recognize(image_bytes) -> list[str].
                    lines = backend.recognize(page_bytes)
                    regions = [OCRTextRegion(str(line).strip()) for line in lines if str(line).strip()]
                if cancel_check and cancel_check():
                    raise PDFOCRCancelled("Распознавание PDF отменено пользователем")
            except PDFOCRCancelled:
                raise
            except Exception as exc:
                # An engine may initialize successfully but fail at runtime
                # (for example on a malformed/unsupported page).  Try one
                # different installed adapter before failing the whole PDF.
                failed_info = backend_info
                failed_backends.add(failed_info.name)
                try:
                    fallback = create_ocr_backend(
                        language=lang, excluded=failed_backends
                    )
                except OCRBackendUnavailable as fallback_error:
                    raise RuntimeError(
                        f"{failed_info.label} не распознал страницу {index}: {exc}"
                    ) from fallback_error
                backend_info, backend = fallback
                # Keep the preferred backend in the cache.  A transient Vision
                # failure (for example outside an interactive macOS session)
                # must not make the fallback primary for later PDFs.
                try:
                    recognize_regions = getattr(backend, "recognize_regions", None)
                    if callable(recognize_regions):
                        regions = recognize_regions(page_bytes)
                    else:
                        lines = backend.recognize(page_bytes)
                        regions = [OCRTextRegion(str(line).strip()) for line in lines if str(line).strip()]
                except Exception as fallback_error:
                    raise RuntimeError(
                        f"OCR не распознал страницу {index}: "
                        f"{failed_info.label}: {exc}; "
                        f"{backend_info.label}: {fallback_error}"
                    ) from fallback_error
            # A backend may return successfully while silently missing a
            # sparse/low-confidence page (passport backs and stamps are a
            # common example).  For such pages, ask one other installed
            # engine and keep whichever result contains more trustworthy
            # Russian/alphanumeric content.  This is a safety net, not a
            # document-specific correction.
            page_backend_name = backend_info.name
            primary_regions = list(regions or [])

            def region_quality(items):
                texts = [str(getattr(item, "text", item) or "").strip() for item in items]
                text = " ".join(value for value in texts if value)
                alnum = sum(char.isalnum() for char in text)
                cyrillic = len(re.findall(r"[А-Яа-яЁё]", text))
                confidences = [
                    float(getattr(item, "confidence", 0.0) or 0.0)
                    for item in items
                    if getattr(item, "confidence", None) is not None
                ]
                mean_conf = sum(confidences) / len(confidences) if confidences else 0.0
                return alnum + min(cyrillic, 200) * 0.35 + mean_conf * 20

            primary_alnum = sum(
                char.isalnum()
                for item in primary_regions
                for char in str(getattr(item, "text", item) or "")
            )
            if primary_alnum < 25:
                try:
                    alternative_info, alternative_backend = create_ocr_backend(
                        language=lang, excluded=failed_backends | {backend_info.name}
                    )
                    alternative_recognize = getattr(alternative_backend, "recognize_regions", None)
                    if callable(alternative_recognize):
                        alternative_regions = alternative_recognize(page_bytes)
                    else:
                        alternative_lines = alternative_backend.recognize(page_bytes)
                        alternative_regions = [
                            OCRTextRegion(str(line).strip())
                            for line in alternative_lines if str(line).strip()
                        ]
                    if region_quality(alternative_regions or []) > region_quality(primary_regions):
                        regions = alternative_regions
                        page_backend_name = alternative_info.name
                        # The alternative wins this page only.  Do not replace
                        # the configured backend for subsequent pages: a
                        # sparse page is not evidence that the primary engine
                        # is broken for the rest of the document.
                except (OCRBackendUnavailable, OCRBackendError):
                    pass
                except Exception:
                    pass
            normalized_regions = []
            for region in regions or []:
                if not isinstance(region, OCRTextRegion):
                    region = OCRTextRegion(str(getattr(region, "text", region)).strip())
                if region.text.strip():
                    normalized_regions.append(replace(region, page=index, text=region.text.strip()))
            pages.append(
                OCRPageResult(
                    page=index,
                    width=int(pixmap.width),
                    height=int(pixmap.height),
                    regions=tuple(normalized_regions),
                )
            )
            result_backend_name = page_backend_name
    finally:
        document.close()

    full_text = "\n\n".join(page.text for page in pages)
    full_text = re.sub(r"(\w)[\-‐‑‒–]\n(\w)", r"\1\2", full_text)
    print(f"[OCR] OCR завершён, длина текста: {len(full_text)} символов")
    return OCRDocumentResult(text=full_text, pages=tuple(pages), backend=result_backend_name)


def reconstruct_structured_paragraphs(cleaned_text: str) -> list[str]:
    """Восстанавливает естественную структуру абзацев, заголовков и списков из OCR/PDF."""
    if not cleaned_text or not cleaned_text.strip():
        return []

    paragraphs = []
    blocks = re.split(r"\n\s*\n+", cleaned_text.strip())

    HEADING_RE = re.compile(
        r"^(?:(?:ДОГОВОР|СОГЛАШЕНИЕ|АКТ|ПРИЛОЖЕНИЕ|УТВЕРЖДАЮ|РАСПОРЯЖЕНИЕ|ПРИКАЗ|РЕШЕНИЕ|ОПРЕДЕЛЕНИЕ|ПОСТАНОВЛЕНИЕ|ПРЕАМБУЛА|ПРЕДМЕТ|ПРАВА И ОБЯЗАННОСТИ|ОТВЕТСТВЕННОСТЬ|РЕКВИЗИТЫ|ПОДПИСИ СТОРОН)\b|[0-9]+(?:\.[0-9]+)*\.?\s+[А-ЯЁA-Z])"
    )
    LIST_ITEM_RE = re.compile(r"^(?:[0-9]+(?:\.[0-9]+)*\.?|[а-яёa-z]\)|\-|\•|\*)\s+")
    LABEL_RE = re.compile(
        r"^(?:(?:Заказчик|Исполнитель|Арендатор|Арендодатель|Покупатель|Продавец|Истец|Ответчик|Судья|Адрес|ИНН|КПП|ОГРН|БИК|Р/с|К/с|Банк|Тел|E-mail|Генеральный директор|Директор|Подпись|М\.П\.)\s*:?)",
        re.I,
    )

    for block in blocks:
        lines = [ln.strip() for ln in block.splitlines() if ln.strip()]
        if not lines:
            continue

        current_para = [lines[0]]

        for line in lines[1:]:
            prev_line = current_para[-1]

            is_new_item = (
                LIST_ITEM_RE.match(line)
                or HEADING_RE.match(line)
                or LABEL_RE.match(line)
                or prev_line.endswith(":")
                or (len(prev_line) < 45 and not prev_line.endswith(","))
                or "|" in line
            )

            if is_new_item:
                paragraphs.append(" ".join(current_para).strip())
                current_para = [line]
            else:
                if prev_line.endswith("-"):
                    current_para[-1] = prev_line[:-1] + line
                else:
                    current_para.append(line)

        if current_para:
            paragraphs.append(" ".join(current_para).strip())

    return [p for p in paragraphs if p.strip()]
