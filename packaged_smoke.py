"""Synthetic round-trip for the installed runtime, with no user document input."""
from contextlib import contextmanager, redirect_stdout, redirect_stderr
import io
import json
import os
from pathlib import Path
import sys
import tempfile

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
import webview


@contextmanager
def _isolated_directory(path):
    previous = Path.cwd()
    os.chdir(path)
    try:
        yield
    finally:
        os.chdir(previous)


def run_keyring():
    """Exercise the real OS vault using one unique, disposable synthetic key."""
    import keyring
    import uuid
    from review_context import SecureContextVault
    service = 'DOCXdodyr.synthetic-packaged-test.' + uuid.uuid4().hex
    account = 'synthetic-only'
    try:
        with tempfile.TemporaryDirectory(prefix='docxdodyr-native-vault-') as folder:
            path = Path(folder) / 'vault.json'
            data = {'synthetic': {'value': 'SYNTHETIC-VAULT-MARKER'}}
            SecureContextVault(path, service=service, account=account).save(data)
            if b'SYNTHETIC-VAULT-MARKER' in path.read_bytes():
                raise RuntimeError('Vault stored plaintext')
            if SecureContextVault(path, service=service, account=account).load() != data:
                raise RuntimeError('Native vault roundtrip failed')
    finally:
        if keyring.get_password(service, account) is not None:
            keyring.delete_password(service, account)
    if keyring.get_password(service, account) is not None:
        raise RuntimeError('Synthetic key cleanup failed')
    return {'status': 'ok', 'native_keyring_vault_roundtrip': True,
            'synthetic_entry_removed': True, 'backend': type(keyring.get_keyring()).__module__}


def run():
    keys=[f'DOCXDODYR_{kind}_DIR' for kind in ('DATA','USER_DATA','CONFIG','CACHE','LOG')]
    before={key:os.environ.get(key) for key in keys}
    api=None
    try:
        with tempfile.TemporaryDirectory(prefix='docxdodyr-smoke-') as directory:
            root=Path(directory)
            for key in keys:os.environ[key]=str(root/key.lower())
            with _isolated_directory(root), redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
                from docx import Document
                import fitz
                import openpyxl
                from backend_api import BackendApi
                from document_restorer import DocumentRestorer
                source=root/'synthetic.docx'
                original='fixture@example.invalid'
                document=Document();document.add_paragraph(original);document.save(source)
                api=BackendApi();api.save_original=True;api.save_decoder=True;api.save_log_files=False
                api.save_pdf=False;api.irreversible_pdf=False
                count=api.process_single_file(str(source),set(),{})
                outputs=[p for p in root.glob('*.docx') if p!=source]
                from document_restorer import find_decoder_near_document
                decoders=[find_decoder_near_document(p) for p in outputs]
                if count<1 or len(outputs)!=1 or len(decoders)!=1:raise RuntimeError('Synthetic processing failed')
                if original in '\n'.join(p.text for p in Document(outputs[0]).paragraphs):raise RuntimeError('Synthetic entity was not replaced')
                mapping = json.loads(decoders[0].read_text(encoding="utf-8"))
                restored=root/'restored.docx'
                ok,_message=DocumentRestorer(mapping).restore_docx(outputs[0],restored)
                if not ok or original not in '\n'.join(p.text for p in Document(restored).paragraphs):raise RuntimeError('Synthetic restoration failed')
                # Exercise the frozen backend's non-DOCX paths with entirely synthetic inputs.
                xlsx = root/'synthetic.xlsx'
                wb = openpyxl.Workbook(); ws = wb.active; ws['A1'] = original; wb.save(xlsx)
                text_pdf = root/'synthetic-text.pdf'
                pdf = fitz.open(); page = pdf.new_page(); page.insert_text((40, 60), original, fontsize=18); pdf.save(text_pdf); pdf.close()
                scan_pdf = root/'synthetic-scan.pdf'
                scan_text = f'Contact email for this synthetic document: {original}. Reply to this address only.'
                source_pdf = fitz.open(); page = source_pdf.new_page(width=1600, height=300); page.insert_text((80, 180), scan_text, fontsize=30); pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                scan_doc = fitz.open(); scan_page = scan_doc.new_page(width=1600, height=300); scan_page.insert_image(scan_page.rect, stream=pix.tobytes('png')); scan_doc.save(scan_pdf); scan_doc.close(); source_pdf.close()
                format_results = {}
                fixtures = [('xlsx', xlsx), ('text_pdf', text_pdf)]
                try:
                    from ocr_backend import create_ocr_backend
                    create_ocr_backend(language='eng')
                    fixtures.append(('scanned_pdf', scan_pdf))
                except Exception:
                    pass
                for name, fixture in fixtures:
                    if name == 'scanned_pdf':
                        api.ocr_lang = 'eng'
                    count = api.process_single_file(str(fixture), set(), {})
                    format_results[name] = {'processed': count > 0, 'count': count}
                if not all(item['processed'] for item in format_results.values()):
                    raise RuntimeError(f'Synthetic XLSX/PDF/OCR processing failed: {format_results}')
                cleaned_xlsx = list(root.glob('*cleaned*.xlsx'))
                if not cleaned_xlsx or openpyxl.load_workbook(cleaned_xlsx[0]).active['A1'].value == original:
                    raise RuntimeError('Synthetic XLSX output was not anonymized')
                restored_xlsx = root/'restored.xlsx'
                xlsx_decoder = find_decoder_near_document(cleaned_xlsx[0])
                xlsx_mapping = json.loads(xlsx_decoder.read_text(encoding='utf-8')) if xlsx_decoder else mapping
                restored_ok, _ = DocumentRestorer(xlsx_mapping).restore_excel(cleaned_xlsx[0], restored_xlsx)
                if not restored_ok or not restored_xlsx.exists() or openpyxl.load_workbook(restored_xlsx).active['A1'].value != original:
                    raise RuntimeError('Synthetic XLSX restoration failed')
                api.shutdown();api=None
            return {'status':'ok','docx_anonymization':True,'decoder':True,'docx_restoration':True,
                    'xlsx_processing': format_results['xlsx']['processed'],
                    'text_pdf_processing': format_results['text_pdf']['processed'],
                    'scanned_pdf_ocr': format_results['scanned_pdf']['processed'] if 'scanned_pdf' in format_results else 'not-tested'}
    finally:
        if api is not None:api.shutdown()
        for key,value in before.items():
            if value is None:os.environ.pop(key,None)
            else:os.environ[key]=value
