import docx
from pathlib import Path
from backend_api import Worker, BackendApi


def test_clean_document_reports_granular_progress(tmp_path):
    doc = docx.Document()
    for i in range(20):
        doc.add_paragraph(f"Параграф {i}")
    tbl = doc.add_table(rows=2, cols=2)
    tbl.rows[0].cells[0].text = "Ячейка 1"
    tbl.rows[0].cells[1].text = "Ячейка 2"

    api = BackendApi.__new__(BackendApi)
    api.anonymize_text_pullenti = lambda text, *args, **kwargs: (text, 0, [])
    api.get_paragraph_text_with_revisions = lambda p: p.text
    api.replace_text_in_paragraph_xml = lambda *args, **kwargs: None
    api.postprocess_placeholder_tails = lambda text: text
    api._write_back_qwen_document = lambda *args, **kwargs: None

    progress_calls = []
    def on_progress(step, total, detail=""):
        progress_calls.append((step, total, detail))

    api.clean_document(doc, set(), set(), progress_callback=on_progress)
    assert len(progress_calls) >= 2
    assert progress_calls[0][0] == 0
    assert progress_calls[0][2] == "Подготовка..."
    assert progress_calls[-1][0] == progress_calls[-1][1]
    assert progress_calls[-1][2] == "Анализ завершен"


def test_worker_reports_intra_file_percentage(tmp_path):
    doc_path = tmp_path / "test.docx"
    doc = docx.Document()
    for i in range(5):
        doc.add_paragraph(f"Текст {i}")
    doc.save(str(doc_path))

    progress_reports = []

    class MockCleaner:
        _window = None
        user_exclusions = set()
        custom_replacements = set()
        batch_id = None
        enable_crash_recovery = False
        save_decoder = False
        emit_audit_sidecars = False
        open_output_folder = False

        def update_progress(self, current_step, max_steps, status_text, percent=None):
            progress_reports.append((current_step, max_steps, status_text, percent))

        def process_single_file(self, file_path, exclusions, replacements, *, batch_mapping=None,
                                batch_entity_seen=None, defer_decoder=False, progress_callback=None):
            if progress_callback:
                progress_callback(10, 20, "Абзац 10/20")
                progress_callback(20, 20, "Анализ завершен")
            return 0

        def is_cancelled(self):
            return False

        def finish_operation(self, _status):
            pass

    cleaner = MockCleaner()
    Worker([str(doc_path)], cleaner).run()

    # Verify that update_progress received the 50% intra-file tick
    percentages = [p[3] for p in progress_reports if p[3] is not None]
    assert 50 in percentages
    assert 100 in percentages


def test_worker_handles_cleaner_without_window_attribute(tmp_path):
    class NoWindowCleaner:
        user_exclusions = set()
        custom_replacements = set()
        batch_id = None
        enable_crash_recovery = False
        save_decoder = False
        emit_audit_sidecars = False
        open_output_folder = False

        def update_progress(self, current_step, max_steps, status_text, percent=None):
            pass

        def process_single_file(self, file_path, exclusions, replacements, **kwargs):
            return 0

        def is_cancelled(self):
            return False

        def finish_operation(self, _status):
            pass

    cleaner = NoWindowCleaner()
    # Cleaner has no _window attribute at all, must not raise AttributeError
    result = Worker([], cleaner).run()
    assert result["status"] == "failed"

    # Also test non-empty batch with file
    test_file = tmp_path / "test.docx"
    test_file.write_text("dummy", encoding="utf-8")
    result_with_file = Worker([str(test_file)], cleaner).run()
    assert result_with_file["status"] == "succeeded"


def test_excel_reports_progress(tmp_path):
    import openpyxl
    xlsx_path = tmp_path / "test.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Лист1"
    ws["A1"] = "Иван Иванов"
    wb.save(str(xlsx_path))

    api = BackendApi.__new__(BackendApi)
    api._pullenti_processor = True
    api.current_placeholders = {}
    api.anonymize_text_pullenti = lambda text, *args, **kwargs: (text, 0, [])
    api.is_cancelled = lambda: False
    api.init_pullenti = lambda: True

    progress_calls = []
    def on_progress(step, total, detail=""):
        progress_calls.append((step, total, detail))

    api.process_single_file(
        str(xlsx_path), set(), set(),
        progress_callback=on_progress,
    )

    assert len(progress_calls) >= 2
    assert progress_calls[0][0] == 0
    assert progress_calls[-1][2] == "Обработка Excel завершена"

