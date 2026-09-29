# -*- coding: utf-8 -*-
from pathlib import Path
import json

from backend_api import Worker
from backend_api import BackendApi
from entity_registry import EntityResolver
from docx import Document


class _BatchCleaner:
    def __init__(self):
        self._window = None
        self.user_exclusions = set()
        self.custom_replacements = set()
        self.save_decoder = False
        self.seen_placeholders = []

    def update_progress(self, *_args):
        pass

    def process_single_file(
        self, file_path, _exclusions, _replacements,
        batch_mapping=None, batch_entity_seen=None, defer_decoder=False,
    ):
        assert defer_decoder is True
        value = "Иванов Иван Иванович" if Path(file_path).stem == "one" else "Иванову И. И."
        placeholder = EntityResolver(batch_entity_seen, batch_mapping).resolve(value, "PER", "ФИО")
        self.seen_placeholders.append(placeholder)
        return 1


def test_worker_shares_entity_session_across_all_files_in_batch(tmp_path):
    first = tmp_path / "one.docx"
    second = tmp_path / "two.docx"
    first.touch()
    second.touch()
    cleaner = _BatchCleaner()

    Worker([str(first), str(second)], cleaner).run()

    assert cleaner.seen_placeholders == ["[ФИО_1]", "[ФИО_1]"]


def test_every_decoder_gets_the_final_batch_wide_mapping(tmp_path):
    first = tmp_path / "one.docx"
    second = tmp_path / "two.docx"
    first.touch()
    second.touch()
    cleaner = _BatchCleaner()
    cleaner.save_decoder = True

    Worker([str(first), str(second)], cleaner).run()

    # Backend сохраняет файл с прописной «Д»: <stem>_Дешифратор.json
    first_map = json.loads((tmp_path / "one_Дешифратор.json").read_text(encoding="utf-8"))
    second_map = json.loads((tmp_path / "two_Дешифратор.json").read_text(encoding="utf-8"))
    assert first_map == second_map
    assert list(first_map) == ["[ФИО_1]"]


def test_real_docx_pipeline_keeps_one_id_across_case_forms(tmp_path):
    first = tmp_path / "first.docx"
    second = tmp_path / "second.docx"
    for path, text in (
        (first, "Иванов Иван Иванович подписал договор."),
        (second, "Согласно Иванову Ивану Ивановичу направлено письмо."),
    ):
        document = Document()
        document.add_paragraph(text)
        document.save(path)

    api = BackendApi()
    api.save_original = True
    api.save_pdf = False
    api.save_decoder = False
    mapping, seen = {}, {}
    for path in (first, second):
        api.process_single_file(
            str(path), set(), set(),
            batch_mapping=mapping, batch_entity_seen=seen,
        )

    assert Document(tmp_path / "first_cleaned.docx").paragraphs[0].text.startswith("[ФИО_1]")
    assert "[ФИО_1]" in Document(tmp_path / "second_cleaned.docx").paragraphs[0].text
    assert list(mapping) == ["[ФИО_1]"]
