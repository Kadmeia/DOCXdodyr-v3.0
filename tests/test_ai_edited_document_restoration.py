# -*- coding: utf-8 -*-
"""Тесты восстановления документов после правок в нейросетях (сценарии LLM / чистовик / правки ПДн)."""

import json
from pathlib import Path
import pytest
from docx import Document

from backend_api import BackendApi
from document_restorer import (
    DocumentRestorer,
    find_decoder_near_document,
    normalize_document_stem,
    is_edited_document_stem,
    restore_document,
)
from decoder_binding import publish_binding


def test_normalize_and_detect_edited_stems():
    """Проверяет нормализацию имен файлов с различными маркерами правок нейросетей."""
    # Сценарий со скриншота пользователя:
    assert normalize_document_stem("Договор_2026_Приложение_1_правки_ПДн_cleaned") == "Договор_2026_Приложение_1"
    assert is_edited_document_stem("Договор_2026_Приложение_1_правки_ПДн_cleaned") is True

    # Сценарий чистовика:
    assert normalize_document_stem("Договор_правки_ПДн_чистовик") == "Договор"
    assert is_edited_document_stem("Договор_правки_ПДн_чистовик") is True

    # Сценарии LLM и других редакций:
    assert normalize_document_stem("Соглашение_cleaned_LLM_правки") == "Соглашение"
    assert is_edited_document_stem("Соглашение_cleaned_LLM_правки") is True

    assert normalize_document_stem("Акт_1_v2_чистовик") == "Акт_1"
    assert is_edited_document_stem("Акт_1_v2_чистовик") is True

    # Обычный обезличенный файл:
    assert normalize_document_stem("договор_cleaned") == "договор"
    assert is_edited_document_stem("договор_cleaned") is False

    # Браузерные дубликаты скачивания из веб-интерфейса LLM:
    assert normalize_document_stem("Договор_cleaned (1)") == "Договор"
    assert is_edited_document_stem("Договор_cleaned (1)") is True
    assert normalize_document_stem("Договор_cleaned (2)") == "Договор"
    assert is_edited_document_stem("Договор_cleaned (2)") is True
    assert normalize_document_stem("Договор (копия)") == "Договор"
    assert is_edited_document_stem("Договор (копия)") is True
    assert normalize_document_stem("Договор - копия") == "Договор"

    # Англоязычные и популярные AI маркеры:
    assert normalize_document_stem("Договор_cleaned_edited") == "Договор"
    assert is_edited_document_stem("Договор_cleaned_edited") is True
    assert normalize_document_stem("Договор_cleaned_revised") == "Договор"
    assert is_edited_document_stem("Договор_cleaned_revised") is True
    assert normalize_document_stem("Договор_gpt") == "Договор"
    assert is_edited_document_stem("Договор_gpt") is True
    assert normalize_document_stem("Договор_deepseek") == "Договор"
    assert is_edited_document_stem("Договор_deepseek") is True
    assert normalize_document_stem("Договор_правки_2") == "Договор"
    assert is_edited_document_stem("Договор_правки_2") is True


def test_restore_ai_edited_document_screenshot_scenario(tmp_path):
    """Сценарий со скриншота пользователя:
    Документ был обезличен, отправлен в нейронку, в нем произошли правки,
    и он сохранен под именем 'Договор_2026_Приложение_1_правки_ПДн_cleaned.docx'.
    Автоматический поиск дешифратора и восстановление должны работать без ошибок.
    """
    # 1. Создаем исходный очищенный документ и дешифратор в папке
    base_stem = "Договор_2026_Приложение_1"
    cleaned_orig = tmp_path / f"{base_stem}_cleaned.docx"
    doc_orig = Document()
    doc_orig.add_paragraph("Стороны договора: [ФИО_1] и [ОРГ_1].")
    doc_orig.save(cleaned_orig)

    mapping = {
        "[ФИО_1]": "Иванов Иван Иванович",
        "[ОРГ_1]": "ООО «Альфа Ромашка»",
    }
    decoder_file = tmp_path / f"{base_stem}_Дешифратор.json"
    decoder_file.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
    publish_binding(decoder_file, [cleaned_orig], "ai-session-001", emit_sidecars=False)

    # 2. Нейросеть вносит правки: добавляет новый параграф, изменяет текст (sha256 меняется!)
    # и сохраняет файл под именем со скриншота пользователя
    edited_doc_path = tmp_path / "Договор_2026_Приложение_1_правки_ПДн_cleaned.docx"
    doc_edited = Document()
    doc_edited.add_paragraph("Стороны настоящего соглашения: [ФИО_1] и [ОРГ_1].")
    doc_edited.add_paragraph("Пункт 1.1 (правка нейросети): стороны пришли к соглашению об изменении условий.")
    doc_edited.save(edited_doc_path)

    # 3. Проверяем автопоиск дешифратора рядом с отредактированным документом
    found_dec = find_decoder_near_document(edited_doc_path, allow_edited=True)
    assert found_dec == decoder_file.resolve()

    # 4. Проверяем восстановление через BackendApi (логика GUI)
    api = BackendApi.__new__(BackendApi)
    api.restore_doc_paths = [str(edited_doc_path)]
    api.restore_json_paths = []
    api.auto_decoder = True
    api.settings = {"auto_decoder": True}
    api._window = None

    result = api._run_restore()
    assert len(result["successes"]) == 1
    assert len(result["errors"]) == 0

    # 5. Проверяем содержимое восстановленного файла
    restored_path = tmp_path / "Договор_2026_Приложение_1_правки_ПДн_cleaned_восстановлено.docx"
    assert restored_path.exists()

    restored_doc = Document(str(restored_path))
    full_text = "\n".join(p.text for p in restored_doc.paragraphs)
    assert "Иванов Иван Иванович" in full_text
    assert "ООО «Альфа Ромашка»" in full_text
    assert "[ФИО_1]" not in full_text
    assert "[ОРГ_1]" not in full_text
    assert "Пункт 1.1 (правка нейросети)" in full_text


def test_restore_ai_chistovik_scenario(tmp_path):
    """Сценарий чистовика:
    Пользователь загружает 'Договор_правки_ПДн_чистовик.docx',
    в котором текст изменен и имя отличается от оригинала.
    """
    base_stem = "Договор"
    cleaned_orig = tmp_path / f"{base_stem}_cleaned.docx"
    doc_orig = Document()
    doc_orig.add_paragraph("Заказчик: [ФИО_1], Исполнитель: [ОРГ_1]")
    doc_orig.save(cleaned_orig)

    mapping = {
        "[ФИО_1]": "Петров Петр Петрович",
        "[ОРГ_1]": "ПАО «Газпром»",
    }
    decoder_file = tmp_path / f"{base_stem}_Дешифратор.json"
    decoder_file.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
    publish_binding(decoder_file, [cleaned_orig], "ai-session-002", emit_sidecars=False)

    chistovik_path = tmp_path / "Договор_правки_ПДн_чистовик.docx"
    doc_chist = Document()
    doc_chist.add_paragraph("Чистовая редакция. Заказчик: [ФИО_1], Исполнитель: [ОРГ_1]")
    doc_chist.save(chistovik_path)

    # Проверяем автопоиск
    found_dec = find_decoder_near_document(chistovik_path, allow_edited=True)
    assert found_dec == decoder_file.resolve()

    # Восстановление через restore_document
    restored = restore_document(chistovik_path, allow_edited=True)
    assert restored.exists()
    doc_res = Document(str(restored))
    text = doc_res.paragraphs[0].text
    assert "Петров Петр Петрович" in text
    assert "ПАО «Газпром»" in text
    assert "[ФИО_1]" not in text


def test_restore_with_explicit_decoder_selection(tmp_path):
    """Проверяет сценарий, когда пользователь вручную перетаскивает дешифратор в поле 2 UI."""
    doc_path = tmp_path / "Произвольный_файл_после_LLM.docx"
    doc = Document()
    doc.add_paragraph("Подписант: [ФИО_1]")
    doc.save(doc_path)

    decoder_path = tmp_path / "Мой_дешифратор.json"
    mapping = {"[ФИО_1]": "Сидоров Сидор Сидорович"}
    decoder_path.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")

    api = BackendApi.__new__(BackendApi)
    api.restore_doc_paths = [str(doc_path)]
    api.restore_json_paths = [str(decoder_path)]
    api.auto_decoder = False
    api.settings = {}
    api._window = None

    result = api._run_restore()
    assert len(result["successes"]) == 1
    assert len(result["errors"]) == 0

    restored_path = tmp_path / "Произвольный_файл_после_LLM_восстановлено.docx"
    assert restored_path.exists()
    assert "Сидоров Сидор Сидорович" in Document(str(restored_path)).paragraphs[0].text


def test_restore_ai_generated_new_document_from_batch_folder_pipeline(tmp_path):
    """Сценарий пакетного обезличивания папки и последующего восстановления собранного ИИ файла:
    1. Пользователь загружает папку с документами на обезличивание.
    2. Пайплайн создает 'Дешифратор.json' в корне и 'Обезличенные документы/' с файлами.
    3. ИИ анализирует обезличенные файлы и собирает совершенно НОВЫЙ документ (например, 'Возражения_на_иск.docx').
    4. Пользователь запускает дешифровку этого собранного файла (как из 'Обезличенные документы', так и из корня папки).
    5. Программа должна автоматически найти пакетный дешифратор и восстановить файл без ошибок.
    """
    import re
    from types import SimpleNamespace
    from folder_pipeline import FolderAnonymizationPipeline, DEFAULT_OUTPUT_DIR_NAME, DEFAULT_DECODER_FILENAME

    class _Occurrence:
        def __init__(self, source, start, end):
            self.sofa = source
            self.begin_char = start
            self.end_char = end - 1

    class _PersonReferent:
        type_name = "PERSON"
        def __init__(self, source, start, end):
            self.occurrence = [_Occurrence(source, start, end)]

    class _MockPersonProcessor:
        _name = re.compile(
            r"(?:Иванов(?:у|ым)?\s+Иван(?:у|ом)?\s+Иванович(?:у|ем)?|"
            r"Петров(?:у|ым)?\s+Петр(?:у|ом)?\s+Петрович(?:у|ем)?)",
            re.I,
        )
        def process(self, source, *_args):
            return SimpleNamespace(
                entities=[_PersonReferent(source, m.start(), m.end()) for m in self._name.finditer(source.text)]
            )

    case_dir = tmp_path / "Судебное_дело_№А40-12345"
    case_dir.mkdir()

    # 1. Создаем два исходных файла в папке
    doc1 = Document()
    doc1.add_paragraph("Истец: Иванов Иван Иванович")
    doc1.save(case_dir / "1_Исковое_заявление.docx")

    doc2 = Document()
    doc2.add_paragraph("Ответчик: Петров Петр Петрович")
    doc2.save(case_dir / "2_Договор_поставки.docx")

    # Инициализируем бэкенд
    api = BackendApi.__new__(BackendApi)
    api._pullenti_processor = _MockPersonProcessor()
    api.current_placeholders = {"PER": "[ФИО]"}
    api.bracket_type = "square"
    api.qwen_settings = SimpleNamespace(enabled=False)
    api._qwen_postprocessor = SimpleNamespace(last_status="disabled")
    api.user_exclusions = set()
    api.custom_replacements = set()
    api.save_original = True
    api.save_pdf = False
    api.save_decoder = True
    api.ocr_pdf = False
    api._window = None
    api.emit_audit_sidecars = False
    api.save_log_files = False
    api.save_audit_files = False
    api.open_output_folder = False

    pipeline = FolderAnonymizationPipeline(api, emit_sidecars=False)
    batch_result = pipeline.process(case_dir)

    assert batch_result.processed_count == 2
    assert batch_result.error_count == 0

    decoder_path = case_dir / DEFAULT_DECODER_FILENAME
    assert decoder_path.is_file()

    output_dir = case_dir / DEFAULT_OUTPUT_DIR_NAME
    assert output_dir.is_dir()

    # 2. ИИ читает обезличенные файлы и создает СОВЕРШЕННО НОВЫЙ документ
    # с новым именем (не совпадающим со стемом исходников и без суффикса _правки)
    ai_doc_in_output = output_dir / "Возражения_на_исковое_заявление.docx"
    ai_doc = Document()
    ai_doc.add_paragraph("В суд: Возражения стороны [ФИО_2] на иск стороны [ФИО_1].")
    ai_doc.save(ai_doc_in_output)

    # 3. Проверяем автопоиск дешифратора для файла, лежащего в 'Обезличенные документы'
    found_dec = find_decoder_near_document(ai_doc_in_output, allow_edited=True)
    assert found_dec == decoder_path.resolve()

    # 4. Проверяем восстановление через BackendApi (логика UI)
    api.restore_doc_paths = [str(ai_doc_in_output)]
    api.restore_json_paths = []
    api.auto_decoder = True
    api.settings = {"auto_decoder": True}
    api._window = None

    restore_res = api._run_restore()
    assert len(restore_res["successes"]) == 1
    assert len(restore_res["errors"]) == 0

    restored_in_output = output_dir / "Возражения_на_исковое_заявление_восстановлено.docx"
    assert restored_in_output.exists()
    restored_text = Document(str(restored_in_output)).paragraphs[0].text
    assert "В суд: Возражения стороны Петров Петр Петрович на иск стороны Иванов Иван Иванович." == restored_text

    # 5. Сценарий, когда ИИ сохранил собранный файл в корень папки дела (рядом с Дешифратор.json)
    ai_doc_in_root = case_dir / "Справка_по_делу_от_нейросети.docx"
    doc_root = Document()
    doc_root.add_paragraph("Справка: участники спора — [ФИО_1] и [ФИО_2].")
    doc_root.save(ai_doc_in_root)

    # Восстановление через restore_document напрямую
    restored_root = restore_document(ai_doc_in_root)
    assert restored_root.exists()
    root_text = Document(str(restored_root)).paragraphs[0].text
    assert "Справка: участники спора — Иванов Иван Иванович и Петров Петр Петрович." == root_text


def test_restore_browser_duplicate_download_scenario(tmp_path):
    """Сценарий скачивания отредактированного документа из браузера с суффиксом ' (1)':
    1. Исходный файл: 'Претензия.docx', обезличенный: 'Претензия_cleaned.docx', дешифратор: 'Претензия_Дешифратор.json'.
    2. Пользователь загружает в ChatGPT, тот возвращает правки.
    3. Браузер сохраняет файл как 'Претензия_cleaned (1).docx'.
    4. Автопоиск находит 'Претензия_Дешифратор.json', и восстанавливает без ошибок.
    """
    orig_cleaned = tmp_path / "Претензия_cleaned.docx"
    doc_orig = Document()
    doc_orig.add_paragraph("Претензия от заявителя [ФИО_1] к компании [ОРГ_1].")
    doc_orig.save(orig_cleaned)

    mapping = {
        "[ФИО_1]": "Смирнов Алексей Владимирович",
        "[ОРГ_1]": "ООО «ТехноСфера»",
    }
    decoder_file = tmp_path / "Претензия_Дешифратор.json"
    decoder_file.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
    publish_binding(decoder_file, [orig_cleaned], "session-browser-duplicate", emit_sidecars=False)

    # Файл, скачанный браузером с дубликатом имени
    downloaded_doc = tmp_path / "Претензия_cleaned (1).docx"
    doc_edited = Document()
    doc_edited.add_paragraph("Претензия от заявителя [ФИО_1] к компании [ОРГ_1].")
    doc_edited.add_paragraph("Дополнительное требование: возврат денежных средств.")
    doc_edited.save(downloaded_doc)

    # 1. Автопоиск через find_decoder_near_document
    found_dec = find_decoder_near_document(downloaded_doc, allow_edited=True)
    assert found_dec == decoder_file.resolve()

    # 2. Восстановление через restore_document
    restored = restore_document(downloaded_doc, allow_edited=True)
    assert restored.exists()
    full_text = "\n".join(p.text for p in Document(str(restored)).paragraphs)
    assert "Смирнов Алексей Владимирович" in full_text
    assert "ООО «ТехноСфера»" in full_text
    assert "[ФИО_1]" not in full_text

    # 3. Восстановление через BackendApi (логика UI)
    api = BackendApi.__new__(BackendApi)
    api.restore_doc_paths = [str(downloaded_doc)]
    api.restore_json_paths = []
    api.auto_decoder = True
    api.settings = {"auto_decoder": True}
    api._window = None

    result = api._run_restore()
    assert len(result["successes"]) == 1
    assert len(result["errors"]) == 0


def test_restore_ai_model_suffix_scenarios(tmp_path):
    """Сценарии сохранения файлов с суффиксами _edited, _deepseek, _gpt, (копия)."""
    orig_cleaned = tmp_path / "Договор_cleaned.docx"
    doc = Document()
    doc.add_paragraph("Заказчик: [ФИО_1]")
    doc.save(orig_cleaned)

    dec_file = tmp_path / "Договор_Дешифратор.json"
    dec_file.write_text(json.dumps({"[ФИО_1]": "Кузнецов Николай"}, ensure_ascii=False), encoding="utf-8")
    publish_binding(dec_file, [orig_cleaned], "session-ai-suffixes", emit_sidecars=False)

    for suffix in ("Договор_cleaned_edited.docx", "Договор_deepseek.docx", "Договор (копия).docx"):
        edited_p = tmp_path / suffix
        ed_doc = Document()
        ed_doc.add_paragraph("Отредактировано: [ФИО_1]")
        ed_doc.save(edited_p)

        found = find_decoder_near_document(edited_p, allow_edited=True)
        assert found == dec_file.resolve(), f"Не найден дешифратор для {suffix}"

        res = restore_document(edited_p, allow_edited=True)
        assert res.exists()
        assert "Кузнецов Николай" in Document(str(res)).paragraphs[0].text


def test_restore_arbitrary_filename_with_single_decoder_fallback(tmp_path):
    """Сценарий правила единственного дешифратора:
    Пользователь сохранил из нейросети файл с абсолютно произвольным именем 'ответ_нейросети.docx'.
    В папке лежит ровно один дешифратор 'Договор_Дешифратор.json'.
    Программа должна безопасно применить этот единственный дешифратор без требования совпадения стема.
    """
    orig_cleaned = tmp_path / "Договор_cleaned.docx"
    doc_orig = Document()
    doc_orig.add_paragraph("Заказчик: [ФИО_1]")
    doc_orig.save(orig_cleaned)

    dec_file = tmp_path / "Договор_Дешифратор.json"
    dec_file.write_text(json.dumps({"[ФИО_1]": "Смирнов Алексей Владимирович"}, ensure_ascii=False), encoding="utf-8")
    publish_binding(dec_file, [orig_cleaned], "session-single-fallback", emit_sidecars=False)

    # Файл с совершенно произвольным именем от нейросети
    ai_arbitrary_doc = tmp_path / "ответ_нейросети.docx"
    doc_ai = Document()
    doc_ai.add_paragraph("Итоговый проект соглашения. Заказчик: [ФИО_1].")
    doc_ai.save(ai_arbitrary_doc)

    # 1. find_decoder_near_document находит единственный дешифратор
    found = find_decoder_near_document(ai_arbitrary_doc, allow_edited=True)
    assert found == dec_file.resolve()

    # 2. restore_document успешно восстанавливает документ
    restored = restore_document(ai_arbitrary_doc, allow_edited=True)
    assert restored.exists()
    assert "Смирнов Алексей Владимирович" in Document(str(restored)).paragraphs[0].text

    # 3. Через BackendApi (UI)
    api = BackendApi.__new__(BackendApi)
    api.restore_doc_paths = [str(ai_arbitrary_doc)]
    api.restore_json_paths = []
    api.auto_decoder = True
    api.settings = {"auto_decoder": True}
    api._window = None

    result = api._run_restore()
    assert len(result["successes"]) == 1
    assert len(result["errors"]) == 0


def test_restore_arbitrary_filename_fails_safely_when_multiple_decoders_present(tmp_path):
    """Проверка безопасности: если в папке 2 или более разных дешифраторов,
    программа НЕ должна угадывать наугад для файла с чужим именем 'ответ_нейросети.docx',
    чтобы исключить подмену чужих персональных данных.
    """
    orig1 = tmp_path / "Договор_Иванов_cleaned.docx"
    Document().save(orig1)
    dec1 = tmp_path / "Договор_Иванов_Дешифратор.json"
    dec1.write_text(json.dumps({"[ФИО_1]": "Иванов Иван"}, ensure_ascii=False), encoding="utf-8")
    publish_binding(dec1, [orig1], "session-multi-1", emit_sidecars=False)

    orig2 = tmp_path / "Договор_Петров_cleaned.docx"
    Document().save(orig2)
    dec2 = tmp_path / "Договор_Петров_Дешифратор.json"
    dec2.write_text(json.dumps({"[ФИО_1]": "Петров Петр"}, ensure_ascii=False), encoding="utf-8")
    publish_binding(dec2, [orig2], "session-multi-2", emit_sidecars=False)

    # Произвольный файл, не связанный ни с Ивановым, ни с Петровым
    ai_doc = tmp_path / "ответ_нейросети.docx"
    Document().save(ai_doc)

    # find_decoder_near_document возвращает None из-за неоднозначности
    found = find_decoder_near_document(ai_doc, allow_edited=True)
    assert found is None

    # restore_document выбрасывает ошибку безопасной привязки вместо подмены чужих ПДн
    with pytest.raises(ValueError, match="отсутствует привязка к документу"):
        restore_document(ai_doc, allow_edited=True)



