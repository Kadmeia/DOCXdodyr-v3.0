import os
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFont


FONT_CANDIDATES = (
    Path('/System/Library/Fonts/Supplemental/Arial Unicode.ttf'),
    Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),
    Path('C:/Windows/Fonts/arial.ttf'),
)


def unicode_font() -> Path:
    for candidate in FONT_CANDIDATES:
        if candidate.is_file():
            return candidate
    raise RuntimeError('Для создания PDF нужен установленный шрифт с кириллицей')

# Define mock PII data
MOCK_PII = {
    "fio_1": "Иванов Иван Иванович",
    "fio_2": "Петров Петр Петрович",
    "fio_3": "Смирнов Алексей Владимирович",
    "phone_1": "+7 (999) 123-45-67",
    "phone_2": "8-800-555-35-35",
    "email_1": "ivanov@example.com",
    "inn_1": "7707083893"
}

def create_docx(path):
    doc = Document()
    
    # Header
    header = doc.sections[0].header
    header.paragraphs[0].text = f"Секретный документ для {MOCK_PII['fio_1']}"
    
    # Paragraphs
    doc.add_heading('Договор оказания услуг', 0)
    doc.add_paragraph(f"Настоящий договор заключен между ООО 'Ромашка' и гражданином {MOCK_PII['fio_1']}.")
    doc.add_paragraph(f"Контактный телефон исполнителя: {MOCK_PII['phone_1']}. Почта: {MOCK_PII['email_1']}.")
    
    # Table
    table = doc.add_table(rows=2, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'ФИО'
    hdr_cells[1].text = 'Телефон'
    hdr_cells[2].text = 'ИНН'
    row_cells = table.rows[1].cells
    row_cells[0].text = MOCK_PII['fio_2']
    row_cells[1].text = MOCK_PII['phone_2']
    row_cells[2].text = MOCK_PII['inn_1']

    doc.save(path)
    print(f"✅ Создан {path}")

def create_xlsx(path):
    wb = Workbook()
    
    # Sheet 1
    ws1 = wb.active
    ws1.title = "Сотрудники"
    ws1.append(["Имя", "Телефон", "Email"])
    ws1.append([MOCK_PII["fio_1"], MOCK_PII["phone_1"], MOCK_PII["email_1"]])
    ws1.append([MOCK_PII["fio_2"], MOCK_PII["phone_2"], "test@example.invalid"])
    
    # Sheet 2
    ws2 = wb.create_sheet("Подрядчики")
    ws2.append(["Контрагент", "ИНН", "Представитель"])
    ws2.append(["ООО Вектор", MOCK_PII["inn_1"], MOCK_PII["fio_3"]])

    wb.save(path)
    print(f"✅ Создан {path}")

def create_pdf_text(path):
    pdf = FPDF()
    pdf.add_font("DejaVu", "", str(unicode_font()), uni=True)
    pdf.add_page()
    pdf.set_font("DejaVu", size=14)
    pdf.cell(200, 10, txt="Соглашение о конфиденциальности", ln=True, align='C')
    pdf.ln(10)
    pdf.set_font("DejaVu", size=12)
    pdf.cell(200, 10, txt=f"Подписант 1: {MOCK_PII['fio_1']}", ln=True)
    pdf.cell(200, 10, txt=f"Телефон: {MOCK_PII['phone_1']}", ln=True)
    pdf.cell(200, 10, txt=f"ИНН: {MOCK_PII['inn_1']}", ln=True)
    pdf.output(path)
    print(f"✅ Создан {path}")

def create_pdf_scan(path):
    # Create an image with text
    img = Image.new('RGB', (800, 600), color='white')
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype(str(unicode_font()), 30)
    except (OSError, RuntimeError):
        font = ImageFont.load_default()
        
    text = f"Скан паспорта\nВыдан: {MOCK_PII['fio_2']}\nНомер телефона: {MOCK_PII['phone_2']}\nИНН: {MOCK_PII['inn_1']}"
    d.text((50, 50), text, font=font, fill=(0, 0, 0))
    
    # Save image to a temporary file
    img_path = path.replace('.pdf', '.png')
    img.save(img_path)
    
    # Create a PDF and embed the image
    pdf = FPDF()
    pdf.add_page()
    pdf.image(img_path, x=10, y=10, w=190)
    pdf.output(path)
    os.remove(img_path)
    print(f"✅ Создан {path}")

if __name__ == "__main__":
    source_dir = Path(__file__).resolve().parent / "source_files"
    source_dir.mkdir(parents=True, exist_ok=True)
    create_docx(str(source_dir / "test_mock.docx"))
    create_xlsx(str(source_dir / "test_mock.xlsx"))
    try:
        create_pdf_text(str(source_dir / "test_mock_text.pdf"))
        create_pdf_scan(str(source_dir / "test_mock_scan.pdf"))
    except Exception as e:
        print(f"Ошибка при создании PDF: {e}")
