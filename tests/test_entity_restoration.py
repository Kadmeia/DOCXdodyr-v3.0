# -*- coding: utf-8 -*-
from document_restorer import DocumentRestorer, is_valid_decoder_structure
from entity_registry import EntityResolver, canonical_key, entity_seen_from_mapping
from docx import Document


def test_person_case_forms_share_stable_id():
    mapping, seen = {}, {}
    resolver = EntityResolver(seen, mapping)

    placeholders = [
        resolver.resolve(value, "PER", "ФИО")
        for value in (
            "Иванов Иван Иванович",
            "Иванову Ивану Ивановичу",
            "Ивановым Иваном Ивановичем",
        )
    ]

    assert placeholders == ["[ФИО_1]"] * 3
    assert mapping["[ФИО_1]"]["id"] == "PER_000001"
    assert mapping["[ФИО_1]"]["forms"]["DATIVE"] == "Иванову Ивану Ивановичу"


def test_context_restoration_inflects_person():
    mapping, seen = {}, {}
    EntityResolver(seen, mapping).resolve("Иванов Иван Иванович", "PER", "ФИО")
    restorer = DocumentRestorer(mapping)

    assert restorer.apply_replacements("С [ФИО_1] заключен договор.") == \
        "С Ивановым Иваном Ивановичем заключен договор."
    assert restorer.apply_replacements("Согласно [ФИО_1], направлено письмо.") == \
        "Согласно Иванову Ивану Ивановичу, направлено письмо."


def test_organization_inflects_head_but_not_quoted_brand():
    mapping, seen = {}, {}
    EntityResolver(seen, mapping).resolve(
        "Коллегия адвокатов «Крупская и партнёры»", "ORG", "Наименование"
    )
    restorer = DocumentRestorer(mapping)

    assert restorer.apply_replacements("С [Наименование_1] заключен договор.") == \
        "С Коллегией адвокатов «Крупская и партнёры» заключен договор."


def test_decoder_accepts_legacy_and_structured_records():
    assert is_valid_decoder_structure({"[ФИО_1]": "Иванов И.И."})
    assert is_valid_decoder_structure({
        "[ФИО_1]": {"id": "PER_000001", "type": "PER", "original": "Иванов И.И."}
    })
    assert is_valid_decoder_structure({"/ФИО_1/": "Иванов И.И."})
    assert not is_valid_decoder_structure({"ФИО_1": "Иванов И.И."})


def test_english_entity_is_not_inflected():
    mapping, seen = {}, {}
    EntityResolver(seen, mapping).resolve("John Smith", "PER", "ФИО")
    assert DocumentRestorer(mapping).apply_replacements("С [ФИО_1] подписан договор.") == \
        "С John Smith подписан договор."


def test_canonical_person_key_ignores_case_and_order():
    assert canonical_key("Иванов Иван Иванович", "PER") == \
        canonical_key("Ивану Ивановичу Иванову", "PER")


def test_full_name_and_two_initials_share_one_cross_document_id():
    mapping, seen = {}, {}
    resolver = EntityResolver(seen, mapping)

    assert resolver.resolve("Иванов Иван Иванович", "PER", "ФИО") == "[ФИО_1]"
    assert resolver.resolve("И. И. Иванов", "PER", "ФИО") == "[ФИО_1]"
    assert resolver.resolve("Иванову И. И.", "PER", "ФИО") == "[ФИО_1]"
    assert len(mapping) == 1


def test_decoder_rebuilds_aliases_and_continues_after_highest_id():
    mapping = {
        "[ФИО_3]": {
            "id": "PER_000003",
            "type": "PER",
            "original": "Иванов Иван Иванович",
            "normalized": canonical_key("Иванов Иван Иванович", "PER")[1],
            "forms": {"DATIVE": "Иванову Ивану Ивановичу"},
        }
    }
    resolver = EntityResolver(entity_seen_from_mapping(mapping), mapping)

    assert resolver.resolve("Иванову Ивану Ивановичу", "PER", "[ФИО]") == "[ФИО_3]"
    assert resolver.resolve("Петров Петр Петрович", "PER", "[ФИО]") == "[ФИО_4]"


def test_numbered_resolver_honours_slash_placeholder_format():
    mapping, seen = {}, {}
    assert EntityResolver(seen, mapping).resolve(
        "Иванов Иван Иванович", "PER", "/ФИО/"
    ) == "/ФИО_1/"


def test_continued_resolver_inherits_existing_wrapper_format():
    mapping = {
        "/ФИО_3/": {
            "id": "PER_000003",
            "type": "PER",
            "original": "Иванов Иван Иванович",
            "normalized": canonical_key("Иванов Иван Иванович", "PER")[1],
        }
    }
    resolver = EntityResolver(entity_seen_from_mapping(mapping), mapping)
    assert resolver.resolve("Петров Петр Петрович", "PER", "[ФИО]") == "/ФИО_4/"


def test_docx_restore_preserves_run_formatting_and_split_token(tmp_path, bind_decoder):
    source = tmp_path / "formatted.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Истец: ").bold = True
    paragraph.add_run("[ФИО_").italic = True
    paragraph.add_run("1]").underline = True
    paragraph.add_run(" обратился в суд.")
    document.save(source)
    bind_decoder(source, {"[ФИО_1]": "Иванов Иван Иванович"})
    ok, _message = DocumentRestorer({"[ФИО_1]": "Иванов Иван Иванович"}).restore_docx(source)
    assert ok
    restored = Document(tmp_path / "formatted_восстановлено.docx")
    paragraph = restored.paragraphs[0]
    assert paragraph.text == "Истец: Иванов Иван Иванович обратился в суд."
    assert paragraph.runs[0].bold is True
    assert paragraph.runs[1].italic is True
    assert paragraph.runs[-1].text == " обратился в суд."
