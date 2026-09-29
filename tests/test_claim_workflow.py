from __future__ import annotations

import json
from pathlib import Path

import pytest

from docx import Document

from claim_workflow import ClaimWorkflowError, create_claim_from_verified_batch, load_verified_batch, validate_claim_text
from document_restorer import DocumentRestorer, find_decoder_near_document
from tests.test_folder_pipeline import _backend, _documents
from folder_pipeline import FolderAnonymizationPipeline


class _SafeQwen:
    last_status = "not_loaded"

    def process(self, text):
        self.last_status = "ok"
        return text.replace("  ", " ")


def test_verified_claim_qwen_and_auto_restore_roundtrip(tmp_path: Path):
    _documents(tmp_path)
    result = FolderAnonymizationPipeline(_backend()).process(tmp_path)
    claim_path = result.output_dir / "Исковое заявление.docx"

    def draft(_corpus, tokens):
        person = next(token for token in tokens if token.startswith("[ФИО_"))
        return f"# ИСКОВОЕ ЗАЯВЛЕНИЕ\n\nВ интересах {person} заявлено требование."

    created = create_claim_from_verified_batch(
        result.manifest_path,
        claim_path,
        draft_builder=draft,
        qwen_processor=_SafeQwen(),
    )
    assert created["qwen_status"] == "ok"
    assert find_decoder_near_document(claim_path) == result.decoder_path

    mapping = json.loads(result.decoder_path.read_text(encoding="utf-8"))
    ok, _message = DocumentRestorer(mapping).restore_docx(claim_path)
    assert ok
    restored = claim_path.with_name("Исковое заявление_восстановлено.docx")
    assert "Иванов" in " ".join(p.text for p in Document(restored).paragraphs)


def test_claim_rejects_unknown_placeholder_id(tmp_path: Path):
    _documents(tmp_path)
    result = FolderAnonymizationPipeline(_backend()).process(tmp_path)
    batch = load_verified_batch(result.manifest_path)
    known = next(token for token in batch.allowed_placeholders if token.startswith("[ФИО_"))

    with pytest.raises(ClaimWorkflowError, match="неизвестные плейсхолдеры"):
        validate_claim_text(f"Истец {known}; представитель [ФИО_999]", batch)


def test_claim_rejects_incomplete_or_failed_manifest(tmp_path: Path):
    _documents(tmp_path)
    result = FolderAnonymizationPipeline(_backend()).process(tmp_path)
    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))

    manifest["complete"] = False
    result.manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(ClaimWorkflowError, match="не помечена"):
        load_verified_batch(result.manifest_path)

    manifest["complete"] = True
    manifest["errors"] = [{"file": "source.pdf", "error": "OCR unavailable"}]
    result.manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(ClaimWorkflowError, match="пачке с ошибками"):
        load_verified_batch(result.manifest_path)

    manifest["errors"] = []
    manifest["expected_document_count"] += 1
    result.manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(ClaimWorkflowError, match="ожидаемым количеством"):
        load_verified_batch(result.manifest_path)


def test_claim_rejects_tampered_output_provenance(tmp_path: Path):
    _documents(tmp_path)
    result = FolderAnonymizationPipeline(_backend()).process(tmp_path)
    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    provenance_path = tmp_path / manifest["documents"][0]["outputs"][0]["provenance"]
    provenance = json.loads(provenance_path.read_text(encoding="utf-8"))
    provenance["decoder_sha256"] = "0" * 64
    provenance_path.write_text(json.dumps(provenance, ensure_ascii=False), encoding="utf-8")

    with pytest.raises(ClaimWorkflowError, match="другим дешифратором"):
        load_verified_batch(result.manifest_path)
