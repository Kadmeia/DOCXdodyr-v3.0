# -*- coding: utf-8 -*-
"""Генерация тестовых файлов с персональными данными и различными типами плейсхолдеров."""
import os
import subprocess
import shutil
from pathlib import Path

# Тексты с персональными данными для проверки обезличивания
SAMPLE_DATA = {
    "fio": "Иванов Иван Иванович",
    "fio_short": "Петрова Мария Сергеевна",
    "org": "ООО Ромашка",
    "org2": "АО Сбербанк",
    "inn": "7707083893",
    "ogrn": "1027700132195",
    "kpp": "773601001",
    "bik": "044525225",
    "account": "40702810000000001234",
    "corr_account": "30101810400000000225",
    "phone": "+7 (495) 123-45-67",
    "phone2": "8-800-555-35-35",
    "email": "ivanov.ivan@example.com",
    "address": "г. Москва, ул. Тверская, д. 1, кв. 10",
    "passport": "4512 123456",
    "birth_cert": "IV-АБ 123456",
}

DOCX_CONTENT = f"""
Договор оказания услуг № 001

Дата: 15.03.2024

Стороны договора:
Исполнитель: {SAMPLE_DATA['fio']}, паспорт {SAMPLE_DATA['passport']}
Заказчик: {SAMPLE_DATA['org']}, ИНН {SAMPLE_DATA['inn']}, ОГРН {SAMPLE_DATA['ogrn']}

Контактные данные исполнителя:
Телефон: {SAMPLE_DATA['phone']}
Email: {SAMPLE_DATA['email']}
Адрес: {SAMPLE_DATA['address']}

Представитель заказчика: {SAMPLE_DATA['fio_short']}

Реквизиты для оплаты:
Банк получателя: {SAMPLE_DATA['org2']}
БИК: {SAMPLE_DATA['bik']}
Корр. счёт: {SAMPLE_DATA['corr_account']}
Расчётный счёт: {SAMPLE_DATA['account']}
КПП: {SAMPLE_DATA['kpp']}

Свидетельство о рождении ребёнка: {SAMPLE_DATA['birth_cert']}
"""

EXCEL_ROWS = [
    ["ФИО", "Организация", "ИНН", "Телефон", "Email"],
    [SAMPLE_DATA["fio"], SAMPLE_DATA["org"], SAMPLE_DATA["inn"], SAMPLE_DATA["phone"], SAMPLE_DATA["email"]],
    [SAMPLE_DATA["fio_short"], SAMPLE_DATA["org2"], "7727563778", SAMPLE_DATA["phone2"], "petrova@example.com"],
    ["Сидоров П.П.", "ИП Сидоров", "123456789012", "+7 999 111-22-33", "sidorov@example.com"],
]


def create_test_docx(out_dir: Path) -> Path:
    """Создаёт тестовый DOCX с персональными данными."""
    from docx import Document
    out_path = out_dir / "test_obfuscation.docx"
    doc = Document()
    for para_text in DOCX_CONTENT.strip().split("\n\n"):
        doc.add_paragraph(para_text)
    doc.save(str(out_path))
    return out_path


def create_test_excel(out_dir: Path) -> Path:
    """Создаёт тестовый Excel с персональными данными."""
    try:
        from openpyxl import Workbook
    except ImportError:
        raise RuntimeError("openpyxl не установлен: pip install openpyxl")
    out_path = out_dir / "test_obfuscation.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Личные данные"
    for r_idx, row in enumerate(EXCEL_ROWS, 1):
        for c_idx, val in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    wb.save(str(out_path))
    return out_path


def create_test_pdf(out_dir: Path) -> Path:
    """Создаёт тестовый PDF с текстом (через DOCX + LibreOffice или fpdf2)."""
    out_path = out_dir / "test_obfuscation.pdf"
    # Сначала создаём DOCX
    from docx import Document
    docx_path = out_dir / "test_obfuscation_for_pdf.docx"
    doc = Document()
    doc.add_paragraph(DOCX_CONTENT.strip())
    doc.save(str(docx_path))

    # Пробуем docx2pdf (Windows/Mac с MS Office)
    try:
        import docx2pdf
        docx2pdf.convert(str(docx_path), str(out_path))
        if out_path.exists():
            docx_path.unlink(missing_ok=True)
            return out_path
    except Exception:
        pass

    # Пробуем конвертировать в PDF через LibreOffice
    lo_path = shutil.which("libreoffice") or shutil.which("soffice")
    if lo_path:
        try:
            subprocess.run([
                lo_path, "--headless", "--convert-to", "pdf",
                "--outdir", str(out_dir), str(docx_path)
            ], capture_output=True, timeout=30)
            gen_pdf = out_dir / "test_obfuscation_for_pdf.pdf"
            if gen_pdf.exists():
                gen_pdf.rename(out_path)
                docx_path.unlink(missing_ok=True)
                return out_path
        except Exception as e:
            print(f"LibreOffice PDF conversion failed: {e}")

    # Fallback: fpdf2 с Unicode-шрифтом (если доступен)
    for fp in [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        str(Path.home() / "Library/Fonts/Arial.ttf"),
    ]:
        font_path = Path(fp)
        if font_path.exists():
            try:
                from fpdf import FPDF
                pdf = FPDF()
                pdf.add_page()
                pdf.add_font("Custom", "", str(font_path), uni=True)
                pdf.set_font("Custom", "", 10)
                for line in DOCX_CONTENT.strip().split("\n"):
                    pdf.multi_cell(0, 8, line)
                pdf.output(str(out_path))
                docx_path.unlink(missing_ok=True)
                return out_path
            except Exception:
                pass

    # DOCX оставляем как fallback
    docx_path.rename(out_dir / "test_obfuscation_fallback.docx")
    return None


def main():
    out_dir = Path(__file__).parent / "test_data"
    out_dir.mkdir(parents=True, exist_ok=True)
    print("Создание тестовых файлов...")
    create_test_docx(out_dir)
    print("  - test_obfuscation.docx")
    create_test_excel(out_dir)
    print("  - test_obfuscation.xlsx")
    pdf_path = create_test_pdf(out_dir)
    if pdf_path:
        print("  - test_obfuscation.pdf")
    else:
        print("  - PDF не создан (нет LibreOffice/fpdf2)")
    print("Готово.")


if __name__ == "__main__":
    main()
