import os
import sys
import json
import shutil
import traceback
import tempfile
from pathlib import Path

# Add project root to path so we can import backend_api
PROJECT_ROOT = Path(__file__).parent.parent
sys.path.append(str(PROJECT_ROOT))
sys.path.insert(0, str(PROJECT_ROOT / "app"))

from backend_api import BackendApi

class TestRunner:
    def __init__(self):
        self.api = BackendApi()
        # Initialize Pullenti required by API
        if hasattr(self.api, '_init_pullenti'):
            self.api._init_pullenti()
        elif hasattr(self.api, '_pullenti_processor') and self.api._pullenti_processor is None:
            # Fallback for manual init
            import pullenti_wrapper.processor
            self.api._pullenti_processor = pullenti_wrapper.processor.Processor(["PERSON", "ORGANIZATION", "PHONE", "URI"])
        
        # Mock convert_to_pdf to avoid hangs on docx2pdf/Word COM
        import backend_api
        import pdf_convert
        def mock_convert_to_pdf(source, target):
            print(f"[MOCK] Converting {source} to {target}")
            import fitz
            document = fitz.open()
            page = document.new_page()
            page.insert_text((72, 72), "Anonymized test document")
            document.save(str(target))
            document.close()
            return True
        backend_api.convert_to_pdf = mock_convert_to_pdf
        
        # Mock the detailed OCR contract so the end-to-end test covers the
        # same page/bbox/confidence path as the UI without requiring hardware.
        def mock_ocr_pdf_to_text(pdf_path, **kwargs):
            print(f"[MOCK] OCR for {pdf_path}")
            return "Скан паспорта\nВыдан: Петров Петр Петрович\nНомер телефона: 8-800-555-35-35\nИНН: 7707083893"
        pdf_convert.ocr_pdf_to_text = mock_ocr_pdf_to_text
        def mock_pdf_to_text_auto_detailed(pdf_path, **kwargs):
            from ocr_backend import OCRDocumentResult, OCRPageResult, OCRTextRegion
            text = mock_ocr_pdf_to_text(pdf_path, **kwargs)
            region = OCRTextRegion(text, page=1, bbox=(0.1, 0.1, 0.9, 0.4), confidence=0.91)
            details = OCRDocumentResult(
                text=text,
                pages=(OCRPageResult(page=1, width=1200, height=1800, regions=(region,)),),
                backend="mock",
            )
            return text, "ocr", details
        backend_api.pdf_to_text_auto_detailed = mock_pdf_to_text_auto_detailed
        
        self.source_dir = PROJECT_ROOT / "tests" / "source_files"
        self.output_dir = PROJECT_ROOT / "tests" / "output_files"
        os.makedirs(self.output_dir, exist_ok=True)
        
        self.results = []
    
    def run_case(self, case_name, file_name, setup_api_fn, validator_fn):
        source_path = self.source_dir / file_name
        if not source_path.exists():
            print(f"❌ Файл {file_name} не найден. Пропуск.")
            self.results.append({"case": case_name, "file": file_name, "status": "SKIPPED", "error": "File not found"})
            return
            
        print(f"\n▶️ Запуск {case_name} для {file_name}...")
        
        # Setup API properties
        setup_api_fn(self.api)
        self.api.user_exclusions = set()
        self.api.custom_replacements = {}
        
        try:
            # Каждый сценарий работает с отдельной копией. Так тест
            # не считывает старые *_cleaned-файлы и не засоряет fixtures.
            with tempfile.TemporaryDirectory(prefix="docxdodyr-auto-") as folder:
                file_path = Path(folder) / file_name
                shutil.copy2(source_path, file_path)

                changes = self.api.process_single_file(
                    str(file_path), self.api.user_exclusions, self.api.custom_replacements
                )
                print(f"Обработано, замен: {changes}")
                validator_fn(self.api, file_path)
            
            print(f"✅ {case_name} ({file_name}) ПРОЙДЕН")
            self.results.append({"case": case_name, "file": file_name, "status": "PASSED", "error": None})
            
        except Exception as e:
            err_trace = traceback.format_exc()
            print(f"❌ {case_name} ({file_name}) УПАЛ:\n{err_trace}")
            self.results.append({"case": case_name, "file": file_name, "status": "FAILED", "error": str(e)})

# --- Validation helpers ---
def assert_no_pii_docx(filepath):
    from docx import Document
    doc = Document(filepath)
    full_text = " ".join([p.text for p in doc.paragraphs])
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text += " " + cell.text
    
    assert "Иванов Иван Иванович" not in full_text, "PII found in docx: Ivanov"
    assert "999" not in full_text and "123-45" not in full_text, "PII found in docx: Phone"
    assert "7707083893" not in full_text, "PII found in docx: INN"
    assert "ivanov@example.com" not in full_text, "PII found in docx: Email"
    assert "[ФИО_" in full_text or "[Наименование_" in full_text or "ФИО" in full_text, "No placeholders found in docx"

def assert_no_pii_xlsx(filepath):
    from openpyxl import load_workbook
    wb = load_workbook(filepath)
    full_text = ""
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell:
                    full_text += str(cell) + " "
    
    assert "Иванов Иван Иванович" not in full_text, "PII found in xlsx: Ivanov"
    assert "7707083893" not in full_text, "PII found in xlsx: INN"
    assert "test@example.invalid" not in full_text, "PII found in xlsx: Email"

def assert_pdf_exists(filepath):
    assert Path(filepath).exists(), f"PDF not generated: {filepath}"

def assert_decoder_exists_and_valid(filepath):
    assert Path(filepath).exists(), f"Decoder JSON not generated: {filepath}"
    with open(filepath, 'r', encoding='utf-8') as f:
        data = json.load(f)
    assert isinstance(data, dict), "Decoder must be a JSON dictionary"
    values = " ".join(
        value if isinstance(value, str) else str(value.get("original", ""))
        for value in data.values()
    )
    assert "Иванов Иван Иванович" in values or "Иванов" in values or "Иванович" in values, "Decoder does not contain expected PII"

# --- Test Cases ---
def run_all_tests():
    runner = TestRunner()
    
    # CASE A: Anonymize to original format
    def setup_a(api):
        api.save_original = True
        api.save_docx = False
        api.save_pdf = False
        api.save_markdown = False
        api.save_decoder = False
        api.ocr_pdf = False

    runner.run_case("Case A (DOCX -> DOCX)", "test_mock.docx", setup_a, 
                   lambda api, p: assert_no_pii_docx(str(p.parent / f"{p.stem}_cleaned.docx")))
    
    runner.run_case("Case A (XLSX -> XLSX)", "test_mock.xlsx", setup_a, 
                   lambda api, p: assert_no_pii_xlsx(str(next(
                       candidate for candidate in p.parent.glob("*.xlsx") if candidate != p
                   ))))

    # CASE B: Anonymize to PDF
    def setup_b(api):
        api.save_original = False
        api.save_pdf = True
        api.save_decoder = False
        api.ocr_pdf = False
        
    runner.run_case("Case B (DOCX -> PDF)", "test_mock.docx", setup_b, 
                   lambda api, p: assert_pdf_exists(str(p.parent / f"{p.stem}_cleaned.pdf")))

    # CASE C: Two-way anonymization
    def setup_c(api):
        api.save_original = True
        api.save_pdf = False
        api.save_decoder = True
        api.ocr_pdf = False
        
    def validate_c(api, p):
        cleaned_path = str(p.parent / f"{p.stem}_cleaned.docx")
        decoder_path = str(p.parent / f"{p.stem}_дешифратор.json")
        assert_no_pii_docx(cleaned_path)
        assert_decoder_exists_and_valid(decoder_path)
        
        # Test Restore
        print(f"Тестирование восстановления для {cleaned_path}")
        # Need to simulate run_restore from BackendApi
        api.restore_doc_paths = [cleaned_path]
        api.restore_json_paths = [decoder_path]
        
        # Override the JS evaluation so it doesn't fail if we don't have a window
        import builtins
        api._window = type('MockWindow', (), {'evaluate_js': lambda self, *args: None})()
        
        api._run_restore()
        
        restored_path = p.parent / f"{p.stem}_cleaned_восстановлено.docx"
        assert restored_path.exists(), "Восстановленный DOCX не создан"
        restored_text = " ".join(par.text for par in __import__("docx").Document(restored_path).paragraphs)
        assert "Иванов" in restored_text, "ФИО не восстановлено"
        print("✅ Файл успешно восстановлен")
            
    runner.run_case("Case C (DOCX Two-way)", "test_mock.docx", setup_c, validate_c)

    # CASE D: PDF Scan OCR
    def setup_d(api):
        api.save_original = True
        api.save_docx = True
        api.save_pdf = False
        api.save_decoder = False
        api.ocr_pdf = True  # OCR is enabled
        
    def validate_d(api, p):
        # Result should be DOCX
        cleaned_path = str(p.parent / f"{p.stem}_cleaned.docx")
        assert Path(cleaned_path).exists(), "DOCX after OCR not found"
        assert_no_pii_docx(cleaned_path)
        
    runner.run_case("Case D (PDF Scan OCR)", "test_mock_scan.pdf", setup_d, validate_d)

    # Print Summary
    print("\n\n" + "="*50)
    print("TEST REPORT SUMMARY")
    print("="*50)
    passed = sum(1 for r in runner.results if r["status"] == "PASSED")
    total = len(runner.results)
    
    for r in runner.results:
        icon = "✅" if r["status"] == "PASSED" else "❌" if r["status"] == "FAILED" else "⚠️"
        print(f"{icon} {r['case']} - {r['file']}")
        if r["error"]:
            print(f"   Ошибка: {r['error']}")
            
    print(f"\nВсего: {total}, Успешно: {passed}, Упало: {total - passed}")
    
    if passed < total:
        sys.exit(1)

if __name__ == "__main__":
    run_all_tests()
