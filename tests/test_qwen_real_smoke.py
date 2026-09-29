# -*- coding: utf-8 -*-
"""Opt-in real-model tests with pinned local weights; never downloads or receives source PII.

Validates:
1. Full SHA-256 integrity of pinned weights.
2. Real inference with positive grammar/punctuation edit (status == "ok").
3. End-to-end scenario: DOCX -> Anonymize -> Qwen -> Decoder -> Restore.
4. Negative safety tests rejecting fact/number changes and placeholder tampering.
"""
import json
import os
from pathlib import Path
import pytest
from docx import Document

from qwen_offline import validate_model
from qwen_postprocessor import (QwenPlaceholderPostprocessor, PlaceholderValidationError,
                                extract_placeholders, _FACT_NUMBER_RE, _MODALITY_RE)
from backend_api import BackendApi
from document_restorer import DocumentRestorer


_model_dir = os.environ.get("QWEN_MODEL_DIR", "").strip()
pytestmark = pytest.mark.skipif(not _model_dir, reason="QWEN_MODEL_DIR is not set; real model smoke is opt-in")


def _get_validated_model_path() -> Path:
    model_path = Path(_model_dir).expanduser().resolve()
    assert validate_model(model_path, verify_hash=True), "Pinned local weights failed SHA-256"
    return model_path


def test_real_local_qwen_sha256_integrity():
    """Verify SHA-256 integrity of all pinned weight files."""
    model_path = Path(_model_dir).expanduser().resolve()
    assert validate_model(model_path, verify_hash=True) is True, "Model hash validation failed"


def test_real_local_qwen_positive_grammar_edit(real_qwen_processor):
    """Verify safe grammar/punctuation correction succeeds with status == 'ok'."""
    model_path = _get_validated_model_path()
    processor = real_qwen_processor
    # Synthetic anonymized text with punctuation/whitespace defect
    input_text = "Уважаемый [ФИО_1] ,направляем Вам проект договора № [НомерДоговора_1] от 15.05.2025."
    result = processor.process(input_text)
    assert result != input_text and " ,направляем" not in result
    assert processor.last_status == "ok", f"Expected 'ok' status, got {processor.last_status}"
    assert "[ФИО_1]" in result
    assert "[НомерДоговора_1]" in result
    assert "15.05.2025" in result


def test_real_local_qwen_structural_writeback_preserves_docx_parts(tmp_path, real_qwen_processor):
    """Exercise real correction across split runs, grouped/table, header/footer."""
    base = "Уважаемый [ФИО_1] ,направляем Вам проект договора № [НомерДоговора_1] от 15.05.2025."
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("Уважаемый [ФИО_1] ")
    paragraph.add_run(",направляем Вам проект договора № [НомерДоговора_1] от 15.05.2025.")
    doc.add_paragraph(base)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = base.replace("ФИО_1", "ФИО_2").replace("НомерДоговора_1", "НомерДоговора_2")
    table.cell(0, 0).add_paragraph(base.replace("ФИО_1", "ФИО_3").replace("НомерДоговора_1", "НомерДоговора_3"))
    doc.sections[0].header.paragraphs[0].text = base.replace("ФИО_1", "ФИО_4").replace("НомерДоговора_1", "НомерДоговора_4")
    doc.sections[0].footer.paragraphs[0].text = base.replace("ФИО_1", "ФИО_5").replace("НомерДоговора_1", "НомерДоговора_5")
    def parts(document):
        return [*document.paragraphs, *document.tables[0].cell(0, 0).paragraphs,
                document.sections[0].header.paragraphs[0],
                document.sections[0].footer.paragraphs[0]]

    def invariants(text):
        return (extract_placeholders(text), _FACT_NUMBER_RE.findall(text),
                _MODALITY_RE.findall(text.casefold()))

    before = [p.text for p in parts(doc)]
    assert len(before) == 6
    api = BackendApi.__new__(BackendApi)
    api.qwen_settings = real_qwen_processor.settings
    api._qwen_postprocessor = real_qwen_processor
    api._write_back_qwen_document(doc, [])
    out = tmp_path / "structural.docx"
    doc.save(out)
    reopened = Document(out)
    after = [p.text for p in parts(reopened)]
    assert len(after) == len(before)
    for original, corrected in zip(before, after):
        assert corrected != original and " ,направляем" not in corrected
        assert invariants(corrected) == invariants(original)
    assert real_qwen_processor.last_status == "ok"


def test_real_local_qwen_e2e_docx_anonymize_qwen_restore(tmp_path, real_qwen_processor):
    """End-to-end pipeline: DOCX -> Anonymize -> Qwen postprocess -> Decoder -> Restore."""
    model_path = _get_validated_model_path()

    # 1. Create synthetic original DOCX
    orig_docx = tmp_path / "original_contract.docx"
    doc = Document()
    doc.add_paragraph("Генеральный директор ООО «СтройТехМонтаж» Кузнецов Василий Петрович заключил договор № 45-89.")
    doc.save(str(orig_docx))

    # 2. Anonymize using BackendApi (Pullenti)
    api = BackendApi(lazy_pullenti=False)
    api.qwen_settings = real_qwen_processor.settings
    api._qwen_postprocessor = real_qwen_processor
    api.save_decoder = True
    output_dir = tmp_path / "output"
    output_dir.mkdir()
    api.save_directory = str(output_dir)

    mapping = {}
    entity_seen = {}
    changes = api.process_single_file(
        str(orig_docx),
        exclusions_list=[],
        custom_replacements_list=[],
        batch_mapping=mapping,
        batch_entity_seen=entity_seen,
        output_dir=str(output_dir),
    )
    assert changes > 0, "No changes detected during anonymization"
    assert len(mapping) > 0, "Empty decoder mapping generated"

    cleaned_docx = output_dir / f"{orig_docx.stem}_cleaned.docx"
    assert cleaned_docx.exists(), "Cleaned docx not found"

    # Read cleaned text from anonymized docx
    cleaned_doc = Document(str(cleaned_docx))
    cleaned_text = "\n".join(p.text for p in cleaned_doc.paragraphs if p.text)
    assert "Кузнецов Василий Петрович" not in cleaned_text
    assert "СтройТехМонтаж" not in cleaned_text

    # Qwen already ran through the production structural write-back, before
    # final output hashing. Never edit the output behind its provenance.
    assert real_qwen_processor.last_status == "ok"

    # 4. Restore document using decoder mapping
    restorer = DocumentRestorer(mapping_dict=mapping)
    restored_docx = tmp_path / "restored_contract.docx"
    rest_success, rest_msg = restorer.restore_docx(str(cleaned_docx), str(restored_docx))
    assert rest_success is True, f"Document restoration failed: {rest_msg}"
    assert restored_docx.exists(), "Restored docx not created"

    # 5. Verify restored document content matches original facts
    restored_doc = Document(str(restored_docx))
    restored_text = "\n".join(p.text for p in restored_doc.paragraphs if p.text)
    assert "Кузнецов Василий Петрович" in restored_text
    assert "СтройТехМонтаж" in restored_text
    assert "45-89" in restored_text


def test_real_local_qwen_negative_fact_number_tamper():
    """Ensure fact number modification is rejected by safety validator."""
    from qwen_postprocessor import protect_placeholders, restore_and_validate
    protected = protect_placeholders("Акт № [НомерАкта_1] от 10.02.2025.")
    # Simulate bad model output changing year from 2025 to 2026
    tampered_output = f"Акт № <<QWENPH0>> от 10.02.2026."
    with pytest.raises(PlaceholderValidationError, match="model changed factual numbers"):
        restore_and_validate(tampered_output, protected)


def test_real_local_qwen_negative_placeholder_tamper():
    """Ensure sentinel modification or placeholder invention is rejected."""
    from qwen_postprocessor import protect_placeholders, restore_and_validate
    protected = protect_placeholders("Сотрудник [ФИО_1] приступил к работе.")
    # Simulate bad model output dropping the sentinel
    tampered_output = "Сотрудник приступил к работе."
    with pytest.raises(PlaceholderValidationError, match="model changed placeholder sentinels"):
        restore_and_validate(tampered_output, protected)
