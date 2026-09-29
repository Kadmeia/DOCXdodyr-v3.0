# -*- coding: utf-8 -*-
"""Fast, deterministic tests for the workflows a user performs in the UI.

These tests deliberately do not create a native WebView window, download a
Qwen model, or call an online service.  File dialogs, the WebView window and
the Pullenti result are all small local fakes; the DOCX/XLSX round-trip still
uses the real application pipeline.
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

import main
from backend_api import BackendApi, Worker
from document_restorer import DocumentRestorer
from qwen_postprocessor import QwenPlaceholderPostprocessor


class _FakeWindow:
    """Minimal pywebview window used for dialog and JavaScript assertions."""

    def __init__(self, selected=()):
        self.selected = tuple(str(path) for path in selected)
        self.dialog_calls = []
        self.js_calls = []

    def create_file_dialog(self, *args, **kwargs):
        self.dialog_calls.append((args, kwargs))
        return self.selected

    def evaluate_js(self, script):
        self.js_calls.append(script)


class _StubApi:
    """Backend substitute for tests that stop at the UI/API boundary."""

    def __init__(self):
        self.settings = {}
        self.qwen_settings = SimpleNamespace(enabled=False)
        self.saved_settings = 0
        self.qwen_updates = []
        self.processed_files = []
        self.processed_folders = []
        self.auto_decoder = True
        self.restore_json_paths = []
        self._window = None

    def set_window(self, window):
        self._window = window

    def _save_settings(self):
        self.saved_settings += 1

    def save_qwen_settings(self, data):
        self.qwen_updates.append(dict(data or {}))
        self.qwen_settings.enabled = bool((data or {}).get("enabled", False))

    def process_files(self, files):
        self.processed_files.append(list(files))

    def process_folder(self, folder):
        self.processed_folders.append(str(folder))


def test_user_can_launch_window_and_select_files_without_native_gui(monkeypatch, tmp_path):
    """The startup contract and file chooser filters are observable with mocks."""

    monkeypatch.setattr(main, "BackendApi", _StubApi)
    created = {}

    def fake_create_window(**kwargs):
        created.update(kwargs)
        return _FakeWindow()

    monkeypatch.setattr(main.webview, "create_window", fake_create_window)
    wrapper = main.ApiWrapper()
    window = main.webview.create_window(
        title="DOCXдодыр v2.0",
        url=str(Path(main.__file__).with_name("web") / "index.html"),
        js_api=wrapper,
        width=840,
        height=740,
        min_size=(780, 700),
        background_color="#020617",
    )
    wrapper.set_window(window)

    selected = tmp_path / "contract.docx"
    selected.touch()
    window.selected = (str(selected),)
    wrapper.open_file_dialog("anonymize_docs")

    assert created["title"] == "DOCXдодыр v2.0"
    assert created["js_api"] is wrapper
    assert created["width"] == 840 and created["height"] == 740
    assert window.dialog_calls[-1][1]["allow_multiple"] is True
    assert window.dialog_calls[-1][1]["file_types"][0].startswith("Все поддерживаемые")
    assert "*.docx" in window.dialog_calls[-1][1]["file_types"][0]
    assert "*.jpg" in window.dialog_calls[-1][1]["file_types"][0]
    assert wrapper._api.processed_files == [[str(selected)]]


def test_user_settings_and_restore_file_selection_are_forwarded(monkeypatch, tmp_path):
    monkeypatch.setattr(main, "BackendApi", _StubApi)
    wrapper = main.ApiWrapper()
    document = tmp_path / "contract_cleaned.docx"
    decoder = tmp_path / "contract_дешифратор.json"
    document.touch()
    decoder.touch()
    window = _FakeWindow([document, decoder])
    wrapper.set_window(window)

    wrapper.update_settings(
        {
            "save_original": True,
            "save_pdf": False,
            "irreversible_pdf": True,
            "save_decoder": True,
            "ocr_lang": "eng",
            "qwen_enabled": True,
        }
    )
    assert wrapper._api.save_original is True
    assert wrapper._api.save_decoder is True
    assert wrapper._api.irreversible_pdf is True
    assert wrapper._api.ocr_lang == "eng"
    assert wrapper._api.qwen_settings.enabled is True
    assert wrapper._api.saved_settings == 1

    wrapper.open_file_dialog("restore_doc")
    assert wrapper._api.restore_doc_paths == [str(document), str(decoder)]
    assert "lbl-restore-doc" in window.js_calls[-1]
    assert window.dialog_calls[-1][1]["file_types"] == ("Документы (*.docx;*.docm;*.xlsx;*.xlsm)",)

    wrapper.open_file_dialog("restore_json")
    assert wrapper._api.restore_json_paths == [str(document), str(decoder)]
    assert "lbl-restore-json" in window.js_calls[-1]
    assert window.dialog_calls[-1][1]["file_types"] == ("Дешифратор (*.json)",)


def test_user_can_select_a_folder_for_batch_anonymization(monkeypatch, tmp_path):
    """The folder picker forwards one root and advertises its output folder."""

    monkeypatch.setattr(main, "BackendApi", _StubApi)
    wrapper = main.ApiWrapper()
    source = tmp_path / "case-files"
    source.mkdir()
    window = _FakeWindow([source])
    wrapper.set_window(window)

    started = []

    class _FakeThread:
        def __init__(self, *, target, daemon):
            self.target = target
            self.daemon = daemon

        def start(self):
            started.append(self)

    monkeypatch.setattr(main.threading, "Thread", _FakeThread)
    wrapper.open_folder_dialog("anonymize_folder")

    assert len(started) == 1
    assert started[0].daemon is True
    # Run the target explicitly so the test remains deterministic while still
    # proving that the bridge returned after starting a background job.
    started[0].target()
    assert wrapper._api.processed_folders == [str(source)]
    assert window.dialog_calls[-1][1] == {}
    assert any("setFolderStatus" in call for call in window.js_calls)
    assert any("Обезличенные документы" in call for call in window.js_calls)


def test_user_can_continue_existing_folder_set(monkeypatch, tmp_path):
    class _ContinueApi(_StubApi):
        def process_folder(self, folder, **kwargs):
            self.processed_folders.append((str(folder), dict(kwargs)))

    monkeypatch.setattr(main, "BackendApi", _ContinueApi)
    wrapper = main.ApiWrapper()
    source = tmp_path / "case-files"
    source.mkdir()
    window = _FakeWindow([source])
    wrapper.set_window(window)

    started = []

    class _FakeThread:
        def __init__(self, *, target, daemon):
            self.target, self.daemon = target, daemon

        def start(self):
            started.append(self)

    monkeypatch.setattr(main.threading, "Thread", _FakeThread)
    wrapper.open_folder_dialog("continue_folder")
    started[0].target()

    folder, options = wrapper._api.processed_folders[0]
    assert folder == str(source)
    assert options["continue_existing"] is True


def test_auto_decoder_setting_is_forwarded_and_can_clear_previous_selection(monkeypatch):
    monkeypatch.setattr(main, "BackendApi", _StubApi)
    wrapper = main.ApiWrapper()
    wrapper.set_restore_auto_decoder(False)
    assert wrapper._api.auto_decoder is False
    assert wrapper._api.settings["auto_decoder"] is False
    wrapper._api.restore_json_paths = ["old.json"]
    wrapper.clear_restore_decoder()
    assert wrapper._api.restore_json_paths == []


class _Occurrence:
    def __init__(self, source, start, end):
        self.sofa = source
        self.begin_char = start
        self.end_char = end - 1  # Pullenti's public end is inclusive.

    def get_text(self):
        return self.sofa.text[self.begin_char : self.end_char + 1]


class _PersonReferent:
    type_name = "PERSON"

    def __init__(self, source, start, end):
        self.occurrence = [_Occurrence(source, start, end)]


class _LocalPersonProcessor:
    """Deterministic Pullenti-shaped result for the two local test names."""

    _name = re.compile(
        r"Иванов(?:у|ым)?\s+Иван(?:у|ом)?\s+Иванович(?:у|ем)?",
        re.IGNORECASE,
    )

    def process(self, source, *_args):
        return SimpleNamespace(
            entities=[
                _PersonReferent(source, match.start(), match.end())
                for match in self._name.finditer(source.text)
            ]
        )


def _offline_backend():
    """Create BackendApi with all heavyweight/network paths disabled."""

    api = BackendApi.__new__(BackendApi)
    api._pullenti_processor = _LocalPersonProcessor()
    api.current_placeholders = {"PER": "[ФИО]"}
    api.qwen_settings = SimpleNamespace(enabled=False)
    api._qwen_postprocessor = SimpleNamespace(last_status="disabled")
    api.user_exclusions = set()
    api.custom_replacements = set()
    api.save_original = True
    api.save_pdf = False
    api.save_decoder = True
    api.ocr_pdf = False
    api._window = None
    return api


def _make_batch_files(folder: Path):
    docx_path = folder / "first.docx"
    doc = Document()
    doc.add_paragraph("Подписант: Иванов Иван Иванович.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "ФИО"
    table.cell(0, 1).text = "Иванов Иван Иванович"
    doc.save(docx_path)

    xlsx_path = folder / "second.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Анкета"
    sheet["A1"] = "ФИО"
    sheet["B1"] = "Иванову Ивану Ивановичу"
    sheet["A2"] = "Комментарий"
    sheet["B2"] = "Документ направлен Иванову Ивану Ивановичу"
    workbook.save(xlsx_path)
    return docx_path, xlsx_path


def test_user_batch_anonymizes_docx_and_xlsx_with_one_decoder_id(tmp_path):
    """A selected DOCX+XLSX batch preserves IDs and table labels."""

    docx_path, xlsx_path = _make_batch_files(tmp_path)
    api = _offline_backend()
    Worker([str(docx_path), str(xlsx_path)], api).run()

    cleaned_docx = tmp_path / "first_cleaned.docx"
    cleaned_xlsx = tmp_path / "second_cleaned.xlsx"
    from document_restorer import find_decoder_near_document
    decoder_docx = find_decoder_near_document(cleaned_docx)
    decoder_xlsx = find_decoder_near_document(cleaned_xlsx)
    assert cleaned_docx.exists() and cleaned_xlsx.exists()
    assert decoder_docx.exists() and decoder_xlsx.exists()

    cleaned_text = " ".join(p.text for p in Document(cleaned_docx).paragraphs)
    assert "Иванов Иван Иванович" not in cleaned_text
    assert "[ФИО_1]" in cleaned_text
    cleaned_table = Document(cleaned_docx).tables[0]
    assert cleaned_table.cell(0, 0).text == "ФИО"
    assert cleaned_table.cell(0, 1).text == "[ФИО_1]"

    sheet = load_workbook(cleaned_xlsx).active
    assert sheet["A1"].value == "ФИО"
    assert sheet["B1"].value == "[ФИО_1]"
    assert "[ФИО_1]" in sheet["B2"].value

    first_map = json.loads(decoder_docx.read_text(encoding="utf-8"))
    second_map = json.loads(decoder_xlsx.read_text(encoding="utf-8"))
    assert first_map == second_map
    assert [k for k in first_map if not k.startswith("__")] == ["[ФИО_1]"]
    assert first_map["[ФИО_1]"]["id"] == "PER_000001"

    restored_docx = tmp_path / "restored.docx"
    ok, _message = DocumentRestorer(first_map).restore_docx(cleaned_docx, restored_docx)
    assert ok
    restored_text = " ".join(p.text for p in Document(restored_docx).paragraphs)
    assert "Иванов Иван Иванович" in restored_text

    restored_xlsx = tmp_path / "restored.xlsx"
    ok, _message = DocumentRestorer(second_map).restore_excel(cleaned_xlsx, restored_xlsx)
    assert ok
    restored_sheet = load_workbook(restored_xlsx).active
    assert "Иванов" in restored_sheet["B1"].value
    assert "Иванов" in restored_sheet["B2"].value


class _InspectingModel:
    def __init__(self):
        self.prompt = None

    def generate(self, prompt, *, max_new_tokens, temperature):
        self.prompt = prompt
        document = prompt.split("ДОКУМЕНТ:\n", 1)[1]
        return document.replace("подготовлен", "подготовлен и проверен")


def test_qwen_postprocess_receives_only_anonymized_text_and_keeps_tokens():
    """The local fake model cannot see decoder values and may only edit grammar."""

    model = _InspectingModel()
    processor = QwenPlaceholderPostprocessor(
        {"enabled": True, "local_files_only": True},
        model_loader=lambda _settings: model,
    )
    source = "Для [ФИО_1] подготовлен документ № 17."
    result = processor.process(source)

    assert "Иванов" not in (model.prompt or "")
    assert "<<QWENPH0>>" in (model.prompt or "") or "\ue000QWENPH0\ue001" in (model.prompt or "")
    assert result == "Для [ФИО_1] подготовлен и проверен документ № 17."
    assert "[ФИО_1]" in result and "№ 17" in result


def test_qwen_model_failure_is_a_safe_noop_without_download(monkeypatch):
    """A missing local fixture returns the anonymized input unchanged."""

    calls = []

    def no_download_loader(settings):
        calls.append(settings.model_id)
        raise FileNotFoundError("local fixture is absent")

    processor = QwenPlaceholderPostprocessor(
        {"enabled": True, "local_files_only": True}, model_loader=no_download_loader
    )
    source = "Для [ФИО_1] подготовлен документ."
    assert processor.process(source) == source
    assert processor.last_status.startswith("unavailable:")
    assert calls == ["Qwen/Qwen3.5-0.8B"]


def test_review_ui_contract_uses_safe_dom_rendering():
    web_root = Path(main.__file__).with_name("web")
    html = (web_root / "index.html").read_text(encoding="utf-8")
    script = (web_root / "script.js").read_text(encoding="utf-8")
    assert 'id="review-overlay"' in html
    assert 'id="btn-review-findings"' in html
    assert "get_review_findings" in script
    assert "decide_review_finding" in script
    assert "context.textContent" in script
    assert "context.innerHTML" not in script


def test_worker_respects_open_output_folder_setting(tmp_path, monkeypatch):
    """Проверяет, что при обработке отдельных файлов папка открывается, если настройка включена, и не открывается при выключенной."""
    import app_paths
    docx_path, xlsx_path = _make_batch_files(tmp_path)
    api = _offline_backend()

    opened_folders = []
    monkeypatch.setattr(app_paths, "open_folder_in_file_manager", lambda p: opened_folders.append(Path(p)) or True)

    # 1. При включенной настройке (по умолчанию)
    api.open_output_folder = True
    Worker([str(docx_path)], api).run()
    assert len(opened_folders) == 1
    assert opened_folders[0].resolve() == tmp_path.resolve()

    # 2. При отключенной настройке
    opened_folders.clear()
    api.open_output_folder = False
    Worker([str(xlsx_path)], api).run()
    assert len(opened_folders) == 0
