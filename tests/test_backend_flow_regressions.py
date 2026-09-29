# -*- coding: utf-8 -*-
"""Regression tests for folder output and decoder trust boundaries."""

import json

import pytest
from docx import Document

from decoder_binding import publish_binding
from document_restorer import restore_document
from folder_pipeline import collect_supported_files
from backend_api import BackendApi


def test_collect_supported_files_rejects_unsafe_output_name(tmp_path):
    with pytest.raises(ValueError):
        collect_supported_files(tmp_path, output_dir_name="../outside")


def test_publish_binding_rejects_output_outside_decoder_directory(tmp_path):
    decoder_dir = tmp_path / "decoder"
    decoder_dir.mkdir()
    decoder = decoder_dir / "дешифратор.json"
    output_dir = tmp_path / "output"
    output_dir.mkdir()
    output = output_dir / "document.docx"
    output.write_bytes(b"not a document")
    decoder.write_text(json.dumps({"[ФИО_1]": "Иванов"}), encoding="utf-8")

    with pytest.raises(ValueError, match="below the decoder directory"):
        publish_binding(decoder, [output], "run-1", emit_sidecars=True)


def test_restore_document_rejects_same_mapping_from_unbound_decoder(tmp_path):
    document = tmp_path / "cleaned.docx"
    doc = Document()
    doc.add_paragraph("[ФИО_1]")
    doc.save(document)

    mapping = {"[ФИО_1]": "Иванов Иван"}
    bound_decoder = document.with_name("cleaned_дешифратор.json")
    bound_decoder.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
    publish_binding(bound_decoder, [document], "run-1", emit_sidecars=True)

    unrelated_decoder = tmp_path / "unrelated.json"
    unrelated_decoder.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")

    with pytest.raises(ValueError, match="не привязан к документу"):
        restore_document(document, decoder_path=unrelated_decoder)


def test_anonymization_fails_closed_when_pullenti_is_unavailable():
    api = BackendApi.__new__(BackendApi)
    api._pullenti_processor = None

    with pytest.raises(RuntimeError, match="обезличивание остановлено"):
        api.anonymize_text_pullenti("Иванов Иван Иванович")


def test_hyperlink_not_destroyed_by_replacement():
    """Плейсхолдер внутри w:hyperlink не должен обнулять URL."""
    from docx import Document
    from docx.oxml.ns import qn
    import backend_api

    doc = Document()
    para = doc.add_paragraph()
    # Создаём гиперссылку программно
    hyperlink = para._p.makeelement(qn("w:hyperlink"), {})
    run = para._p.makeelement(qn("w:r"), {})
    t = para._p.makeelement(qn("w:t"), {})
    t.text = "example.com"
    run.append(t)
    hyperlink.append(run)
    para._p.append(hyperlink)
    # Обычный текст
    normal_run = para.add_run(" контакт Иванов И.И.")

    replacements = [(9, 20, "[ФИО]")]
    backend_api.replace_spans_in_paragraph_xml(para, replacements)

    # Проверяем, что текст внутри hyperlink не изменён
    hl_texts = [
        node.text for node in hyperlink.iter(qn("w:t")) if node.text
    ]
    assert "example.com" in " ".join(hl_texts)
    assert "[ФИО]" in para.text

