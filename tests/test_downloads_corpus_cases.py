# -*- coding: utf-8 -*-
"""Regression tests using synthetic legal document examples only.

Identifiers, names and dates were replaced during the 2026-09-16 privacy
review. They are invented test values, not records from a source document;
coincidental overlap with public identifiers is possible. No source documents
are opened by this module. Preserve the formats and negative cases below:
- Complex multi-segment court case numbers in arbitral decisions
- Rospatent software registration certificates and applications
- False positive elimination for "без доверенности", visa text, and job title signature prompts
- Contract number variations (ГК, контракты, договоры)
"""

from pullenti_legal.analyzer import iter_candidates
from backend_api import BackendApi
from legal_pullenti import initialize_ner


def test_court_case_numbers_with_russian_prefixes_and_subdivisions():
    # Synthetic arbitral ruling retaining a multi-segment case format.
    text1 = "Решение Арбитражного суда г. Москвы от 17.02.2026 по делу N А40-900731/26-908-9076"
    cands1 = [c for c in iter_candidates(text1) if c[0] == "COURT_CASE_NUMBER"]
    assert len(cands1) == 1
    assert cands1[0][1] == "А40-900731/26-908-9076"
    assert text1[cands1[0][2]:cands1[0][3]] == "А40-900731/26-908-9076"

    # Variations with № and different case declensions
    text2 = "рассмотрев в судебном заседании дело № 2-9073/2026"
    cands2 = [c for c in iter_candidates(text2) if c[0] == "COURT_CASE_NUMBER"]
    assert len(cands2) == 1
    assert cands2[0][1] == "2-9073/2026"

    text3 = "по административному делу № А56-90731/2026"
    cands3 = [c for c in iter_candidates(text3) if c[0] == "COURT_CASE_NUMBER"]
    assert len(cands3) == 1
    assert cands3[0][1] == "А56-90731/2026"


def test_rospatent_software_certificates_and_applications():
    text1 = "Свидетельство о государственной регистрации программы для ЭВМ № 2026990731"
    cands1 = [c for c in iter_candidates(text1) if c[0] == "SOFTWARE_REGISTRATION_NUMBER"]
    assert len(cands1) == 1
    assert cands1[0][1] == "2026990731"

    text2 = "Заявка № 2026990732\nДата поступления 17 февраля 2026 г."
    cands2 = [c for c in iter_candidates(text2) if c[0] == "PATENT_APPLICATION_NUMBER"]
    assert len(cands2) == 1
    assert cands2[0][1] == "2026990732"

    text3 = "зарегистрировано в Реестре программ для ЭВМ № 2026990733"
    cands3 = [c for c in iter_candidates(text3) if c[0] == "SOFTWARE_REGISTRATION_NUMBER"]
    assert len(cands3) == 1
    assert cands3[0][1] == "2026990733"


def test_power_of_attorney_false_positive_elimination_and_valid_cases():
    # Extract from EGRUL: "без доверенности действовать" must NOT trigger POA candidate
    egrul_text = "Лицо, имеющее право без доверенности действовать от имени юридического лица"
    poa_cands = [c for c in iter_candidates(egrul_text) if c[0] == "POWER_OF_ATTORNEY_NUMBER"]
    assert poa_cands == []

    # Valid POA with number and date
    valid_poa = "на основании доверенности № 907/26 от 18.02.2026"
    cands_valid = [c for c in iter_candidates(valid_poa) if c[0] == "POWER_OF_ATTORNEY_NUMBER"]
    assert len(cands_valid) == 1
    assert cands_valid[0][1] == "907/26"


def test_visa_number_false_positive_elimination_and_valid_cases():
    # Common contract text containing genitive/plural forms of "виза"
    contract_text = "по согласованию визы любому клиенту, оформление визы директором"
    visa_cands = [c for c in iter_candidates(contract_text) if c[0] == "VISA_NUMBER"]
    assert visa_cands == []

    # Valid visa number
    valid_visa = "Номер визы: 907319073"
    cands_valid = [c for c in iter_candidates(valid_visa) if c[0] == "VISA_NUMBER"]
    assert len(cands_valid) == 1
    assert cands_valid[0][1] == "907319073"


def test_job_title_signature_prompts_sanitization():
    # Job title followed by signature placeholders
    text1 = "Должность: Генеральный директор (подпись) (расшифровка подписи)"
    cands1 = [c for c in iter_candidates(text1) if c[0] == "JOB_TITLE"]
    assert len(cands1) == 1
    assert cands1[0][1] == "ГЕНЕРАЛЬНЫЙ ДИРЕКТОР"
    assert text1[cands1[0][2]:cands1[0][3]] == "Генеральный директор"

    # Bare signature line placeholder should be skipped
    text2 = "Должность: ______________ (подпись) (Ф.И.О.)"
    cands2 = [c for c in iter_candidates(text2) if c[0] == "JOB_TITLE"]
    assert cands2 == []


def test_contract_numbers_with_gk_and_prefixes():
    text1 = "ТЗ Приложение №1 к ГК 90731_ЭК"
    cands1 = [c for c in iter_candidates(text1) if c[0] == "CONTRACT_NUMBER"]
    assert len(cands1) == 1
    assert cands1[0][1] == "90731_ЭК"

    text2 = "Договор № 907-73 от 19.02.2026"
    cands2 = [c for c in iter_candidates(text2) if c[0] == "CONTRACT_NUMBER"]
    assert len(cands2) == 1
    assert cands2[0][1] == "907-73"


def test_end_to_end_anonymization_of_court_decision():
    initialize_ner()
    api = BackendApi()

    court_decision_sample = (
        "Решение Арбитражного суда г. Москвы от 17.02.2026 по делу N А40-900731/26-908-9076.\n"
        "Заявитель: ООО «Синтетический Пример», ИНН: 9900907310, ОГРН: 1269900907319.\n"
        "Представитель по доверенности № 907/26: Орлов Роман Петрович.\n"
        "Судья: Р.П. Орлов."
    )

    cleaned, repl_count, logs = api.anonymize_text_pullenti(court_decision_sample)

    assert "А40-900731/26-908-9076" not in cleaned
    assert "9900907310" not in cleaned
    assert "1269900907319" not in cleaned
    assert "907/26" not in cleaned
    assert "Орлов Роман Петрович" not in cleaned
    assert repl_count >= 5


def test_inn_kpp_pairs_with_role_headers():
    text = "ИНН/КПП покупателя: 9900907327 / 990001001"
    inns = [c for c in iter_candidates(text) if c[0] == "INN"]
    kpps = [c for c in iter_candidates(text) if c[0] == "KPP"]
    assert len(inns) == 1 and inns[0][1] == "9900907327"
    assert len(kpps) == 1 and kpps[0][1] == "990001001"


def test_bik_with_bank_and_branch_qualifiers():
    for phrase in [
        "БИК банка: 049990731",
        "БИК отделения 049990732",
        "БИК Банка: 049990733",
        "БИК ТОФК: 049990734",
    ]:
        biks = [c for c in iter_candidates(phrase) if c[0] == "BIK"]
        assert len(biks) == 1, f"Failed for {phrase}"


def test_bank_and_treasury_accounts_with_slash_and_latin_chars():
    c1 = [c for c in iter_candidates("к/счет 30101810999990090731") if c[0] == "RU_CORR_ACCOUNT"]
    assert len(c1) == 1 and c1[0][1] == "30101810999990090731"

    c2 = [c for c in iter_candidates("К/c: 30101810999990090732") if c[0] == "RU_CORR_ACCOUNT"]
    assert len(c2) == 1 and c2[0][1] == "30101810999990090732"

    a1 = [c for c in iter_candidates("р\\с 40702810999990090731") if c[0] == "RU_ACCOUNT"]
    assert len(a1) == 1 and a1[0][1] == "40702810999990090731"

    a2 = [c for c in iter_candidates("Счет банка получателя: 40102810999990090732") if c[0] == "RU_ACCOUNT"]
    assert len(a2) == 1 and a2[0][1] == "40102810999990090732"


def test_egrip_and_egrul_registration_records():
    egrip_text = "о чем внесена запись в ЕГРИП 326990090731905"
    cands = [c for c in iter_candidates(egrip_text) if c[0] == "OGRNIP"]
    assert len(cands) == 1 and cands[0][1] == "326990090731905"


def test_snils_full_official_phrasing():
    snils_text = "страхового номера индивидуального лицевого счета | 900-000-001 82"
    cands = [c for c in iter_candidates(snils_text) if c[0] == "SNILS"]
    assert len(cands) == 1 and cands[0][1] == "90000000182"


def test_person_role_prefix_preservation():
    initialize_ner()
    api = BackendApi()

    text = (
        "Договор заключен в лице Генерального директора Соколова Антона Сергеевича с одной стороны, "
        "и Индивидуальный предприниматель Морозов Павел Викторович с другой стороны.\n"
        "Судья Р.П. Орлов\n"
        "Генеральный директор _______________________ /Соколов А. С."
    )
    cleaned, count, _ = api.anonymize_text_pullenti(text)

    assert "Генерального директора [ФИО" in cleaned
    assert "Индивидуальный предприниматель [ФИО" in cleaned
    assert "Судья [ФИО" in cleaned
    assert "Генеральный директор _______________________ /[ФИО" in cleaned
    assert "Соколова Антона Сергеевича" not in cleaned
    assert "Морозов Павел Викторович" not in cleaned
    assert "Орлов" not in cleaned
    assert "Орлов" not in cleaned
    assert "Соколов А. С." not in cleaned


def test_tracked_revisions_with_newlines_are_safely_anonymized():
    import docx
    from docx.oxml import parse_xml
    initialize_ner()
    api = BackendApi()

    doc = docx.Document()
    p = doc.add_paragraph()
    # Имитируем режим рецензирования Word с w:ins и внутренним переносом строки
    ins_xml = parse_xml(
        '<w:ins xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:id="100" w:author="Юрист">'
        '<w:r><w:t xml:space="preserve">ПОДПИСИ СТОРОН:\nОт Туроператора: Орлов Роман Петрович\nЗаказчик: Белов Олег Сергеевич</w:t></w:r>'
        '</w:ins>'
    )
    p._element.append(ins_xml)

    total_repl, logs = api.clean_document(doc, [], [], {})
    assert total_repl >= 2
    final_text = api.get_paragraph_text_with_revisions(p)
    assert "Орлов Роман Петрович" not in final_text
    assert "Белов Олег Сергеевич" not in final_text
    assert "[ФИО]" in final_text
    assert "\n" in final_text


def test_tour_operator_efrt_registry_numbers_are_anonymized():
    from pullenti_legal.analyzer import iter_candidates
    from table_entity_context import TableCellContext

    initialize_ner()
    api = BackendApi()

    # Проверка текста с префиксом ЕФРТ -> [НомерРТО]
    text1 = "Реестровый номер в ЕФРТ: РТО 990731 "
    res1, _, _ = api.anonymize_text_pullenti(text1)
    assert "990731" not in res1
    assert "РТО 990731" not in res1
    assert "Реестровый номер в ЕФРТ: [НомерРТО] " == res1

    # Проверка табличного контекста (колонка/строка 'Реестровый номер')
    ctx = TableCellContext(row=0, column=1, value_text="РТО 990732", label="Реестровый номер", label_row=0, label_column=0)
    res2, _, _ = api.anonymize_text_pullenti("РТО 990732", semantic_context=ctx)
    assert "990732" not in res2
    assert "РТО 990732" not in res2
    assert res2 == "[НомерРТО]"

    # Проверка серий в тексте туроператоров
    candidates = list(iter_candidates("Туроператор ООО «Синтетический Тур» (РТО 990731), лицензия МВТ 990733"))
    rto_nums = [val for label, val, _, _ in candidates if label == "TOUR_OPERATOR_REGISTRY_NUMBER"]
    assert "РТО 990731" in rto_nums
    assert "МВТ 990733" in rto_nums

    # Проверка защиты от ложных срабатываний (техническое РТО, товарооборот, закупки, законы)
    negatives = [
        "проведение РТО 1 раз в квартал",
        "график РТО оборудования на 2026 год",
        "РТО составил 1500000 руб.",
        "РТО за месяц 250000 руб",
        "план РТО 50000 шт",
        "реестровый номер закупки 0999900907326000001",
        "реестровый номер адвоката 99/9073",
        "в соответствии со ст. 4.1 Федерального закона от 24.11.1996 № 132-ФЗ",
        "карточка спортсмена",
    ]
    for neg in negatives:
        cands = [val for label, val, _, _ in iter_candidates(neg) if label == "TOUR_OPERATOR_REGISTRY_NUMBER"]
        assert not cands, f"Ложное срабатывание на тексте: {neg}"
