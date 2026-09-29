# -*- coding: utf-8 -*-
"""Автоматические тесты матрицы форматов сохранения, чистой папки, сверки XLSM и безопасного восстановления."""

from __future__ import annotations

import json
import os
from pathlib import Path
import sys
from types import SimpleNamespace
import zipfile
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw

from backend_api import BackendApi
from document_restorer import (
    DocumentRestorer,
    find_decoder_near_document,
    restore_document,
)
from folder_pipeline import (
    DEFAULT_DECODER_FILENAME,
    DEFAULT_OUTPUT_DIR_NAME,
    FolderAnonymizationPipeline,
)
from pdf_convert import convert_docx_text_to_pdf


class _SimpleMockPullenti:
    """Mock-процессор Pullenti для детерминированных тестов."""

    def __init__(self, entities_map=None):
        self.entities_map = entities_map or {
            "Иванов Иван Иванович": "PERSON",
            "Иванова Ивана Ивановича": "PERSON",
            "Петров Петр Петрович": "PERSON",
            "Сидоров Сидор Сидорович": "PERSON",
            "ООО «Вектор»": "ORGANIZATION",
            "ООО «Ромашка»": "ORGANIZATION",
            "40702810302410001062": "RU_ACCOUNT",
            "7701234567": "INN",
        }

    def process(self, source, *_args):
        text = getattr(source, "text", "") or str(source)
        entities = []
        for name, type_name in self.entities_map.items():
            start = 0
            while True:
                idx = text.find(name, start)
                if idx < 0:
                    break
                end = idx + len(name)
                occ = SimpleNamespace(sofa=source, begin_char=idx, end_char=end - 1)
                referent = SimpleNamespace(type_name=type_name, occurrence=[occ])
                entities.append(referent)
                start = end
        return SimpleNamespace(entities=entities)


def _create_mock_backend(monkeypatch, custom_entities=None):
    import pdf_convert
    monkeypatch.setattr("backend_api.convert_to_pdf", pdf_convert.convert_to_pdf)
    api = BackendApi.__new__(BackendApi)
    api._pullenti_processor = _SimpleMockPullenti(custom_entities)
    api.current_placeholders = {
        "PER": "[ФИО]",
        "ORG": "[Организация]",
        "RU_ACCOUNT": "[Р/с]",
        "INN": "[ИНН]",
    }
    api.qwen_settings = SimpleNamespace(enabled=False)
    api._qwen_postprocessor = SimpleNamespace(last_status="disabled")
    api.user_exclusions = set()
    api.custom_replacements = set()
    api.save_original = False
    api.save_docx = False
    api.save_pdf = False
    api.save_markdown = False
    api.save_decoder = False
    api.ocr_pdf = False
    api._window = None
    api.emit_audit_sidecars = False
    api.save_log_files = False
    api.save_audit_files = False
    api.open_output_folder = False
    api._last_reconcile_temp = None
    api._last_generated_outputs = []
    api.anonymize_cache = {}

    def _mock_ocr_image_to_result(path, **kwargs):
        from ocr_backend import OCRDocumentResult, OCRPageResult, OCRTextRegion
        return OCRDocumentResult(
            text="ООО «Ромашка» и Иванов Иван Иванович",
            pages=(OCRPageResult(page=1, width=800, height=200, regions=(
                OCRTextRegion("ООО «Ромашка»", 1, (0.1, 0.1, 0.5, 0.2), 0.99),
                OCRTextRegion("Иванов Иван Иванович", 1, (0.1, 0.3, 0.8, 0.5), 0.99),
            )),),
            backend="mock_deterministic",
        )
    monkeypatch.setattr("ocr_backend.ocr_image_to_result", _mock_ocr_image_to_result)
    return api


def _create_sample_file(path: Path, text: str = "Договор подписал Иванов Иван Иванович"):
    suffix = path.suffix.lower()
    path.parent.mkdir(parents=True, exist_ok=True)
    if suffix in {".docx", ".docm"}:
        doc = Document()
        doc.add_paragraph(text)
        if suffix == ".docx":
            doc.save(str(path))
        else:
            base = path.with_suffix(".docx")
            doc.save(str(base))
            with zipfile.ZipFile(base) as source_zip:
                parts = {name: source_zip.read(name) for name in source_zip.namelist()}
            parts["word/vbaProject.bin"] = b"TEST_VBA_PAYLOAD"
            rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
            rels = ET.fromstring(parts["word/_rels/document.xml.rels"])
            ET.SubElement(rels, f"{{{rel_ns}}}Relationship", {
                "Id": "rIdTestVba", "Type": "http://schemas.microsoft.com/office/2006/relationships/vbaProject",
                "Target": "vbaProject.bin",
            })
            parts["word/_rels/document.xml.rels"] = ET.tostring(rels, encoding="utf-8")
            ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
            types = ET.fromstring(parts["[Content_Types].xml"])
            ET.SubElement(types, f"{{{ct_ns}}}Default", {
                "Extension": "bin", "ContentType": "application/vnd.ms-office.vbaProject",
            })
            for node in types:
                if node.attrib.get("PartName") == "/word/document.xml":
                    node.set("ContentType", "application/vnd.ms-word.document.macroEnabled.main+xml")
            parts["[Content_Types].xml"] = ET.tostring(types, encoding="utf-8")
            with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as output_zip:
                for name, data in parts.items():
                    output_zip.writestr(name, data)
            base.unlink()
    elif suffix in {".xlsx", ".xlsm"}:
        wb = Workbook()
        ws = wb.active
        ws["A1"] = text
        wb.save(str(path))
    elif suffix == ".pdf":
        docx_temp = path.with_suffix(".docx")
        doc = Document()
        doc.add_paragraph(text)
        doc.save(str(docx_temp))
        convert_docx_text_to_pdf(docx_temp, path)
        docx_temp.unlink(missing_ok=True)
    elif suffix in {".png", ".jpg", ".jpeg"}:
        img = Image.new("RGB", (800, 200), color=(255, 255, 255))
        draw = ImageDraw.Draw(img)
        draw.text((10, 10), text, fill=(0, 0, 0))
        img.save(str(path))
    return path


# =========================================================================
# 1. Матрица форматов по умолчанию (все переключатели выключены)
# =========================================================================
@pytest.mark.parametrize("ext,expected_ext", [
    (".docx", ".docx"),
    (".docm", ".docm"),
    (".xlsx", ".xlsx"),
    (".xlsm", ".xlsm"),
    (".pdf", ".pdf"),
    (".png", ".pdf"),  # Изображения по умолчанию строго в PDF
])
def test_default_formats_matrix_single_file(tmp_path, monkeypatch, ext, expected_ext):
    api = _create_mock_backend(monkeypatch)
    src = _create_sample_file(tmp_path / f"sample{ext}")
    out_dir = tmp_path / f"out_{ext.strip('.')}"
    out_dir.mkdir()

    count = api.process_single_file(str(src), set(), set(), output_dir=out_dir)
    assert count >= 0

    all_files = [f for f in out_dir.iterdir() if f.is_file()]
    # Ровно один результат на входной файл, никаких скрытых или временных файлов!
    assert len(all_files) == 1
    assert all_files[0].name.endswith(expected_ext)
    assert not all_files[0].name.startswith(".")


# =========================================================================
# 2. Селективные переключатели форматов (MD-only, PDF-only, Word-only, Multi)
# =========================================================================
@pytest.mark.parametrize("input_ext", [".docx", ".xlsx", ".pdf", ".png"])
def test_markdown_only_format(tmp_path, monkeypatch, input_ext):
    api = _create_mock_backend(monkeypatch)
    api.save_markdown = True
    src = _create_sample_file(tmp_path / f"sample{input_ext}")
    out_dir = tmp_path / f"out_md_{input_ext.strip('.')}"
    out_dir.mkdir()

    api.process_single_file(str(src), set(), set(), output_dir=out_dir)
    all_files = [f for f in out_dir.iterdir() if f.is_file()]
    assert len(all_files) == 1
    assert all_files[0].name.endswith(".md")
    assert not all_files[0].name.startswith(".")


@pytest.mark.parametrize("input_ext", [".docx", ".xlsx", ".pdf", ".png"])
def test_pdf_only_format(tmp_path, monkeypatch, input_ext):
    api = _create_mock_backend(monkeypatch)
    api.save_pdf = True
    src = _create_sample_file(tmp_path / f"sample{input_ext}")
    out_dir = tmp_path / f"out_pdf_{input_ext.strip('.')}"
    out_dir.mkdir()

    api.process_single_file(str(src), set(), set(), output_dir=out_dir)
    all_files = [f for f in out_dir.iterdir() if f.is_file()]
    assert len(all_files) == 1
    assert all_files[0].name.endswith(".pdf")
    assert not all_files[0].name.startswith(".")


@pytest.mark.parametrize("input_ext", [".docx", ".xlsx", ".xlsm", ".pdf", ".png"])
def test_word_only_format_gives_docx(tmp_path, monkeypatch, input_ext):
    """Word для любого формата (включая Excel .xlsx/.xlsm) должен давать .docx!"""
    api = _create_mock_backend(monkeypatch)
    api.save_docx = True
    src = _create_sample_file(tmp_path / f"sample{input_ext}")
    out_dir = tmp_path / f"out_word_{input_ext.strip('.')}"
    out_dir.mkdir()

    api.process_single_file(str(src), set(), set(), output_dir=out_dir)
    all_files = [f for f in out_dir.iterdir() if f.is_file()]
    assert len(all_files) == 1
    assert all_files[0].name.endswith(".docx")
    assert not all_files[0].name.startswith(".")


def test_multi_formats_selected_exact_list(tmp_path, monkeypatch):
    api = _create_mock_backend(monkeypatch)
    api.save_docx = True
    api.save_pdf = True
    api.save_markdown = True

    src = _create_sample_file(tmp_path / "sample.xlsx")
    out_dir = tmp_path / "out_multi"
    out_dir.mkdir()

    api.process_single_file(str(src), set(), set(), output_dir=out_dir)
    all_files = [f for f in out_dir.iterdir() if f.is_file()]
    # Ровно выбранные форматы: Word (.docx), PDF (.pdf), Markdown (.md), без XLSX!
    assert len(all_files) == 3
    exts = {f.suffix for f in all_files}
    assert exts == {".docx", ".pdf", ".md"}
    for f in all_files:
        assert not f.name.startswith(".")


# =========================================================================
# 3. Дешифратор для одиночного файла и набора: только при включённой настройке
# =========================================================================
def test_single_file_decoder_setting(tmp_path, monkeypatch):
    api = _create_mock_backend(monkeypatch)
    src = _create_sample_file(tmp_path / "contract.docx", "Иванов Иван Иванович")

    # Без настройки save_decoder: дешифратора нет
    out1 = tmp_path / "out1"
    out1.mkdir()
    api.save_decoder = False
    api.process_single_file(str(src), set(), set(), output_dir=out1)
    files1 = [f.name for f in out1.iterdir() if f.is_file()]
    assert not any("дешифратор" in f.casefold() for f in files1)

    # С настройкой save_decoder: дешифратор есть (со встроенной привязкой v2, без sidecars)
    out2 = tmp_path / "out2"
    out2.mkdir()
    api.save_decoder = True
    api.process_single_file(str(src), set(), set(), output_dir=out2)
    files2 = [f.name for f in out2.iterdir() if f.is_file()]
    assert any("дешифратор.json" in f.casefold() for f in files2)
    assert not any(f.endswith(".provenance.json") for f in files2)
    assert not any(f.endswith(".manifest.json") for f in files2)
    assert not any(f.endswith(".audit.json") for f in files2)


def test_no_empty_decoder_when_zero_replacements(tmp_path, monkeypatch):
    api = _create_mock_backend(monkeypatch)
    api.save_decoder = True
    # Текст без персональных данных
    src = _create_sample_file(tmp_path / "empty_ner.docx", "Просто нейтральный текст без имен")
    out = tmp_path / "out_zero"
    out.mkdir()

    count = api.process_single_file(str(src), set(), set(), output_dir=out)
    assert count == 0
    files = [f.name for f in out.iterdir() if f.is_file()]
    assert not any("дешифратор" in f.casefold() for f in files)


# =========================================================================
# 4. Папочный режим: чистая папка, общий дешифратор, точный состав файлов
# =========================================================================
def test_folder_pipeline_clean_output_exact_files(tmp_path, monkeypatch):
    _create_sample_file(tmp_path / "doc1.docx", "Клиент: Иванов Иван Иванович")
    _create_sample_file(tmp_path / "doc2.xlsx", "Директор: Петров Петр Петрович")

    api = _create_mock_backend(monkeypatch)
    # По умолчанию сохраняются исходные форматы
    pipeline = FolderAnonymizationPipeline(api, emit_sidecars=False)
    result = pipeline.process(tmp_path)

    assert result.error_count == 0
    out_dir = tmp_path / DEFAULT_OUTPUT_DIR_NAME
    assert out_dir.is_dir()

    out_files = sorted([f.name for f in out_dir.iterdir() if f.is_file()])
    # В папке результатов ТОЛЬКО выбранные документы!
    assert out_files == ["doc1_cleaned.docx", "doc2_cleaned.xlsx"]

    # Общий дешифратор папки хранится РЯДОМ с ней в корне задания
    root_files = [f.name for f in tmp_path.iterdir() if f.is_file()]
    assert DEFAULT_DECODER_FILENAME in root_files
    assert not any(f.endswith(".manifest.json") for f in root_files)


# =========================================================================
# 5. P1. Финальная сверка XLSM и сохранность макросов
# =========================================================================
def test_xlsm_final_reconciliation_cross_document(tmp_path, monkeypatch):
    """Сущность (Р/с), открытая только в позднем docx, должна обезличиться в раннем xlsm."""
    account = "40702810302410001062"

    # Создаем ранний XLSM с аккаунтом
    xlsm_path = tmp_path / "01_account.xlsm"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Расчетный счет"
    ws["B1"] = account
    wb.save(str(xlsm_path))
    wb.close()

    # Добавляем в zip файл фиктивный vbaProject.bin для проверки сохранения частей VBA
    with zipfile.ZipFile(xlsm_path, "a") as zf:
        zf.writestr("xl/vbaProject.bin", b"VBA_MOCK_PAYLOAD_12345")

    # Создаем поздний DOCX, который классифицирует этот аккаунт
    _create_sample_file(tmp_path / "02_contract.docx", f"В договоре указан Р/с {account}")

    api = _create_mock_backend(monkeypatch)
    pipeline = FolderAnonymizationPipeline(api, emit_sidecars=False)
    result = pipeline.process(tmp_path)

    assert result.error_count == 0
    out_dir = tmp_path / DEFAULT_OUTPUT_DIR_NAME
    cleaned_xlsm = out_dir / "01_account_cleaned.xlsm"
    assert cleaned_xlsm.is_file()

    # Проверяем, что в очищенном XLSM исходный номер счета заменен!
    cleaned_wb = load_workbook(str(cleaned_xlsm), keep_vba=True)
    val_b1 = str(cleaned_wb.active["B1"].value)
    from xlsx_semantic import close_workbook
    close_workbook(cleaned_wb)
    assert account not in val_b1
    assert "[Р/с" in val_b1

    # Проверяем сохранность части vbaProject.bin
    with zipfile.ZipFile(cleaned_xlsm, "r") as zf:
        assert "xl/vbaProject.bin" in zf.namelist()
        assert zf.read("xl/vbaProject.bin") == b"VBA_MOCK_PAYLOAD_12345"


def test_xlsm_reconciled_in_derived_pdf_and_markdown(tmp_path, monkeypatch):
    """Если для XLSM включены PDF и MD, после сверки они не должны содержать исходное значение."""
    account = "40702810302410001062"
    xlsm_path = tmp_path / "01_account.xlsm"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = f"Реквизиты: {account}"
    wb.save(str(xlsm_path))
    wb.close()

    _create_sample_file(tmp_path / "02_contract.docx", f"Р/с {account}")

    api = _create_mock_backend(monkeypatch)
    api.save_markdown = True
    api.save_pdf = True

    pipeline = FolderAnonymizationPipeline(api, emit_sidecars=False)
    result = pipeline.process(tmp_path)
    assert result.error_count == 0

    out_dir = tmp_path / DEFAULT_OUTPUT_DIR_NAME
    cleaned_md = out_dir / "01_account_cleaned.md"
    assert cleaned_md.is_file()
    md_text = cleaned_md.read_text(encoding="utf-8")
    assert account not in md_text
    assert "[Р/с" in md_text


# =========================================================================
# 6. P1. Вложенные папки с одинаковыми именами и продолжение чистой обработки
# =========================================================================
def test_nested_folders_duplicate_filenames(tmp_path, monkeypatch):
    """Вложенные папки с одинаковыми именами файлов не должны конфликтовать."""
    dir_a = tmp_path / "dept_a"
    dir_b = tmp_path / "dept_b"
    _create_sample_file(dir_a / "act.docx", "Подписант: Иванов Иван Иванович")
    _create_sample_file(dir_b / "act.docx", "Подписант: Петров Петр Петрович")

    api = _create_mock_backend(monkeypatch)
    pipeline = FolderAnonymizationPipeline(api, emit_sidecars=False)
    result = pipeline.process(tmp_path)

    assert result.error_count == 0
    out_dir = tmp_path / DEFAULT_OUTPUT_DIR_NAME
    act_a = out_dir / "dept_a" / "act_cleaned.docx"
    act_b = out_dir / "dept_b" / "act_cleaned.docx"
    assert act_a.is_file() and act_b.is_file()

    # Оба файла восстанавливаются безопасно
    res_a = restore_document(act_a)
    res_b = restore_document(act_b)
    assert "Иванов Иван Иванович" in Document(str(res_a)).paragraphs[0].text
    assert "Петров Петр Петрович" in Document(str(res_b)).paragraphs[0].text


def test_continuation_v2_without_new_replacements(tmp_path, monkeypatch):
    """Продолжение комплекта v2, когда новый файл не содержит замен, всё равно обновляет привязку."""
    _create_sample_file(tmp_path / "doc1.docx", "Иванов Иван Иванович")
    api = _create_mock_backend(monkeypatch)
    pipeline = FolderAnonymizationPipeline(api, emit_sidecars=False)
    res1 = pipeline.process(tmp_path)
    assert res1.error_count == 0
    assert (tmp_path / DEFAULT_DECODER_FILENAME).is_file()

    # Добавляем файл без ПДн
    _create_sample_file(tmp_path / "doc2.docx", "Просто текст без сущностей")
    res2 = pipeline.process(tmp_path, continue_existing=True)
    assert res2.error_count == 0
    # Даже при 0 новых замен дешифратор и привязка должны быть сохранены!
    assert res2.decoder_path and res2.decoder_path.is_file()
    meta = json.loads(res2.decoder_path.read_text(encoding="utf-8"))["__docxdodyr_meta__"]
    doc_paths = [d["path"] for d in meta["documents"]]
    assert any("doc1_cleaned.docx" in p for p in doc_paths)
    assert any("doc2_cleaned.docx" in p for p in doc_paths)


def test_continuation_v1_compatibility(tmp_path, monkeypatch):
    """Продолжение комплекта v1 (с манифестом и sidecars)."""
    from decoder_binding import publish_binding
    doc = Document()
    doc.add_paragraph("Тест: [ФИО_1]")
    out_dir = tmp_path / DEFAULT_OUTPUT_DIR_NAME
    out_dir.mkdir()
    doc_path = out_dir / "legacy_cleaned.docx"
    doc.save(str(doc_path))

    dec_path = tmp_path / "дешифратор.json"
    dec_path.write_text(json.dumps({"[ФИО_1]": {"original": "Иванов Иван Иванович", "type": "PER"}}, ensure_ascii=False), encoding="utf-8")
    publish_binding(dec_path, [doc_path], "legacy_run_01", emit_sidecars=True)

    _create_sample_file(tmp_path / "new_doc.docx", "Иванов Иван Иванович")

    api = _create_mock_backend(monkeypatch)
    api.emit_audit_sidecars = True
    pipeline = FolderAnonymizationPipeline(api, emit_sidecars=True)
    res = pipeline.process(tmp_path, continue_existing=True)
    assert res.error_count == 0


# =========================================================================
# 7. P1. Безопасный выбор дешифратора v2: подмена, неверный путь, неоднозначность
# =========================================================================
def test_decoder_tampered_document_rejected(tmp_path, monkeypatch):
    api = _create_mock_backend(monkeypatch)
    api.save_decoder = True
    src = _create_sample_file(tmp_path / "doc.docx", "Иванов Иван Иванович")
    out_dir = tmp_path / "out"
    out_dir.mkdir()

    api.process_single_file(str(src), set(), set(), output_dir=out_dir)
    cleaned = out_dir / "doc_cleaned.docx"

    # Подделка содержимого
    with open(cleaned, "ab") as f:
        f.write(b"TAMPER")

    with pytest.raises(ValueError, match="Хеш документа не совпадает|файл был изменён"):
        restore_document(cleaned)


def test_decoder_tampered_path_rejected(tmp_path, monkeypatch):
    api = _create_mock_backend(monkeypatch)
    api.save_decoder = True
    src = _create_sample_file(tmp_path / "doc.docx", "Иванов Иван Иванович")
    out_dir = tmp_path / "out"
    out_dir.mkdir()

    api.process_single_file(str(src), set(), set(), output_dir=out_dir)
    cleaned = out_dir / "doc_cleaned.docx"
    dec = next(path for path in out_dir.glob("*.json") if "дешифратор" in path.name.casefold())

    # Подменяем относительный путь в дешифраторе
    data = json.loads(dec.read_text(encoding="utf-8"))
    data["__docxdodyr_meta__"]["documents"][0]["path"] = "some/other/path.docx"
    dec.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    with pytest.raises(ValueError, match="отсутствует привязка к документу"):
        restore_document(cleaned)


def test_equivalent_duplicate_decoders_are_accepted(tmp_path, monkeypatch):
    """Эквивалентные пакетные дешифраторы не должны блокировать восстановление."""
    api = _create_mock_backend(monkeypatch)
    api.save_decoder = True
    src = _create_sample_file(tmp_path / "doc.docx", "Иванов Иван Иванович")
    out_dir = tmp_path / "out"
    out_dir.mkdir()

    api.process_single_file(str(src), set(), set(), output_dir=out_dir)
    cleaned = out_dir / "doc_cleaned.docx"
    dec1 = next(out_dir.glob("*[Дд]ешифратор*.json"))

    # Дублируем дешифратор с другим именем в той же папке
    dec2 = out_dir / "дешифратор_дубликат.json"
    dec2.write_text(dec1.read_text(encoding="utf-8"), encoding="utf-8")

    assert find_decoder_near_document(cleaned) in {dec1.resolve(), dec2.resolve()}
    assert restore_document(cleaned).is_file()


def test_ambiguous_decoders_rejected(tmp_path, monkeypatch):
    """При разных подходящих дешифраторах нельзя молча выбирать первый."""
    api = _create_mock_backend(monkeypatch)
    api.save_decoder = True
    src = _create_sample_file(tmp_path / "doc.docx", "Иванов Иван Иванович")
    out_dir = tmp_path / "out"
    out_dir.mkdir()

    api.process_single_file(str(src), set(), set(), output_dir=out_dir)
    cleaned = out_dir / "doc_cleaned.docx"
    dec1 = next(out_dir.glob("*[Дд]ешифратор*.json"))

    data = json.loads(dec1.read_text(encoding="utf-8"))
    mapping = data.get("mapping", data)
    token = next(key for key in mapping if key.startswith("[") or key.startswith("/"))
    record = mapping[token]
    if isinstance(record, dict):
        record["original"] = "Петров Петр Петрович"
    else:
        mapping[token] = "Петров Петр Петрович"
    dec2 = out_dir / "дешифратор_конфликт.json"
    dec2.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    with pytest.raises(ValueError, match="Неоднозначность"):
        find_decoder_near_document(cleaned)
    with pytest.raises(ValueError, match="Неоднозначность"):
        restore_document(cleaned)


# =========================================================================
# 8. P2. Восстановление Markdown и отказ PDF
# =========================================================================
def test_restore_markdown_with_binding(tmp_path, monkeypatch):
    api = _create_mock_backend(monkeypatch)
    api.save_markdown = True
    api.save_decoder = True
    src = _create_sample_file(tmp_path / "contract.docx", "Стороны: ООО «Вектор» и Иванов Иван Иванович")
    out_dir = tmp_path / "out_md"
    out_dir.mkdir()

    api.process_single_file(str(src), set(), set(), output_dir=out_dir)
    cleaned_md = out_dir / "contract_cleaned.md"
    assert cleaned_md.is_file()

    restored_md = restore_document(cleaned_md)
    assert restored_md.is_file()
    restored_text = restored_md.read_text(encoding="utf-8")
    assert "Иванов Иван Иванович" in restored_text
    assert "ООО «Вектор»" in restored_text


def test_pdf_restoration_explicitly_unsupported(tmp_path, monkeypatch):
    api = _create_mock_backend(monkeypatch)
    api.save_pdf = True
    api.save_decoder = True
    src = _create_sample_file(tmp_path / "contract.docx", "Иванов Иван Иванович")
    out_dir = tmp_path / "out_pdf"
    out_dir.mkdir()

    api.process_single_file(str(src), set(), set(), output_dir=out_dir)
    cleaned_pdf = out_dir / "contract_cleaned.pdf"
    assert cleaned_pdf.is_file()

    with pytest.raises(ValueError, match="не поддерживается: формат является необратимым"):
        restore_document(cleaned_pdf)


# =========================================================================
# 9. P2. Очистка при сбоях и сохранение предыдущего комплекта
# =========================================================================
def test_no_temp_files_on_pdf_conversion_failure(tmp_path, monkeypatch):
    api = _create_mock_backend(monkeypatch)
    api.save_pdf = True
    src = _create_sample_file(tmp_path / "bad.docx", "Иванов Иван Иванович")
    out_dir = tmp_path / "out_err"
    out_dir.mkdir()

    def _failing_converter(*_args, **_kwargs):
        raise RuntimeError("Simulated conversion failure")

    monkeypatch.setattr("backend_api.convert_to_pdf", _failing_converter)
    monkeypatch.setattr("backend_api.convert_docx_text_to_pdf", _failing_converter)

    with pytest.raises(IOError):
        api.process_single_file(str(src), set(), set(), output_dir=out_dir)

    # В out_dir не должно остаться ни одного файла
    remaining = list(out_dir.iterdir())
    assert len(remaining) == 0


def test_atomic_folder_publication_preserves_archive(tmp_path, monkeypatch):
    """При падении публикации предыдущий комплект не теряется."""
    api = _create_mock_backend(monkeypatch)
    _create_sample_file(tmp_path / "doc.docx", "Иванов Иван Иванович")

    pipeline = FolderAnonymizationPipeline(api, emit_sidecars=False)
    res1 = pipeline.process(tmp_path)
    assert res1.error_count == 0

    pub_dir = tmp_path / DEFAULT_OUTPUT_DIR_NAME
    assert pub_dir.is_dir()
    original_marker = pub_dir / "marker.txt"
    original_marker.write_text("INITIAL_RUN", encoding="utf-8")
    initial_decoders = {
        path.name for path in tmp_path.glob("*.json")
        if path.name.casefold().startswith("дешифратор")
    }

    # Имитируем сбой при второй публикации (только первая попытка заменяет pub_dir неудачно)
    orig_replace = os.replace
    failed = False
    def _failing_replace(src, dst):
        nonlocal failed
        if str(dst) == str(pub_dir) and not failed:
            failed = True
            raise OSError("Simulated atomic replace failure")
        return orig_replace(src, dst)

    monkeypatch.setattr(os, "replace", _failing_replace)

    with pytest.raises(OSError):
        pipeline.process(tmp_path)

    # Исходная папка осталась на месте и содержит первоначальный маркер
    assert pub_dir.is_dir()
    assert original_marker.is_file()
    assert original_marker.read_text(encoding="utf-8") == "INITIAL_RUN"
    assert {
        path.name for path in tmp_path.glob("*.json")
        if path.name.casefold().startswith("дешифратор")
    } == initial_decoders


def test_cancel_before_folder_publication_leaves_no_orphan_decoder(tmp_path, monkeypatch):
    _create_sample_file(tmp_path / "doc.docx", "Иванов Иван Иванович")
    api = _create_mock_backend(monkeypatch)
    checks = 0

    def cancel_after_reconciliation():
        nonlocal checks
        checks += 1
        return checks >= 3

    api.is_cancelled = cancel_after_reconciliation
    result = FolderAnonymizationPipeline(api, emit_sidecars=False).process(tmp_path)

    assert result.error_count > 0
    assert result.decoder_path is None
    assert not [
        path for path in tmp_path.glob("*.json")
        if path.name.casefold().startswith("дешифратор")
    ]
    assert not (tmp_path / DEFAULT_OUTPUT_DIR_NAME).exists()


# =========================================================================
# 10. OCR Smoke-тест для интерактивной сессии macOS
# =========================================================================
@pytest.mark.skipif(sys.platform != "darwin", reason="Apple Vision доступен только на macOS")
def test_apple_vision_real_smoke_interactive():
    """Настоящий smoke-тест Apple Vision OCR."""
    from ocr_backend import _AppleVisionBackend
    import io
    try:
        backend = _AppleVisionBackend(language="rus+eng")
        img = Image.new("RGB", (400, 100), color=(255, 255, 255))
        draw = ImageDraw.Draw(img)
        draw.text((10, 10), "TEST OCR", fill=(0, 0, 0))
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        results = backend.recognize(buf.getvalue())
        assert any("TEST" in r for r in results)
    except Exception as exc:
        pytest.skip(f"Apple Vision требует интерактивную оконную сессию macOS: {exc}")
