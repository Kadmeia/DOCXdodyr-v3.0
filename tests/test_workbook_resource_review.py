"""Workbook resources must be released deterministically, including VBA ZIPs."""

import zipfile

import openpyxl
import pytest

from batch_reconcile import reconcile_xlsx
from document_restorer import DocumentRestorer
from tests.test_output_format_and_clean_restoration_matrix import _create_mock_backend


@pytest.fixture
def tracked_workbooks(tmp_path, monkeypatch):
    source = tmp_path / "synthetic.xlsm"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "Synthetic original value"
    workbook.save(source)
    workbook.close()
    with zipfile.ZipFile(source, "a") as archive:
        archive.writestr("xl/vbaProject.bin", b"SYNTHETIC VBA")
    loaded = []
    original_load = openpyxl.load_workbook
    def track(*args, **kwargs):
        result = original_load(*args, **kwargs)
        loaded.append(result)
        return result
    monkeypatch.setattr(openpyxl, "load_workbook", track)
    yield source, loaded
    # Cleanup remains safe even when an assertion fails.
    from xlsx_semantic import close_workbook
    for item in loaded:
        close_workbook(item)


def assert_closed(loaded):
    assert loaded
    assert all(item.vba_archive is not None for item in loaded)
    assert all(item.vba_archive.fp is None for item in loaded)


@pytest.mark.parametrize("fail", [False, True])
def test_reconciliation_closes_vba_on_success_and_failure(tracked_workbooks, monkeypatch, fail):
    source, loaded = tracked_workbooks
    def reconcile(value, mapping):
        if fail:
            raise RuntimeError("synthetic processing failure")
        return "[VALUE_1]", 1
    monkeypatch.setattr("batch_reconcile.reconcile_text", reconcile)
    if fail:
        with pytest.raises(RuntimeError, match="synthetic processing failure"):
            reconcile_xlsx(source, {})
    else:
        assert reconcile_xlsx(source, {}) == 1
        with zipfile.ZipFile(source) as archive:
            assert archive.read("xl/vbaProject.bin") == b"SYNTHETIC VBA"
    assert_closed(loaded)


@pytest.mark.parametrize("fail", [False, True])
def test_restoration_closes_vba_on_success_and_save_failure(tracked_workbooks, tmp_path, monkeypatch, fail):
    source, loaded = tracked_workbooks
    restorer = DocumentRestorer({})
    monkeypatch.setattr(restorer, "_verify_document", lambda path: None)
    if fail:
        def fail_save(*args, **kwargs):
            raise OSError("synthetic save failure")
        monkeypatch.setattr(openpyxl.Workbook, "save", fail_save)
    ok, message = restorer.restore_excel(source, tmp_path / "restored.xlsm")
    assert ok is (not fail)
    if fail:
        assert "synthetic save failure" in message
    assert_closed(loaded)


@pytest.mark.parametrize("fail", [False, True])
def test_backend_closes_vba_on_success_and_processing_failure(tracked_workbooks, tmp_path, monkeypatch, fail):
    source, loaded = tracked_workbooks
    api = _create_mock_backend(monkeypatch)
    if fail:
        def fail_profile(*args, **kwargs):
            raise RuntimeError("synthetic processing failure")
        monkeypatch.setattr("xlsx_semantic.sheet_profile", fail_profile)
        with pytest.raises(RuntimeError, match="synthetic processing failure"):
            api.process_single_file(str(source), set(), set(), output_dir=tmp_path / "out")
    else:
        api.process_single_file(str(source), set(), set(), output_dir=tmp_path / "out")
    assert_closed(loaded)
