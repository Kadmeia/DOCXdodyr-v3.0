# -*- coding: utf-8 -*-

import pytest
from docx import Document

from backend_api import BackendApi
from table_entity_context import build_table_contexts


@pytest.fixture(scope="module")
def api():
    return BackendApi()


def _clean(api, text):
    return api.anonymize_text_pullenti(
        text, set(), set(), mapping_dict={}, entity_seen={}
    )[0]


def test_email_wins_over_nested_bare_website(api):
    source = "электронная почта: ivanov@example.invalid"
    assert _clean(api, source) == "электронная почта: [Email_1]"


def test_unlabelled_compound_inn_kpp_has_two_semantic_identifiers(api):
    assert _clean(api, "7704747169 / 770401001") == "[ИНН_1] / [КПП_1]"


def test_product_model_and_technical_measurements_are_not_organizations(api):
    source = "MSPOS-SE-Ф (ГК МультиСофт)"
    assert _clean(api, source) == "MSPOS-SE-Ф (ГК [Наименование_1])"
    specs = "Память оперативная –1 Гб.; Память постоянная –8 Гб."
    assert _clean(api, specs) == specs


def test_person_record_number_after_birth_date_sex_and_country_is_sensitive(api):
    source = (
        "Смирнову Марию Ивановну 25.02.1950 жен РФ 4501123456 "
        "(ИНВАЛИД) и сопровождающее лицо Васильеву Ольгу Петровну "
        "21.10.1975 жен РФ 4501654321"
    )
    cleaned = _clean(api, source)
    assert "4501123456" not in cleaned
    assert "4501654321" not in cleaned
    assert cleaned.count("[СерияПаспорта_") == 2
    assert cleaned.count("[НомерПаспорта_") == 2


def test_resume_month_year_period_and_graduation_year_are_dates(api):
    assert _clean(api, "Сентябрь 2024 — настоящее время\n5 месяцев") == (
        "[Дата_1] — настоящее время\n5 месяцев"
    )
    rows = [
        ["Образование", "Образование", "Образование"],
        ["2028", "2028", "Государственная академия"],
    ]
    context = next(
        item for item in build_table_contexts(rows)
        if (item.row, item.column) == (1, 0)
    )
    assert context.label == "Дата образования"
    assert api.anonymize_text_pullenti(
        "2028", set(), set(), mapping_dict={}, entity_seen={},
        semantic_context=context,
    )[0] == "[Дата_1]"


def test_resume_employer_and_educational_legal_entity_are_redacted(api):
    assert _clean(api, "АБСОЛЮТ, Группа") == "[Наименование_1]"
    source = (
        "Российская академия народного хозяйства и государственной службы "
        "при Президенте Российской Федерации, Москва"
    )
    assert _clean(api, source) == "[Наименование_1], Москва"


def test_education_level_category_is_preserved(api):
    assert _clean(api, "Образование: Неоконченное высшее") == (
        "Образование: Неоконченное высшее"
    )


def test_luna_new5_semantic_boundaries(api):
    assert _clean(api, "book@example.invalid") == "[Email_1]"
    assert _clean(api, "ФИО Туристов, проживавших") == "ФИО Туристов, проживавших"
    assert _clean(api, 'ОП "Гостиничный комплекс"') == 'ОП "Гостиничный комплекс"'
    assert _clean(api, 'ОП "Гостиничный комплекс Кинопарк"') == (
        'ОП "Гостиничный комплекс [Наименование_1]"'
    )
    company = '30.03.2023 г. между ООО ПСК «Монолит» (далее – «ответчик»)'
    assert _clean(api, company) == '[Дата_1] г. между ООО ПСК «[Наименование_1]» (далее – «ответчик»)'


def test_public_law_dates_and_private_order_requisites(api):
    assert _clean(api, "ч.6 ст. 13 ФЗ от 27.07.2006 № 149-ФЗ") == (
        "ч.6 ст. 13 ФЗ от 27.07.2006 № 149-ФЗ"
    )
    assert _clean(api, "зарегистрирован Минюстом России 14.05.2021 № 63416") == (
        "зарегистрирован Минюстом России 14.05.2021 № 63416"
    )
    assert _clean(api, "на основании приказа № 6/П от 01.08.2024") == (
        "на основании приказа № [НомерДокумента_1] от [Дата_1]"
    )
    assert _clean(api, "Дата") == "Дата"


def test_city_only_after_a_completed_full_address_sentence_is_preserved(api):
    source = (
        "г. Приморск, Северной области, ул. Примерная 10, подъезд 1, "
        "этаж 3. Есть грузовой лифт - г. Москва, Ярославский вокзал, "
        "Поезд отправление 18.10.2024 01:00."
    )
    cleaned = _clean(api, source)
    assert "г. Москва, Ярославский вокзал" in cleaned
    assert "ул. Примерная" not in cleaned


def test_all_header_and_footer_variants_are_processed(api):
    doc = Document()
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.header.paragraphs[0].text = (
        "Документ предоставлен КонсультантПлюс\nДата сохранения: 04.06.2025"
    )
    section.first_page_header.paragraphs[0].text = (
        "Документ предоставлен КонсультантПлюсДата сохранения: 04.06.2025"
    )
    section.footer.paragraphs[0].text = (
        "Иванова Мария • Резюме обновлено 10 января 2025 в 14:05"
    )
    section.first_page_footer.paragraphs[0].text = (
        "Резюме обновлено 10 января 2025 в 14:05"
    )

    api.clean_document(doc, set(), set(), mapping_dict={}, entity_seen={})

    assert "04.06.2025" in section.header.paragraphs[0].text
    assert "04.06.2025" in section.first_page_header.paragraphs[0].text
    assert "10 января 2025" not in section.footer.paragraphs[0].text
    assert "10 января 2025" not in section.first_page_footer.paragraphs[0].text
