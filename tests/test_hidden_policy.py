"""Contracts for aggregate hidden-data inspection and explicit policy actions."""

import zipfile
from pathlib import Path

import pytest
from docx import Document

from hidden_data import HiddenDataError, apply_hidden_data_policy


def _package_with_hidden_parts(path: Path, *, signed: bool = False) -> None:
    document = Document()
    document.add_paragraph("видимый текст")
    document.save(path)
    temp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp, "w") as target:
        for info in source.infolist():
            target.writestr(info, source.read(info.filename))
        target.writestr(
            "word/comments.xml",
            b"<w:comments xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:comment w:id='0'><w:p><w:r><w:t>secret comment</w:t></w:r></w:p></w:comment></w:comments>",
        )
        target.writestr(
            "word/footnotes.xml",
            b"<w:footnotes xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:footnote w:id='1'><w:p><w:r><w:t>secret note</w:t></w:r></w:p></w:footnote></w:footnotes>",
        )
        if signed:
            target.writestr("_xmlsignatures/sig1.xml", b"<sig />")
    temp.replace(path)


def test_policy_anonymizes_textual_hidden_parts_without_exposing_values(tmp_path: Path):
    source = tmp_path / "source.docx"
    output = tmp_path / "clean.docx"
    _package_with_hidden_parts(source)

    report = apply_hidden_data_policy(
        source,
        output,
        {"comments": "anonymize", "footnotes": "anonymize"},
        text_transform=lambda text, kind: text.replace("secret", "[СКРЫТО]") ,
    )
    assert report.detected["comments"] == 1
    assert "anonymized_comments" in report.removed
    assert "anonymized_footnotes" in report.removed
    with zipfile.ZipFile(output) as out_zip:
        comments_xml = out_zip.read("word/comments.xml")
    assert b"secret" not in comments_xml
    assert b"[\xd0\xa1\xd0\x9a\xd0\xa0\xd0\xab\xd0\xa2\xd0\x9e]" in comments_xml


def test_signed_office_policy_fails_closed(tmp_path: Path):
    source = tmp_path / "signed.docx"
    _package_with_hidden_parts(source, signed=True)
    with pytest.raises(HiddenDataError):
        apply_hidden_data_policy(source, tmp_path / "out.docx", {"comments": "remove"})


def test_backend_policy_requires_explicit_confirmation_for_detected_removal():
    from backend_api import BackendApi

    api = BackendApi.__new__(BackendApi)
    api.hidden_data_policy = api._default_hidden_data_policy()
    api.settings = {}
    api._save_settings = lambda: None
    pending = api.save_hidden_data_policy({"footnotes": "remove"})
    assert pending["requires_confirmation"] is True
    accepted = api.save_hidden_data_policy({"footnotes": "remove"}, confirm_destructive=True)
    assert accepted["success"] is True


def test_hidden_data_ui_exposes_aggregate_actions_and_confirm_gate():
    root = Path(__file__).parents[1]
    html = (root / "web" / "index.html").read_text(encoding="utf-8")
    script = (root / "web" / "script.js").read_text(encoding="utf-8")
    for marker in ("hidden-data-overlay", "hidden-data-risks", "hidden-data-confirm", "hidden-data-apply"):
        assert marker in html
    assert "inspect_hidden_file" in script
    assert "apply_hidden_data_policy" in script
    assert "hiddenDataHasDestructiveChoice" in script
    assert "textContent" in script and "hidden-data-risks" in script


def test_excel_hidden_sheets_anonymized(tmp_path: Path):
    """BUG-08: Скрытые листы XLSX корректно обезличиваются."""
    from openpyxl import Workbook, load_workbook
    from backend_api import BackendApi

    source = tmp_path / "hidden.xlsx"
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Visible"
    ws1["A1"] = "Иванов Иван"

    ws2 = wb.create_sheet("Hidden")
    ws2.sheet_state = "hidden"
    ws2["A1"] = "Петров Петр"

    ws3 = wb.create_sheet("VeryHidden")
    ws3.sheet_state = "veryHidden"
    ws3["A1"] = "Сидоров Сидор"

    wb.save(str(source))
    wb.close()

    api = BackendApi()
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    api.save_original = False
    api.save_pdf = False
    api.save_markdown = False
    api.save_docx = False

    count = api.process_single_file(str(source), set(), [], output_dir=out_dir)
    assert count >= 3

    cleaned = list(out_dir.glob("*_cleaned.xlsx"))
    assert len(cleaned) == 1

    wb_clean = load_workbook(str(cleaned[0]))
    c_ws1 = wb_clean["Visible"]
    c_ws2 = wb_clean["Hidden"]
    c_ws3 = wb_clean["VeryHidden"]

    assert c_ws1.sheet_state == "visible"
    assert c_ws2.sheet_state == "hidden"
    assert c_ws3.sheet_state == "veryHidden"

    assert "Иванов Иван" not in str(c_ws1["A1"].value)
    assert "[ФИО" in str(c_ws1["A1"].value)

    assert "Петров Петр" not in str(c_ws2["A1"].value)
    assert "[ФИО" in str(c_ws2["A1"].value)

    assert "Сидоров Сидор" not in str(c_ws3["A1"].value)
    assert "[ФИО" in str(c_ws3["A1"].value)
    wb_clean.close()
