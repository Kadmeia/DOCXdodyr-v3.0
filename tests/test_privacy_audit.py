from __future__ import annotations

import json
import os
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter

from privacy_audit import (
    ReviewFinding,
    ReviewQueue,
    build_audit_certificate,
    clean_file_metadata,
    write_audit_certificate,
)


def test_review_queue_whitelists_fields_and_tracks_decisions(tmp_path):
    queue = ReviewQueue([{
        "document_ref": "contract_cleaned.docx",
        "location": "Док.Пар.3",
        "entity_type": "person",
        "placeholder": "[ФИО_1]",
        "confidence": 1.5,
        "source": "pullenti",
        "redacted_context": "Подписант: [ФИО_1].",
        "original": "Иванов Иван Иванович",  # must never enter the record
    }])
    item = queue.pending()[0]
    assert item.confidence == 1.0
    assert "original" not in item.to_dict()
    queue.decide(item.finding_id, "accepted", "Проверено")
    assert queue.summary() == {"pending": 0, "accepted": 1, "rejected": 0, "skipped": 0, "total": 1}

    saved = queue.save(tmp_path / "review.json")
    payload = json.loads(saved.read_text(encoding="utf-8"))
    assert "Иванов" not in saved.read_text(encoding="utf-8")
    assert payload["items"][0]["status"] == "accepted"


def test_docx_metadata_is_cleaned_without_touching_body(tmp_path):
    source = tmp_path / "source.docx"
    cleaned = tmp_path / "cleaned.docx"
    doc = Document()
    doc.add_paragraph("Текст: [ФИО_1]")
    doc.core_properties.author = "Иванов Иван"
    doc.core_properties.title = "Секретный договор"
    doc.core_properties.subject = "Персональные данные"
    doc.save(source)

    report = clean_file_metadata(source, cleaned)
    assert report.changed is True
    assert Document(cleaned).paragraphs[0].text == "Текст: [ФИО_1]"
    with zipfile.ZipFile(cleaned) as package:
        core = package.read("docProps/core.xml").decode("utf-8")
        assert "Иванов" not in core
        assert "Секретный договор" not in core
        relationships = package.read("_rels/.rels")
        content_types = package.read("[Content_Types].xml")
        assert b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"' in relationships
        assert b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"' in content_types


def test_xlsx_metadata_is_cleaned_and_cells_remain(tmp_path):
    source = tmp_path / "source.xlsx"
    cleaned = tmp_path / "cleaned.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "[ФИО_1]"
    workbook.properties.creator = "Иванов Иван"
    workbook.properties.title = "Договор с персональными данными"
    workbook.save(source)

    report = clean_file_metadata(source, cleaned)
    assert report.changed is True
    assert load_workbook(cleaned, read_only=True).active["A1"].value == "[ФИО_1]"
    with zipfile.ZipFile(cleaned) as package:
        core = package.read("docProps/core.xml").decode("utf-8")
        assert "Иванов" not in core


def test_pdf_info_metadata_is_cleaned(tmp_path):
    source = tmp_path / "source.pdf"
    cleaned = tmp_path / "cleaned.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_metadata({"/Author": "Иванов Иван", "/Title": "Личное дело"})
    with source.open("wb") as stream:
        writer.write(stream)

    report = clean_file_metadata(source, cleaned)
    assert report.changed is True
    metadata = PdfReader(str(cleaned)).metadata
    assert not metadata or not metadata.author
    assert not metadata or not metadata.title


def test_certificate_contains_hashes_and_aggregates_but_no_pii(tmp_path):
    source = tmp_path / "source.docx"
    result = tmp_path / "result.docx"
    source.write_bytes(b"source")
    result.write_bytes(b"result")
    queue = ReviewQueue([ReviewFinding("contract.docx", "p.1", "PER", "[ФИО_1]", 0.95)])
    certificate = build_audit_certificate(
        source,
        result,
        counts={
            "files": 1,
            "replacements": 2,
            "entities_by_type": {"PER": 1},
            "original": "Иванов Иван Иванович",
        },
        review_queue=queue,
    )
    out = write_audit_certificate(certificate, tmp_path / "audit.json")
    text = out.read_text(encoding="utf-8")
    assert "sha256" in text
    assert "Иванов" not in text
    assert certificate["counts"]["entities_by_type"] == {"PER": 1}
    assert certificate["counts"]["review"]["pending"] == 1


def test_audit_certificate_is_private_on_posix(tmp_path):
    if os.name != "posix":
        return
    source = tmp_path / "source"
    result = tmp_path / "result"
    source.write_bytes(b"source")
    result.write_bytes(b"result")
    out = write_audit_certificate(build_audit_certificate(source, result), tmp_path / "audit.json")
    assert out.stat().st_mode & 0o777 == 0o600
