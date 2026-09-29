# -*- coding: utf-8 -*-
"""Проверки скрытых частей Office/PDF и необратимого PDF-режима."""

from __future__ import annotations

import io
import os
import zipfile
from pathlib import Path

import fitz
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from types import SimpleNamespace

from hidden_data import (
    HiddenDataError,
    _transform_word_xml,
    apply_hidden_data_policy,
    flatten_pdf_to_images,
    inspect_hidden_data,
    sanitize_hidden_data,
)


def test_word_transform_keeps_mc_ignorable_prefixes_well_formed():
    source = (
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        b'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        b'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
        b'xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml" '
        b'mc:Ignorable="w14 w15"><w:body><w:p w14:paraId="12345678"/></w:body></w:document>'
    )
    transformed, _ = _transform_word_xml(
        source, comments=False, notes=False, revisions=False, hidden_text=False
    )
    assert b'mc:Ignorable="w14"' in transformed
    assert b'xmlns:w14=' in transformed
    assert b'w15' not in transformed


def _add_office_parts(source: Path) -> None:
    temp = source.with_suffix(".tmp.docx")
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(temp, "w") as zout:
        for info in zin.infolist():
            if info.filename != "word/document.xml":
                zout.writestr(info, zin.read(info.filename))
        zout.writestr("word/comments.xml", b"<w:comments xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:comment w:id='0'/></w:comments>")
        zout.writestr("word/footnotes.xml", b"<w:footnotes xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'/>")
        zout.writestr("word/embeddings/oleObject1.bin", b"private attachment")
        # Include references in the document body to check that the scrubber
        # removes links while retaining a valid OOXML package.
        body = b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body><w:p><w:r><w:commentReference w:id='0'/><w:footnoteReference w:id='1'/><w:t>visible</w:t></w:r></w:p></w:body></w:document>"
        zout.writestr("word/document.xml", body)
    temp.replace(source)


def test_office_hidden_data_report_and_strict_scrub(tmp_path: Path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("visible")
    document.save(source)
    _add_office_parts(source)

    report = inspect_hidden_data(source)
    assert report.detected["comments"] == 1
    assert report.detected["footnotes"] == 1
    assert report.detected["embedded_objects"] == 1

    output = tmp_path / "clean.docx"
    scrub = sanitize_hidden_data(source, output, remove_footnotes=True)
    assert output.exists()
    assert "comments" in scrub.removed
    assert "footnotes" in scrub.removed
    with zipfile.ZipFile(output) as package:
        names = set(package.namelist())
        assert "word/comments.xml" not in names
        assert "word/footnotes.xml" not in names
        assert not any(name.startswith("word/embeddings/") for name in names)
        assert b"commentReference" not in package.read("word/document.xml")
    # It remains a readable Word package after removing hidden parts.
    assert Document(output).paragraphs[0].text == "visible"


def test_sanitize_reuses_matching_preflight_report(monkeypatch, tmp_path: Path):
    import hidden_data

    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("visible")
    document.save(source)
    report = hidden_data.inspect_hidden_data(source)
    monkeypatch.setattr(
        hidden_data,
        "inspect_hidden_data",
        lambda _path: (_ for _ in ()).throw(AssertionError("duplicate scan")),
    )

    output = tmp_path / "clean.docx"
    result = hidden_data.sanitize_hidden_data(
        source, output, preflight_report=report
    )

    assert result is report
    assert Document(output).paragraphs[0].text == "visible"


def test_sanitized_office_output_is_private_on_posix(tmp_path: Path):
    if os.name != "posix":
        return
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("private")
    document.save(source)
    source.chmod(0o644)
    output = tmp_path / "clean.docx"
    sanitize_hidden_data(source, output)
    assert output.stat().st_mode & 0o777 == 0o600


def test_sanitize_rejects_preflight_report_for_another_file(tmp_path: Path):
    first = tmp_path / "first.docx"
    second = tmp_path / "second.docx"
    Document().save(first)
    Document().save(second)
    report = inspect_hidden_data(first)

    with __import__("pytest").raises(HiddenDataError, match="другому файлу"):
        sanitize_hidden_data(
            second, tmp_path / "clean.docx", preflight_report=report
        )


def test_hidden_text_anonymization_never_reprocesses_visible_body(tmp_path: Path):
    source = tmp_path / "source.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("visible 24.02.2021")
    hidden = paragraph.add_run(" hidden Иванов")
    hidden.font.hidden = True
    document.save(source)

    output = tmp_path / "clean.docx"
    apply_hidden_data_policy(
        source,
        output,
        {"hidden_text": "anonymize"},
        text_transform=lambda text, _kind: text.replace("Иванов", "[ФИО]"),
    )

    cleaned = Document(output).paragraphs[0]
    assert cleaned.runs[0].text == "visible 24.02.2021"
    assert cleaned.runs[1].text == " hidden [ФИО]"


def test_external_hyperlink_target_is_removed_but_visible_text_survives(tmp_path: Path):
    source = tmp_path / "linked.docx"
    document = Document()
    paragraph = document.add_paragraph()
    relation_id = paragraph.part.relate_to(
        "https://private.example/path", RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = "[Сайт_1]"
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    document.save(source)

    output = tmp_path / "clean.docx"
    apply_hidden_data_policy(source, output, {"external_links": "remove"})
    assert Document(output).paragraphs[0].text == "[Сайт_1]"
    with zipfile.ZipFile(output) as package:
        rels = package.read("word/_rels/document.xml.rels")
        assert b"private.example" not in rels
        assert b'TargetMode="External"' not in rels


def test_tracked_changes_are_accepted_without_losing_inserted_text(tmp_path: Path):
    source = tmp_path / "revisions.docx"
    document = Document()
    document.add_paragraph("prefix")
    document.save(source)
    temp = tmp_path / "revisions-with-xml.docx"
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(temp, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "word/document.xml":
                data = data.replace(
                    b"<w:r><w:t>prefix</w:t></w:r>",
                    b"<w:r><w:t>prefix</w:t></w:r><w:del><w:r><w:delText> old</w:delText></w:r></w:del>"
                    b"<w:ins><w:r><w:t> new</w:t></w:r></w:ins>",
                )
            zout.writestr(info, data)
    temp.replace(source)

    output = tmp_path / "accepted.docx"
    apply_hidden_data_policy(source, output, {"tracked_changes": "remove"})
    assert Document(output).paragraphs[0].text == "prefix new"


def test_pdf_flatten_removes_text_annotations_and_metadata(tmp_path: Path):
    source = tmp_path / "source.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Иванов Иван Иванович")
    page.add_text_annot((72, 100), "private comment")
    document.set_metadata({"title": "private title", "author": "private author"})
    document.save(source)
    document.close()

    output = tmp_path / "flattened.pdf"
    report = flatten_pdf_to_images(source, output, dpi=120)
    assert report.page_count == 1
    check = fitz.open(output)
    assert check.page_count == 1
    assert check[0].get_text().strip() == ""
    assert check[0].get_images()
    assert list(check[0].annots() or ()) == []
    assert all(value not in {"private title", "private author"} for value in check.metadata.values())
    check.close()
    assert "Иванов" not in output.read_bytes().decode("latin-1", errors="ignore")


def test_pdf_scrub_report_is_strict_for_encrypted_pdf(tmp_path: Path):
    source = tmp_path / "encrypted.pdf"
    document = fitz.open()
    document.new_page().insert_text((72, 72), "secret")
    document.save(source, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user")
    document.close()
    with __import__("pytest").raises(HiddenDataError):
        sanitize_hidden_data(source, tmp_path / "clean.pdf")


def test_backend_irreversible_pdf_setting_flattens_generated_pdf(monkeypatch, tmp_path: Path):
    import backend_api

    source = tmp_path / "contract.docx"
    document = Document()
    document.add_paragraph("Обезличенный договор [ФИО_1]")
    document.save(source)

    def fake_convert(_source, target):
        pdf = fitz.open()
        pdf.new_page().insert_text((72, 72), "extractable placeholder text")
        pdf.save(str(target))
        pdf.close()
        return True

    class EmptyProcessor:
        def process(self, *_args):
            return SimpleNamespace(entities=[])

    monkeypatch.setattr(backend_api, "PDF_CONVERSION_AVAILABLE", True)
    monkeypatch.setattr(backend_api, "convert_to_pdf", fake_convert)
    api = backend_api.BackendApi.__new__(backend_api.BackendApi)
    api._pullenti_processor = EmptyProcessor()
    api.current_placeholders = {}
    api.qwen_settings = SimpleNamespace(enabled=False)
    api._qwen_postprocessor = SimpleNamespace(last_status="disabled")
    api.save_original = False
    api.save_pdf = True
    api.save_decoder = False
    api.irreversible_pdf = True
    api.save_audit_files = True
    api.ocr_lang = "rus"
    api._window = None
    api.process_single_file(str(source), set(), set())

    output = tmp_path / "contract_cleaned.pdf"
    assert output.exists()
    check = fitz.open(output)
    assert check[0].get_text().strip() == ""
    check.close()
    audit = output.with_name(output.name + ".audit.json")
    assert '"irreversible_pdf": true' in audit.read_text(encoding="utf-8")


def test_chart_embeddings_preserved_when_sanitizing_hidden_data(tmp_path: Path):
    source = tmp_path / "chart_doc.docx"
    doc = Document()
    doc.add_paragraph("Document with chart")
    doc.save(source)

    temp = tmp_path / "chart_doc_packaged.docx"
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(temp, "w") as zout:
        for info in zin.infolist():
            zout.writestr(info, zin.read(info.filename))
        # Add chart with its rels pointing to chart_data.xlsx
        chart_xml = b"""<?xml version="1.0" encoding="utf-8"?>
<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <c:externalData r:id="rId1"/>
</c:chartSpace>"""
        chart_rels = b"""<?xml version="1.0" encoding="utf-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/package" Target="../embeddings/chart_data.xlsx"/>
</Relationships>"""
        zout.writestr("word/charts/chart1.xml", chart_xml)
        zout.writestr("word/charts/_rels/chart1.xml.rels", chart_rels)
        zout.writestr("word/embeddings/chart_data.xlsx", b"excel chart data")
        zout.writestr("word/embeddings/attachment.bin", b"private standalone file")

    temp.replace(source)

    report = inspect_hidden_data(source)
    # Only the standalone attachment is detected as embedded_objects, not chart data
    assert report.detected.get("embedded_objects") == 1

    clean = tmp_path / "chart_doc_clean.docx"
    sanitize_hidden_data(source, clean, remove_attachments=True)

    with zipfile.ZipFile(clean) as zclean:
        clean_names = set(zclean.namelist())
        assert "word/embeddings/chart_data.xlsx" in clean_names
        assert "word/embeddings/attachment.bin" not in clean_names
        assert "word/charts/chart1.xml" in clean_names
        assert "word/charts/_rels/chart1.xml.rels" in clean_names

