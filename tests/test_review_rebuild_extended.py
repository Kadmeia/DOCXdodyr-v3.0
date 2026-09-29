from __future__ import annotations

import sys
from types import SimpleNamespace
import zipfile

import pytest
from docx import Document
from docx.shared import Inches
from openpyxl import load_workbook

from review_rebuild import rebuild_docx, rebuild_pdf
from review_context import SecureContextError, SecureContextVault
import review_context


def test_docx_header_footer_and_table_coordinates(tmp_path):
    path = tmp_path / "headers.docx"
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "Шапка [ФИО_1]"
    section.footer.add_table(rows=1, cols=1, width=Inches(5)).cell(0, 0).text = "Подвал [АДРЕС_1]"
    doc.save(path)
    assert rebuild_docx(path, {"scope": "header_footer", "block": "header", "section": 0,
                               "kind": "paragraph", "index": 0}, "[ФИО_1]", "Иванов")
    assert rebuild_docx(path, {"scope": "header_footer", "block": "footer", "section": 0,
                               "kind": "table_cell", "table": 0, "row": 0, "column": 0},
                        "[АДРЕС_1]", "Москва")
    reopened = Document(path)
    assert reopened.sections[0].header.paragraphs[0].text == "Шапка Иванов"
    assert reopened.sections[0].footer.tables[0].cell(0, 0).text == "Подвал Москва"


def test_docx_footnote_xml_coordinate(tmp_path):
    path = tmp_path / "notes.docx"
    doc = Document()
    doc.add_paragraph("Текст [ФИО_1]")
    doc.save(path)
    # Inject a minimal notes part for a package-level regression test.
    with zipfile.ZipFile(path, "a") as package:
        package.writestr("word/footnotes.xml", '''<?xml version="1.0"?><w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="7"><w:p><w:r><w:t>[ФИО_1]</w:t></w:r></w:p></w:footnote></w:footnotes>'''.encode("utf-8"))
    assert rebuild_docx(path, {"scope": "footnotes", "id": 7}, "[ФИО_1]", "Иванов")
    with zipfile.ZipFile(path) as package:
        assert "Иванов" in package.read("word/footnotes.xml").decode("utf-8")


def test_secure_context_roundtrip_and_corruption(tmp_path):
    pytest.importorskip("cryptography")
    key = b"k" * 32
    vault = SecureContextVault(tmp_path / ".ctx", key_provider=lambda: key)
    vault.put("finding", {"original": "Иванов Иван", "coordinate": {"page": 1}})
    assert vault.load()["finding"]["original"] == "Иванов Иван"
    assert "Иванов".encode("utf-8") not in (tmp_path / ".ctx").read_bytes()
    (tmp_path / ".ctx").write_bytes(b"corrupt")
    with pytest.raises(SecureContextError):
        vault.load()


def test_secure_context_missing_key_fails_closed(tmp_path):
    pytest.importorskip("cryptography")
    vault = SecureContextVault(tmp_path / ".ctx", key_provider=lambda: b"bad")
    with pytest.raises(SecureContextError):
        vault.save({"x": {"original": "secret"}})


def test_existing_secure_context_never_rotates_missing_key(monkeypatch, tmp_path):
    """Потеря доступа к Keychain не должна уничтожать ключ существующего vault."""
    pytest.importorskip("cryptography")
    path = tmp_path / ".ctx"
    original = SecureContextVault(path, key_provider=lambda: b"k" * 32)
    original.save({"finding": {"original": "Иванов"}})
    before = path.read_bytes()
    calls = []

    def unavailable(_service, _account, *, create_if_missing=True):
        calls.append(create_if_missing)
        raise SecureContextError("key unavailable")

    monkeypatch.setattr(review_context, "_keychain_key", unavailable)
    reopened = SecureContextVault(path)

    with pytest.raises(SecureContextError):
        reopened.load()
    with pytest.raises(SecureContextError):
        reopened.save({"replacement": {"original": "Петров"}})

    assert calls == [False, False]
    assert path.read_bytes() == before


def test_new_secure_context_may_create_key(monkeypatch, tmp_path):
    pytest.importorskip("cryptography")
    calls = []

    def available(_service, _account, *, create_if_missing=True):
        calls.append(create_if_missing)
        return b"z" * 32

    monkeypatch.setattr(review_context, "_keychain_key", available)
    path = tmp_path / ".ctx"
    SecureContextVault(path).save({"finding": {"original": "Иванов"}})

    assert calls == [True]
    assert path.exists()


def test_keyring_backend_error_fails_closed_without_native_fallback(monkeypatch, tmp_path):
    """Ошибка активного keyring backend не должна менять ключ или тихо выбирать другой backend."""
    pytest.importorskip("cryptography")
    path = tmp_path / ".ctx"
    SecureContextVault(path, key_provider=lambda: b"k" * 32).save(
        {"finding": {"original": "Иванов"}}
    )
    before = path.read_bytes()

    def unavailable(*_args, **_kwargs):
        raise RuntimeError("keychain denied")

    monkeypatch.setitem(
        sys.modules,
        "keyring",
        SimpleNamespace(get_password=unavailable, set_password=unavailable),
    )
    monkeypatch.setattr(
        review_context.subprocess,
        "run",
        lambda *_args, **_kwargs: pytest.fail("native fallback must not run after keyring backend error"),
    )

    vault = SecureContextVault(path)
    with pytest.raises(SecureContextError, match="системного хранилища"):
        vault.load()

    assert path.read_bytes() == before


def test_pdf_rebuild_requires_exact_text_coordinate(tmp_path):
    fitz = pytest.importorskip("fitz")
    path = tmp_path / "result.pdf"
    document = fitz.open()
    page = document.new_page(width=300, height=150)
    page.insert_text((40, 80), "Name: [FIO_1]", fontsize=12)
    document.save(path)
    document.close()
    with fitz.open(path) as doc_reader:
        page = doc_reader[0]
        span = next(span for block in page.get_text("dict")["blocks"] for line in block.get("lines", []) for span in line.get("spans", []) if "[FIO_1]" in span["text"])
    assert rebuild_pdf(path, {"page": 0, "page_base": 0, "bbox": span["bbox"], "text": span["text"]}, "[FIO_1]", "Ivanov")
    with fitz.open(path) as doc_reader:
        assert "Ivanov" in doc_reader[0].get_text()
