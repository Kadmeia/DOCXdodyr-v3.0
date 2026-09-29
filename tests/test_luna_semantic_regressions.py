# -*- coding: utf-8 -*-

import pytest
from docx import Document

from backend_api import BackendApi
from legal_pullenti import iter_legal_identifier_spans


@pytest.fixture(scope="module")
def api():
    return BackendApi()


def _clean(api, text):
    return api.anonymize_text_pullenti(
        text, set(), set(), mapping_dict={}, entity_seen={}
    )[0]


def test_legal_common_nouns_are_not_treated_as_people(api):
    text = (
        "12.3 Договор должен быть зарегистрирован Заказчиком "
        "в Реестре Договоров, заключенных заказчиками."
    )

    assert _clean(api, text) == text


def test_punctuation_is_not_treated_as_a_person(api):
    text = "получателем) в соответствии с договором"

    assert _clean(api, text) == text


def test_internal_structural_unit_is_not_an_organization(api):
    text = "Начальник Управления развития"

    assert _clean(api, text) == text


@pytest.mark.parametrize(
    "source",
    [
        "члена Совета директоров Общества",
        "иск подан в Варненский окружной суд Республики Болгарии",
    ],
)
def test_governing_body_and_public_court_are_not_private_organizations(api, source):
    assert _clean(api, source) == source


def test_inflected_legal_education_term_is_not_a_personal_education_field(api):
    source = (
        "договор с адвокатским образованием Республики Болгарии на "
        "представление интересов в суде"
    )

    assert _clean(api, source) == source


def test_legal_code_citation_is_not_an_address(api):
    source = "толкование ст. 1119 ГК. В частности, применяется ст. 168 ГК."

    assert _clean(api, source) == source


@pytest.mark.parametrize(
    ("source", "expected"),
    [
        ("ООО «Ромашка»", "ООО «[Наименование_1]»"),
        ("ПАО СБЕРБАНК", "ПАО [Наименование_1]"),
        (
            "Государственного автономного учреждения культуры города Москвы "
            '"Московское агентство организации отдыха и туризма"',
            "Государственного автономного учреждения культуры города Москвы "
            '"[Наименование_1]"',
        ),
    ],
)
def test_organization_form_is_preserved_while_name_is_masked(api, source, expected):
    assert _clean(api, source) == expected


def test_full_document_date_is_detected_but_public_year_is_preserved(api):
    source = "Акт от «09» января 2025 г.; форма действует в 2026 году."
    cleaned = _clean(api, source)

    assert "2025" not in cleaned
    assert "2026 году" in cleaned
    assert cleaned.count("[Дата_") == 1


def test_npa_dates_and_public_authority_are_preserved(api):
    source = (
        "Положение утверждено Постановлением Правительства РФ от 16.03.2009 N 228. "
        "Требования утверждены Приказом Роскомнадзора от 24.02.2021 N 18."
    )

    assert _clean(api, source) == source


def test_judicial_act_date_is_preserved_and_case_number_is_complete(api):
    source = (
        "Определение Судебной коллегии по экономическим спорам ВС РФ "
        "от 11 июня 2020 г. N 306-ЭС19-24912."
    )

    cleaned = _clean(api, source)

    assert "11 июня 2020 г." in cleaned
    assert "N [НомерДела_1]" in cleaned
    assert "N 30[" not in cleaned


def test_salutation_masks_the_complete_three_part_name(api):
    source = "Уважаемый Сергей Александрович Смирнов!"

    assert _clean(api, source) == "Уважаемый [ФИО_1]!"


def test_reference_and_historical_years_are_preserved(api):
    source = (
        "Готовое решение (КонсультантПлюс, 2025). "
        "Правило сформулировано в относительно далеком 2001 году, "
        "а в 2014 году подход изменился."
    )

    assert _clean(api, source) == source


def test_strict_eis_procurement_number_is_detected_without_nearby_label(api):
    source = "Услуги/ № 0373200101025000008 - 2 911 244,00 ₽"
    cleaned = _clean(api, source)

    assert "0373200101025000008" not in cleaned
    assert "[НомерЗакупки_1]" in cleaned


def test_contract_number_can_follow_a_contract_date():
    source = "гражданско-правовой договор от 17 июля 2024 г. № КП7-КС"
    spans = list(iter_legal_identifier_spans(source))

    assert any(
        span.label == "CONTRACT_NUMBER" and span.text == "КП7-КС"
        for span in spans
    )


def test_email_uses_canonical_placeholder(api):
    assert _clean(api, "E-mail: name@example.com") == "E-mail: [Email_1]"


def test_consecutive_organization_names_keep_descriptors(api):
    source = 'ГАУК "МОСГОРТУР" "Гостиничный комплекс Кинопарк"'

    assert _clean(api, source) == (
        'ГАУК "[Наименование_1]" "Гостиничный комплекс [Наименование_2]"'
    )


@pytest.mark.parametrize(
    "source",
    [
        "Менеджер: Екатерина",
        "Вернуть Семенову в сумме",
        "Семенов будет вынужден обратиться",
    ],
)
def test_single_person_tokens_require_semantic_person_context(api, source):
    assert "[ФИО_1]" in _clean(api, source)


def test_standalone_city_and_jurisdiction_are_preserved(api):
    assert _clean(api, "г. Москва                  дата") == "г. Москва                  дата"
    assert _clean(api, "«Газпромбанк» (Акционерное общество) г.Москва").endswith("г.Москва")
    assert _clean(api, "сборные команды города Москвы") == "сборные команды города Москвы"


def test_city_is_masked_when_it_is_an_explicit_or_full_address(api):
    assert _clean(api, "Адрес регистрации: г. Москва") == "Адрес регистрации: [Адрес_1]"
    assert _clean(api, "г. Москва, ул. Тверская, д. 1") == "[Адрес_1]"


def test_semantic_table_address_masks_full_value_but_not_a_bare_bank_city(api):
    from table_entity_context import build_table_contexts

    rows = [
        ["Адрес Банка", "117420, г. Москва, ул. Наметкина, д.16, корпус 1"],
        ["Наименование Банка", "Банк ГБ (АО) г. Москва"],
    ]
    contexts = {
        (context.row, context.column): context
        for context in build_table_contexts(rows)
    }

    address = api.anonymize_text_pullenti(
        rows[0][1], mapping_dict={}, entity_seen={},
        semantic_context=contexts[(0, 1)],
    )[0]
    bank = api.anonymize_text_pullenti(
        rows[1][1], mapping_dict={}, entity_seen={},
        semantic_context=contexts[(1, 1)],
    )[0]

    assert address == "[Адрес_1]"
    assert bank.endswith("г. Москва")


def test_dotted_fio_table_header_masks_a_single_given_name(api):
    from table_entity_context import build_table_contexts

    rows = [["Должность", "Ф.И.О"], ["Администратор", "Света ( новый сотрудник)"]]
    contexts = {
        (context.row, context.column): context
        for context in build_table_contexts(rows)
    }

    cleaned = api.anonymize_text_pullenti(
        rows[1][1], mapping_dict={}, entity_seen={},
        semantic_context=contexts[(1, 1)],
    )[0]

    assert cleaned == "[ФИО_1] ( новый сотрудник)"


def test_fio_header_does_not_leak_into_neighbouring_size_column(api):
    from table_entity_context import build_table_contexts

    rows = [
        ["Должность", "Ф.И.О", "Размер"],
        ["Администратор", "Света", "Верх - 52, низ - 52"],
    ]
    contexts = {
        (context.row, context.column): context
        for context in build_table_contexts(rows)
    }

    header = api.anonymize_text_pullenti(
        rows[0][2], mapping_dict={}, entity_seen={},
        semantic_context=contexts[(0, 2)],
    )[0]
    size = api.anonymize_text_pullenti(
        rows[1][2], mapping_dict={}, entity_seen={},
        semantic_context=contexts[(1, 2)],
    )[0]

    assert header == "Размер"
    assert size == "Верх - 52, низ - 52"


def test_inflected_contract_marker_is_supported(api):
    cleaned = _clean(api, "К Агентскому договору № 5730 от 10.01.2023 г.")

    assert cleaned == "К Агентскому договору № [НомерДоговора_1] от [Дата_1] г."


def test_docx_replacement_fails_closed_when_text_cannot_be_written(api, monkeypatch):
    class RefusingParagraph:
        text = "Иванов"

        def __setattr__(self, name, value):
            if name != "text":
                super().__setattr__(name, value)

    monkeypatch.setattr("backend_api.replace_diff_in_paragraph_xml", lambda *args: False)

    with pytest.raises(RuntimeError, match="fallback"):
        api.replace_text_in_paragraph_xml(
            RefusingParagraph(), "Иванов", "[ФИО_1]", [], "[test] "
        )


def test_labelled_address_covers_the_complete_field(api):
    source = (
        "Адрес регистрации: 101000, г. Москва, пос. Первомайское, "
        "ул. Примерная, 3 к 1, кв. 25"
    )

    assert _clean(api, source) == "Адрес регистрации: [Адрес_1]"


def test_judicial_determination_number_is_a_case_number(api):
    source = (
        "Определением Судебной коллегии по гражданским делам Верховного Суда РФ "
        "от 13 октября 2015 года № 57-КГ15-7"
    )
    cleaned = _clean(api, source)

    assert "57-КГ15-7" not in cleaned
    assert "[НомерДела_1]" in cleaned


def test_internal_faculty_and_information_system_are_preserved(api):
    assert _clean(api, "на юридическом факультете им. М.М. Сперанского") == (
        "на юридическом факультете им. М.М. Сперанского"
    )
    assert _clean(api, "АИС «Мосгортур»") == "АИС «Мосгортур»"


def test_inflected_public_authorities_are_preserved(api):
    assert _clean(api, "Приказом Роскомнадзора") == "Приказом Роскомнадзора"
    assert _clean(api, "Постановления Пленума ВС РФ") == "Постановления Пленума ВС РФ"


def test_generic_job_field_instruction_is_not_a_job_value(api):
    assert _clean(api, "Должность сотрудника") == "Должность сотрудника"


def test_job_title_value_is_preserved(api):
    assert _clean(api, "Должность: профессор") == "Должность: профессор"


def test_quoted_bank_and_standalone_city_keep_stable_context(api):
    assert _clean(api, "«Газпромбанк» (Акционерное общество) г.Москва") == (
        "«[Наименование_1]» (Акционерное общество) г.Москва"
    )
    assert _clean(api, "Банк ГБ (АО) г. Москва") == (
        "Банк [Наименование_1] (АО) г. Москва"
    )


def test_lowercase_academic_titles_are_not_people(api):
    source = "Семенова Дмитрия Петровича, профессора и кандидата наук"
    cleaned = _clean(api, source)

    assert cleaned == "[ФИО_1], профессора и кандидата наук"


def test_full_fio_with_adjectival_surname_is_masked(api):
    assert _clean(api, "Николаевская Елена Федоровна  ( кандидат)") == (
        "[ФИО_1]  ( кандидат)"
    )


def test_salutation_masks_the_complete_name(api):
    assert _clean(api, "Уважаемый Сергей Смирнов") == "Уважаемый [ФИО_1]"


def test_long_organization_keeps_opf_alias_and_surrounding_clause(api):
    source = (
        'оспаривание в судебном порядке договоров аренды гостиничного комплекса, '
        'заключенных от имени Единоличного акционерного общества '
        '"Санаторно-оздоровительный комплекс "Восток" '
        '(далее – ЕАО "СОК "Восток", Общество);'
    )
    cleaned = _clean(api, source)

    assert cleaned.startswith(
        "оспаривание в судебном порядке договоров аренды гостиничного комплекса, "
        "заключенных от имени Единоличного акционерного общества "
    )
    assert '(далее – ЕАО "[Наименование_' in cleaned
    assert cleaned.endswith('", Общество);')


def test_multiline_company_name_in_table_cell_is_masked(api):
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Параметры"
    cell = table.cell(0, 1)
    cell.paragraphs[0].text = "Наименование юридического лица"
    cell.add_paragraph("Общество с ограниченной ответственностью")
    cell.add_paragraph("«Медиа-Групп»")

    api.clean_document(doc, set(), set(), mapping_dict={}, entity_seen={})

    assert cell.text == (
        "Наименование юридического лица\n"
        "Общество с ограниченной ответственностью\n"
        "«[Наименование_1]»"
    )


def test_labelled_company_identifiers_and_account_in_table_are_masked(api):
    doc = Document()
    table = doc.add_table(rows=4, cols=2)
    values = [
        ("ОГРН", "1197746012342", "[ОГРН_1]"),
        ("ИНН", "7702012345", "[ИНН_1]"),
        ("КПП", "771701001", "[КПП_1]"),
        ("Расчетный счет", "40702810500000015075", "[Р/с_1]"),
    ]
    for row, (label, value, _expected) in zip(table.rows, values):
        row.cells[0].text = label
        row.cells[1].text = value

    api.clean_document(doc, set(), set(), mapping_dict={}, entity_seen={})

    assert [row.cells[1].text for row in table.rows] == [item[2] for item in values]
