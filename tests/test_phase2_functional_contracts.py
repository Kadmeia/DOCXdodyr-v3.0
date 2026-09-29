# -*- coding: utf-8 -*-
"""Regression tests for Phase 2: Functional contracts, list persistence, crash recovery, and OOXML structures."""

import json
import os
from pathlib import Path
import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from backend_api import BackendApi
import crash_recovery
from constants import FORMAT_REGISTRY
from document_restorer import DocumentRestorer
from folder_pipeline import FolderAnonymizationPipeline, _load_continuation_state


def test_save_lists_disk_persistence(tmp_path, monkeypatch):
    """Проверяет, что save_lists гарантированно перезаписывает файлы исключений и замен на диске."""
    excl_file = tmp_path / "Исключения.txt"
    repl_file = tmp_path / "Замены.txt"
    excl_file.write_text("# Old exclusions\nСтароеСлово\n", encoding="utf-8")
    repl_file.write_text("# Old replacements\nСтараяЗамена\n", encoding="utf-8")

    api = BackendApi.__new__(BackendApi)
    api.exclusions_file_path = excl_file
    api.replacements_file_path = repl_file
    api.user_exclusions = {"СтароеСлово"}
    api.custom_replacements = {"СтараяЗамена"}
    api._window = None

    save_res = api.save_lists({
        "exclusions": ["НовоеИсключение1", "НовоеИсключение2"],
        "replacements": ["НоваяЗамена"],
    })
    assert save_res.get("success") is True

    # Проверяем файлы на диске
    excl_disk = excl_file.read_text(encoding="utf-8")
    assert "НовоеИсключение1" in excl_disk
    assert "НовоеИсключение2" in excl_disk
    assert "СтароеСлово" not in excl_disk

    repl_disk = repl_file.read_text(encoding="utf-8")
    assert "НоваяЗамена" in repl_disk
    assert "СтараяЗамена" not in repl_disk


def test_update_settings_patch_semantics():
    """Проверяет строгое патч-обновление: непереданные ключи не перезаписываются дефолтами."""
    api = BackendApi.__new__(BackendApi)
    api.settings = {
        "save_original": True,
        "save_docx": True,
        "save_pdf": True,
        "save_markdown": True,
        "ocr_lang": "rus+eng",
    }
    api.save_original = True
    api.save_docx = True
    api.save_pdf = True
    api.save_markdown = True
    api.ocr_lang = "rus+eng"
    api._save_settings = lambda: None

    import pdf_convert
    pdf_convert._OCR_BACKEND_CACHE["test"] = "dummy"

    # Частичное обновление только одного поля
    api.update_settings({"ocr_lang": "eng"})

    assert api.ocr_lang == "eng"
    assert len(pdf_convert._OCR_BACKEND_CACHE) == 0, "OCR кэш должен быть очищен при смене ocr_lang"
    assert api.save_pdf is True  # Не должно было сброситься в False
    assert api.save_markdown is True
    assert api.save_docx is True


def test_crash_recovery_encrypted_vault_and_resume(tmp_path, monkeypatch):
    """Проверяет безопасное хранение дешифратора в зашифрованном vault и продолжение сессии."""
    monkeypatch.setattr(crash_recovery, "get_recovery_dir", lambda: tmp_path)

    batch_id = "test-batch-uuid-1234"
    files = [str(tmp_path / "doc1.docx"), str(tmp_path / "doc2.docx")]

    # 1. Старт чекпоинта
    checkpoint = crash_recovery.start_batch_checkpoint(batch_id, files, mode="files")
    assert checkpoint.batch_id == batch_id

    # 2. Обновление прогресса с маппингом
    mapping_delta = {"[ФИО_1]": "Иванов Иван Иванович", "[ИНН_1]": "7701234567"}
    crash_recovery.update_batch_checkpoint(
        batch_id,
        processed_file=files[0],
        mapping_delta=mapping_delta,
        replacements=2,
    )

    # В незашифрованном чекпоинт-JSON не должно быть исходных PII
    raw_checkpoint = (tmp_path / crash_recovery.CHECKPOINT_FILENAME).read_text(encoding="utf-8")
    assert "Иванов Иван Иванович" not in raw_checkpoint
    assert "7701234567" not in raw_checkpoint

    # Проверяем, что зашифрованный vault существует
    assert (tmp_path / "active_batch_vault.bin").exists()

    # 3. Чтение прерванной сессии
    info = crash_recovery.get_interrupted_batch()
    assert info is not None
    assert info["batch_id"] == batch_id
    assert info["remaining_count"] == 1
    assert info["remaining_files"] == [str(Path(files[1]).resolve())]

    # 4. Расшифровка маппинга
    decrypted = crash_recovery.load_interrupted_batch_mapping(batch_id)
    assert decrypted["[ФИО_1]"] == "Иванов Иван Иванович"
    assert decrypted["[ИНН_1]"] == "7701234567"

    # 5. Сброс
    crash_recovery.discard_interrupted_batch()
    assert not (tmp_path / crash_recovery.CHECKPOINT_FILENAME).exists()
    assert not (tmp_path / "active_batch_vault.bin").exists()


def test_folder_continuation_without_manifest_rejected(tmp_path):
    """A decoder without a completed manifest is not a trusted folder run."""
    decoder_file = tmp_path / "дешифратор.json"
    mapping = {
        "[ФИО_1]": {
            "original": "Петров Петр Петрович",
            "type": "PER",
            "gender": "MASCULINE",
        }
    }
    decoder_file.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")

    with pytest.raises(ValueError, match="нет manifest завершённого запуска"):
        _load_continuation_state(tmp_path)


def test_ooxml_textbox_and_sdt_restoration(tmp_path, bind_decoder):
    """Проверяет корректное восстановление плейсхолдеров в Text Box и SDT элементах."""
    doc = Document()
    doc.add_paragraph("Обычный текст: [ФИО_1]")

    # Добавляем XML с w:txbxContent (Text Box)
    tb_xml = parse_xml(
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:pict><v:textbox><w:txbxContent>'
        '<w:p><w:r><w:t>Текст в рамке: [ФИО_1]</w:t></w:r></w:p>'
        '</w:txbxContent></v:textbox></w:pict></w:r></w:p>'
    )
    doc.element.body.append(tb_xml)

    # Добавляем XML с w:sdt (Structured Document Tag)
    sdt_xml = parse_xml(
        '<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:sdtContent>'
        '<w:p><w:r><w:t>Текст в поле: [ФИО_1]</w:t></w:r></w:p>'
        '</w:sdtContent></w:sdt>'
    )
    doc.element.body.append(sdt_xml)

    doc_path = tmp_path / "test_complex.docx"
    doc.save(str(doc_path))

    mapping = {"[ФИО_1]": "Сидоров Сидор Сидорович"}
    bind_decoder(doc_path, mapping)
    restorer = DocumentRestorer(mapping)
    ok, msg = restorer.restore_docx(str(doc_path))
    assert ok is True

    # Проверяем восстановленный документ
    restored_doc = Document(str(tmp_path / "test_complex_восстановлено.docx"))
    all_text = "".join(node.text or "" for node in restored_doc.element.iter())
    assert "Сидоров Сидор Сидорович" in all_text
    assert "[ФИО_1]" not in all_text


def test_format_registry_completeness():
    """Проверяет наполненность и валидность FORMAT_REGISTRY."""
    assert "anonymize" in FORMAT_REGISTRY
    assert "ocr" in FORMAT_REGISTRY
    assert "restore" in FORMAT_REGISTRY
    assert "hidden_inspect" in FORMAT_REGISTRY

    assert ".docx" in FORMAT_REGISTRY["anonymize"]["extensions"]
    assert ".xlsx" in FORMAT_REGISTRY["anonymize"]["extensions"]
    assert ".pdf" in FORMAT_REGISTRY["anonymize"]["extensions"]
    assert ".jpg" in FORMAT_REGISTRY["anonymize"]["extensions"]
