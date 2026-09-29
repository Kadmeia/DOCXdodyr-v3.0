from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook


PII = {
    "person": "Иванов Иван Иванович",
    "phone": "+7 900 123-45-67",
    "email": "ivanov@example.invalid",
}


def generate_fixtures(root: Path) -> dict[str, str]:
    """Create tiny local fixtures; existing large test fixtures are reused."""

    root.mkdir(parents=True, exist_ok=True)
    docx_path = root / "antigravity_contract.docx"
    if not docx_path.exists():
        document = Document()
        document.add_paragraph(f"Подписант: {PII['person']}; телефон: {PII['phone']}")
        document.sections[0].header.paragraphs[0].text = f"Сторона: {PII['person']}"
        document.sections[0].footer.paragraphs[0].text = "Служебная копия"
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Email"
        table.cell(0, 1).text = PII["email"]
        document.save(docx_path)

    xlsx_path = root / "antigravity_sheet.xlsx"
    if not xlsx_path.exists():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Данные"
        sheet["A1"] = "ФИО"
        sheet["B1"] = PII["person"]
        sheet["A2"] = "Телефон"
        sheet["B2"] = PII["phone"]
        workbook.save(xlsx_path)

    pdf_path = root / "antigravity_scan.pdf"
    if not pdf_path.exists():
        try:
            import fitz
            document = fitz.open()
            page = document.new_page(width=300, height=160)
            page.insert_text((30, 70), "Name: [FIO_1]", fontsize=12)
            document.save(pdf_path)
            document.close()
        except Exception:
            # PDF-dependent scenarios skip cleanly when PyMuPDF is absent.
            pass

    decoder_path = root / "antigravity_contract_дешифратор.json"
    if not decoder_path.exists():
        decoder_path.write_text(json.dumps({"[ФИО_1]": PII["person"]}, ensure_ascii=False), encoding="utf-8")
    return {"docx": str(docx_path), "xlsx": str(xlsx_path), "pdf": str(pdf_path), "decoder": str(decoder_path)}

