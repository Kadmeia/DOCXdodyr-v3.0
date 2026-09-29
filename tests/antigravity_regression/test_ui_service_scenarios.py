from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

import main
from backend_api import BackendApi
from hidden_data import flatten_pdf_to_images
from privacy_audit import ReviewFinding, ReviewQueue
from qwen_postprocessor import QwenPlaceholderPostprocessor

from .fixtures import generate_fixtures


class Window:
    def __init__(self, selected=()):
        self.selected = tuple(str(item) for item in selected)
        self.dialog_calls = []
        self.js_calls = []

    def create_file_dialog(self, *args, **kwargs):
        self.dialog_calls.append((args, kwargs))
        return self.selected

    def evaluate_js(self, script):
        self.js_calls.append(str(script))


class ApiStub:
    def __init__(self):
        self.settings = {}
        self.qwen_settings = SimpleNamespace(enabled=False)
        self.irreversible_pdf = False
        self.save_original = True
        self.save_pdf = False
        self.save_decoder = False
        self.saved = 0
        self.files = []
        self.folders = []
        self.auto_decoder = True

    def set_window(self, window):
        self.window = window

    def _save_settings(self):
        self.saved += 1

    def save_qwen_settings(self, data):
        self.qwen_settings.enabled = bool((data or {}).get("enabled", False))
        return {"success": True}

    def process_files(self, files):
        self.files.append(list(files))

    def process_folder(self, folder, **_kwargs):
        self.folders.append(str(folder))


@pytest.mark.antigravity_ui
def test_launch_theme_and_preflight_controls(monkeypatch):
    web = Path(__file__).parents[2] / "web"
    html = (web / "index.html").read_text(encoding="utf-8")
    script = (web / "script.js").read_text(encoding="utf-8")
    for control in ("btn-theme-light", "btn-theme-dark", "btn-theme-system", "cb-save-original", "cb-save-pdf", "cb-irreversible-pdf", "cb-qwen-postprocess"):
        assert f'id="{control}"' in html
    assert 'id="cb-continue-folder"' in html
    assert "'continue_folder'" in script
    assert "function setTheme" in script and "localStorage.setItem('theme'" in script
    assert "setGlobalProgress" in script

    monkeypatch.setattr(main, "BackendApi", ApiStub)
    wrapper = main.ApiWrapper()
    wrapper.update_settings({"save_original": False, "save_pdf": True, "irreversible_pdf": True, "qwen_enabled": False})
    assert wrapper._api.irreversible_pdf is True
    assert wrapper._api.save_pdf is True
    assert wrapper._api.saved == 1


@pytest.mark.antigravity_ui
def test_files_folder_and_restore_dropzones(monkeypatch, tmp_path):
    monkeypatch.setattr(main, "BackendApi", ApiStub)
    wrapper = main.ApiWrapper()
    source = tmp_path / "case.docx"
    source.write_bytes(b"fixture")
    window = Window([source])
    wrapper.set_window(window)
    wrapper.open_file_dialog("anonymize_docs")
    assert wrapper._api.files == [[str(source)]]
    assert window.dialog_calls[-1][1]["allow_multiple"] is True
    assert window.dialog_calls[-1][1]["file_types"][0].startswith("Все поддерживаемые")
    assert "*.docx" in window.dialog_calls[-1][1]["file_types"][0]
    assert "*.jpg" in window.dialog_calls[-1][1]["file_types"][0]

    started = []

    class Thread:
        def __init__(self, *, target, daemon):
            self.target, self.daemon = target, daemon

        def start(self):
            started.append(self)

    monkeypatch.setattr(main.threading, "Thread", Thread)
    wrapper.process_folder_dialog(tmp_path)
    assert started and started[0].daemon is True
    started[0].target()
    assert wrapper._api.folders == [str(tmp_path)]
    assert any("setFolderStatus" in call for call in window.js_calls)


def _review_api():
    api = BackendApi.__new__(BackendApi)
    api.review_queue = ReviewQueue([
        ReviewFinding("contract.docx", "body.0", "FIO", "[ФИО_1]", .95, position={"scope": "body", "kind": "paragraph", "index": 0}),
        ReviewFinding("table.xlsx", "sheet!B2", "PHONE", "[ТЕЛЕФОН_1]", .61, position={"scope": "sheet", "kind": "cell", "sheet": "Лист1", "cell": "B2"}),
        ReviewFinding("scan.pdf", "page.0", "ADDRESS", "[АДРЕС_1]", .88, position={"scope": "pdf", "kind": "ocr", "page": 0, "bbox": [.1, .2, .3, .1]}),
    ])
    api._persist_review_state = lambda: None
    return api


@pytest.mark.antigravity_ui
def test_review_filter_navigation_and_batch_action():
    api = _review_api()
    payload = api.get_review_findings(entity_type="FIO", min_confidence=90)
    assert payload["filtered_summary"]["total"] == 1
    finding_id = api.review_queue.items()[2].finding_id
    preview = api.get_review_finding_preview(finding_id)
    assert preview["position"]["bbox"] == [.1, .2, .3, .1]
    assert "redacted_context" not in preview and "original" not in preview
    result = api.batch_decide_review_findings([], "accepted", document_ref=".pdf", min_confidence=80)
    assert result["count"] == 1 and api.review_queue.summary()["accepted"] == 1


@pytest.mark.antigravity_ui
def test_hidden_content_choices_require_confirmation_and_are_aggregate_only(tmp_path):
    source = tmp_path / "hidden.docx"
    doc = Document()
    doc.add_paragraph("Открытый текст")
    doc.save(source)
    api = BackendApi.__new__(BackendApi)
    api.hidden_data_policy = api._default_hidden_data_policy()
    api.settings = {}
    api._save_settings = lambda: None
    api._hidden_inspections = {}
    payload = api.inspect_hidden_file(source)
    assert payload["inspection_token"] and payload["source_type"] == "docx"
    assert "original" not in json.dumps(payload, ensure_ascii=False)
    policy = api.save_hidden_data_policy({"comments": "remove"})
    assert policy["requires_confirmation"] is True
    assert api.save_hidden_data_policy({"comments": "remove"}, confirm_destructive=True)["success"] is True
    assert api.save_hidden_data_policy({"embedded_files": "remove"})["requires_confirmation"] is True
    assert api.save_hidden_data_policy({"embedded_files": "remove"}, confirm_destructive=True)["success"] is True


@pytest.mark.antigravity_ui
def test_restore_auto_decoder_finds_nearby_decoder(tmp_path, bind_decoder):
    source = tmp_path / "contract_cleaned.docx"
    doc = Document()
    doc.add_paragraph("Подписант: [ФИО_1]")
    doc.save(source)
    decoder = tmp_path / "contract_дешифратор.json"
    decoder.write_text(json.dumps({"[ФИО_1]": "Иванов Иван Иванович"}, ensure_ascii=False), encoding="utf-8")
    bind_decoder(source, {"[ФИО_1]": "Иванов Иван Иванович"})
    api = BackendApi.__new__(BackendApi)
    api.restore_doc_paths = [str(source)]
    api.restore_json_paths = []
    api.auto_decoder = True
    api.settings = {"auto_decoder": True}
    api._window = Window()
    api._run_restore()
    restored = tmp_path / "contract_cleaned_восстановлено.docx"
    assert restored.exists()
    assert "Иванов" in Document(restored).paragraphs[0].text


@pytest.mark.antigravity_ui
def test_irreversible_pdf_is_explicit_and_text_layer_is_removed(tmp_path):
    fitz = pytest.importorskip("fitz")
    source = tmp_path / "source.pdf"
    target = tmp_path / "flat.pdf"
    document = fitz.open()
    page = document.new_page(width=300, height=160)
    page.insert_text((30, 70), "Sensitive text", fontsize=12)
    document.save(source)
    document.close()
    report = flatten_pdf_to_images(source, target)
    assert report.text_layer_removed is True
    assert "Sensitive text" not in fitz.open(target)[0].get_text()


@pytest.mark.antigravity_ui
def test_qwen_offline_preflight_never_loads_missing_weights():
    loaded = []
    processor = QwenPlaceholderPostprocessor({"enabled": True, "local_files_only": True, "model_path": "missing"}, model_loader=lambda settings: loaded.append(settings) or (_ for _ in ()).throw(FileNotFoundError("cache miss")))
    text = "Для [ФИО_1] готов документ."
    assert processor.process(text) == text
    assert processor.process(text) == text
    assert len(loaded) == 1 and processor.last_status.startswith("unavailable:")
