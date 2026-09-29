import json

import pytest
from docx import Document

from document_restorer import DocumentRestorer, find_decoder_near_document


@pytest.mark.parametrize("tamper", ["missing", "output", "decoder", "run", "schema", "state", "foreign"])
def test_rejects_unbound_or_tampered_decoder(tmp_path, bind_decoder, tamper):
    source = tmp_path / "synthetic.docx"
    doc = Document()
    doc.add_paragraph("Hello [NAME_1]")
    doc.save(source)
    mapping = {"[NAME_1]": "Synthetic Person"}
    decoder = bind_decoder(source, mapping)
    sidecar = source.with_name(source.name + ".provenance.json")
    manifest = decoder.with_name(decoder.stem + ".manifest.json")
    if tamper == "missing":
        sidecar.unlink()
    elif tamper == "output":
        with source.open("ab") as stream:
            stream.write(b"tampered")
    elif tamper == "decoder":
        decoder.write_text('{}')
    elif tamper in ("run", "schema"):
        data = json.loads(sidecar.read_text())
        data["run_id" if tamper == "run" else "schema"] = "foreign"
        sidecar.write_text(json.dumps(data))
    elif tamper == "state":
        data = json.loads(manifest.read_text())
        data["complete"] = False
        manifest.write_text(json.dumps(data))
    else:
        mapping = {"[NAME_1]": "Foreign Person"}
    with pytest.raises(ValueError):
        DocumentRestorer(mapping).restore_docx(source)
    assert not source.with_name("synthetic_восстановлено.docx").exists()


def test_bound_roundtrip_changes_actual_text(tmp_path, bind_decoder):
    source = tmp_path / "synthetic.docx"
    doc = Document()
    doc.add_paragraph("Hello [NAME_1]")
    doc.save(source)
    mapping = {"[NAME_1]": "Synthetic Person"}
    decoder = bind_decoder(source, mapping)
    assert find_decoder_near_document(source) == decoder
    output = tmp_path / "restored.docx"
    assert DocumentRestorer(mapping).restore_docx(source, output)[0]
    assert Document(output).paragraphs[0].text == "Hello Synthetic Person"
