from __future__ import annotations

import json
import re
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from backend_api import Worker
from document_restorer import find_decoder_near_document
from folder_pipeline import (
    DEFAULT_DECODER_FILENAME,
    DEFAULT_OUTPUT_DIR_NAME,
    FolderAnonymizationPipeline,
    MAX_INPUT_BYTES,
    _secure_input_snapshot,
    collect_supported_files,
)


def test_secure_input_snapshot_rejects_symlink(tmp_path):
    target = tmp_path / "real.docx"
    target.write_bytes(b"private")
    link = tmp_path / "input.docx"
    try:
        link.symlink_to(target)
    except OSError:
        pytest.skip("Symlink creation requires administrative privileges on Windows")
    with pytest.raises(OSError):
        _secure_input_snapshot(link, tmp_path / "snapshots")


def test_oversized_input_is_reported_not_silently_omitted(tmp_path):
    source = tmp_path / "large.pdf"
    with source.open("wb") as stream:
        stream.truncate(MAX_INPUT_BYTES + 1)
    assert collect_supported_files(tmp_path) == [source.resolve()]
    with pytest.raises(ValueError, match="превышает допустимый размер"):
        _secure_input_snapshot(source, tmp_path / "snapshots")


def test_symlink_swap_after_scan_does_not_process_private_file(tmp_path):
    source = tmp_path / "input.docx"
    Document().save(source)
    private_dir = tmp_path / ".private"
    private_dir.mkdir()
    private_file = private_dir / "secret.docx"
    Document().save(private_file)

    def swap_after_scan(index, _total, _message):
        if index == 0 and source.is_file() and not source.is_symlink():
            source.unlink()
            try:
                source.symlink_to(private_file)
            except OSError:
                pytest.skip("Symlink creation requires administrative privileges on Windows")

    result = FolderAnonymizationPipeline(_backend()).process(
        tmp_path, progress_callback=swap_after_scan
    )
    assert result.processed_count == 0
    assert result.error_count == 1
    assert not (tmp_path / DEFAULT_OUTPUT_DIR_NAME).exists()


def test_parent_directory_symlink_swap_after_scan_is_rejected(tmp_path):
    incoming = tmp_path / "incoming"
    incoming.mkdir()
    Document().save(incoming / "input.docx")
    private_dir = tmp_path / ".private"
    private_dir.mkdir()
    Document().save(private_dir / "input.docx")

    def swap_parent_after_scan(index, _total, _message):
        if index == 0 and incoming.is_dir() and not incoming.is_symlink():
            (incoming / "input.docx").unlink()
            incoming.rmdir()
            try:
                incoming.symlink_to(private_dir, target_is_directory=True)
            except OSError:
                pytest.skip("Symlink creation requires administrative privileges on Windows")

    result = FolderAnonymizationPipeline(_backend()).process(
        tmp_path, progress_callback=swap_parent_after_scan
    )
    assert result.processed_count == 0
    assert result.error_count == 1
    assert not (tmp_path / DEFAULT_OUTPUT_DIR_NAME).exists()


class _Occurrence:
    def __init__(self, source, start, end):
        self.sofa = source
        self.begin_char = start
        self.end_char = end - 1


class _PersonReferent:
    type_name = "PERSON"

    def __init__(self, source, start, end):
        self.occurrence = [_Occurrence(source, start, end)]


class _LocalPersonProcessor:
    _name = re.compile(
        r"(?:Иванов(?:у|ым)?\s+Иван(?:у|ом)?\s+Иванович(?:у|ем)?|"
        r"Петров(?:у|ым)?\s+Петр(?:у|ом)?\s+Петрович(?:у|ем)?)",
        re.I,
    )

    def process(self, source, *_args):
        return SimpleNamespace(
            entities=[_PersonReferent(source, m.start(), m.end()) for m in self._name.finditer(source.text)]
        )


def _backend():
    from backend_api import BackendApi

    api = BackendApi.__new__(BackendApi)
    api._pullenti_processor = _LocalPersonProcessor()
    api.current_placeholders = {"PER": "[ФИО]"}
    api.qwen_settings = SimpleNamespace(enabled=False)
    api._qwen_postprocessor = SimpleNamespace(last_status="disabled")
    api.user_exclusions = set()
    api.custom_replacements = set()
    api.save_original = True
    api.save_pdf = False
    api.save_decoder = True
    api.ocr_pdf = False
    api._window = None
    api.emit_audit_sidecars = True
    api.save_log_files = True
    api.save_audit_files = True
    return api


def _documents(root: Path):
    first = root / "договор.docx"
    doc = Document()
    doc.add_paragraph("Подписант: Иванов Иван Иванович")
    doc.save(first)

    nested = root / "входящие"
    nested.mkdir()
    second = nested / "анкета.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "ФИО"
    ws["B1"] = "Иванову Ивану Ивановичу"
    wb.save(second)
    return first, second


def test_collect_supported_files_excludes_output_subtree(tmp_path):
    first, second = _documents(tmp_path)
    output = tmp_path / DEFAULT_OUTPUT_DIR_NAME
    output.mkdir()
    Document().save(output / "already.docx")
    (tmp_path / "readme.txt").write_text("не документ", encoding="utf-8")

    found = collect_supported_files(tmp_path)
    assert found == sorted([first.resolve(), second.resolve()], key=lambda p: str(p.relative_to(tmp_path)).casefold())


@pytest.mark.parametrize("unsafe_name", ["../outside", "/tmp/outside"])
def test_output_dir_name_rejects_path_escape(unsafe_name):
    with pytest.raises(ValueError, match="Имя выходной папки"):
        FolderAnonymizationPipeline(_backend(), output_dir_name=unsafe_name)


def test_output_dir_name_accepts_simple_name():
    pipeline = FolderAnonymizationPipeline(_backend(), output_dir_name="Безопасные результаты")
    assert pipeline.output_dir_name == "Безопасные результаты"


def test_folder_cancel_during_final_reconciliation_is_not_success(tmp_path):
    """Cancellation after file processing must be represented as an error."""
    source = tmp_path / "договор.docx"
    doc = Document()
    doc.add_paragraph("Подписант: Иванов Иван Иванович")
    doc.save(source)
    api = _backend()
    checks = 0

    def is_cancelled():
        nonlocal checks
        checks += 1
        # First check starts the file; second check is the final-sweep guard.
        return checks >= 2

    api.is_cancelled = is_cancelled
    result = FolderAnonymizationPipeline(api).process(tmp_path)

    assert result.processed_count == 1
    assert result.error_count > 0
    assert result.errors
    assert not (tmp_path / DEFAULT_OUTPUT_DIR_NAME).exists()


@pytest.mark.parametrize('tamper', ['run_id', 'output', 'path'])
def test_continuation_requires_bound_outputs(tmp_path, tamper):
    from folder_pipeline import _load_continuation_state
    _documents(tmp_path)
    result = FolderAnonymizationPipeline(_backend()).process(tmp_path)
    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    if tamper == 'run_id':
        manifest.pop('run_id')
    elif tamper == 'path':
        manifest['documents'][0]['outputs'][0]['path'] = '../foreign.docx'
    else:
        output = tmp_path / manifest['documents'][0]['outputs'][0]['path']
        output.write_bytes(b'synthetic tamper')
    result.manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(ValueError):
        _load_continuation_state(tmp_path)


def test_folder_pipeline_writes_nested_outputs_and_one_root_decoder(tmp_path):
    first, second = _documents(tmp_path)
    api = _backend()
    result = FolderAnonymizationPipeline(api).process(tmp_path)

    assert result.processed_count == 2
    assert result.error_count == 0
    assert result.decoder_path == tmp_path / DEFAULT_DECODER_FILENAME
    assert (tmp_path / DEFAULT_OUTPUT_DIR_NAME / "договор_cleaned.docx").exists()
    assert (tmp_path / DEFAULT_OUTPUT_DIR_NAME / "входящие" / "анкета_cleaned.xlsx").exists()
    assert first.exists() and second.exists()
    mapping = json.loads(result.decoder_path.read_text(encoding="utf-8"))
    assert list(mapping) == ["[ФИО_1]"]
    assert "Иванов" not in " ".join(p.text for p in Document(
        tmp_path / DEFAULT_OUTPUT_DIR_NAME / "договор_cleaned.docx"
    ).paragraphs)
    assert load_workbook(tmp_path / DEFAULT_OUTPUT_DIR_NAME / "входящие" / "анкета_cleaned.xlsx").active["B1"].value == "[ФИО_1]"
    docx_audit = tmp_path / DEFAULT_OUTPUT_DIR_NAME / "договор_cleaned.docx.audit.json"
    xlsx_audit = tmp_path / DEFAULT_OUTPUT_DIR_NAME / "входящие" / "анкета_cleaned.xlsx.audit.json"
    assert docx_audit.exists() and xlsx_audit.exists()
    audit_text = docx_audit.read_text(encoding="utf-8")
    assert '"sha256"' in audit_text
    assert "Иванов" not in audit_text
    audit_payload = json.loads(audit_text)
    assert audit_payload["counts"]["entities_by_type"]["PER"] >= 1
    assert audit_payload["counts"]["review"]["pending"] >= 1
    log_text = (tmp_path / DEFAULT_OUTPUT_DIR_NAME / "договор_cleaned_log.txt").read_text(encoding="utf-8")
    assert "Иванов Иван Иванович" not in log_text
    assert result.manifest_path.exists()
    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    assert manifest["schema"] == "docxdodyr.batch-manifest/v1"
    assert manifest["complete"] is True
    assert manifest["expected_document_count"] == 2
    assert manifest["decoder"]["sha256"]
    provenance = tmp_path / DEFAULT_OUTPUT_DIR_NAME / "договор_cleaned.docx.provenance.json"
    assert provenance.exists()
    assert find_decoder_near_document(tmp_path / DEFAULT_OUTPUT_DIR_NAME / "договор_cleaned.docx") == result.decoder_path


def test_existing_invalid_provenance_never_falls_back_to_newest_decoder(tmp_path):
    _documents(tmp_path)
    result = FolderAnonymizationPipeline(_backend()).process(tmp_path)
    document = tmp_path / DEFAULT_OUTPUT_DIR_NAME / "договор_cleaned.docx"
    provenance_path = document.with_name(document.name + ".provenance.json")
    provenance = json.loads(provenance_path.read_text(encoding="utf-8"))
    provenance["schema"] = "invalid"
    provenance_path.write_text(json.dumps(provenance, ensure_ascii=False), encoding="utf-8")

    with pytest.raises(ValueError, match="безопасная привязка"):
        find_decoder_near_document(document)


def test_broken_provenance_symlink_never_uses_legacy_fallback(tmp_path):
    _documents(tmp_path)
    FolderAnonymizationPipeline(_backend()).process(tmp_path)
    document = tmp_path / DEFAULT_OUTPUT_DIR_NAME / "договор_cleaned.docx"
    provenance_path = document.with_name(document.name + ".provenance.json")
    provenance_path.unlink()
    try:
        provenance_path.symlink_to(tmp_path / "missing-provenance.json")
    except OSError:
        pytest.skip("Symlink creation requires administrative privileges on Windows")

    with pytest.raises(ValueError, match="provenance является ссылкой"):
        find_decoder_near_document(document)


def test_folder_pipeline_exposes_redacted_review_queue_only(tmp_path):
    _documents(tmp_path)
    api = _backend()
    FolderAnonymizationPipeline(api).process(tmp_path)

    payload = api.get_review_findings()
    assert payload["summary"]["pending"] >= 2
    serialized = json.dumps(payload, ensure_ascii=False)
    assert "Иванов" not in serialized
    assert "[ФИО_1]" in serialized
    finding_id = payload["items"][0]["finding_id"]
    decision = api.decide_review_finding(finding_id, "accepted")
    assert decision["item"]["status"] == "accepted"


def test_rerun_preserves_old_decoder_and_auto_find_chooses_newest(tmp_path):
    first, _second = _documents(tmp_path)
    old_mapping = {"[ФИО_99]": "Старое лицо"}
    (tmp_path / "дешифратор.json").write_text(json.dumps(old_mapping), encoding="utf-8")
    result = FolderAnonymizationPipeline(_backend()).process(tmp_path)

    assert result.decoder_path.name in ("дешифратор_1.json", "Дешифратор_1.json")
    assert json.loads((tmp_path / "дешифратор.json").read_text(encoding="utf-8")) == old_mapping
    found = find_decoder_near_document(
        tmp_path / DEFAULT_OUTPUT_DIR_NAME / "договор_cleaned.docx"
    )
    assert found == result.decoder_path


def test_continue_folder_reuses_ids_and_allocates_only_new_ids(tmp_path):
    first, second = _documents(tmp_path)
    initial = FolderAnonymizationPipeline(_backend()).process(tmp_path)
    initial_mapping = json.loads(initial.decoder_path.read_text(encoding="utf-8"))
    assert list(initial_mapping) == ["[ФИО_1]"]

    doc = Document()
    doc.add_paragraph("Участники: Иванову Ивану Ивановичу и Петров Петр Петрович")
    doc.save(first)
    second.unlink()

    continued = FolderAnonymizationPipeline(_backend()).process(
        tmp_path, continue_existing=True,
    )

    assert continued.decoder_path == tmp_path / DEFAULT_DECODER_FILENAME.replace(".json", "_1.json")
    assert json.loads(initial.decoder_path.read_text(encoding="utf-8")) == initial_mapping
    assert list(continued.mapping) == ["[ФИО_1]", "[ФИО_2]"]
    output_text = " ".join(
        paragraph.text for paragraph in Document(
            tmp_path / DEFAULT_OUTPUT_DIR_NAME / "договор_cleaned.docx"
        ).paragraphs
    )
    assert "[ФИО_1]" in output_text
    assert "[ФИО_2]" in output_text
    assert "Иванов" not in output_text and "Петров" not in output_text
    manifest = json.loads(continued.manifest_path.read_text(encoding="utf-8"))
    assert manifest["continuation"]["run_id"]
    assert manifest["continuation"]["decoder_path"] == initial.decoder_path.name


def test_continue_folder_requires_verified_completed_run(tmp_path):
    _documents(tmp_path)
    with pytest.raises(ValueError, match="нет manifest завершённого запуска"):
        FolderAnonymizationPipeline(_backend()).process(tmp_path, continue_existing=True)
    assert not list(tmp_path.glob(".docxdodyr-incomplete-*"))


def test_continue_folder_rejects_tampered_decoder(tmp_path):
    _documents(tmp_path)
    initial = FolderAnonymizationPipeline(_backend()).process(tmp_path)
    initial.decoder_path.write_text("{}", encoding="utf-8")

    with pytest.raises(ValueError, match="контрольная сумма"):
        FolderAnonymizationPipeline(_backend()).process(tmp_path, continue_existing=True)
    assert not list(tmp_path.glob(".docxdodyr-incomplete-*"))


def test_restore_auto_decoder_can_be_disabled(tmp_path):
    from backend_api import BackendApi

    document = tmp_path / "cleaned.docx"
    doc = Document()
    doc.add_paragraph("[ФИО_1]")
    doc.save(document)
    (tmp_path / "дешифратор.json").write_text(
        json.dumps({"[ФИО_1]": "Иванов Иван Иванович"}), encoding="utf-8"
    )

    class _Window:
        def __init__(self):
            self.calls = []

        def evaluate_js(self, script):
            self.calls.append(script)

    api = BackendApi.__new__(BackendApi)
    api.restore_doc_paths = [str(document)]
    api.restore_json_paths = []
    api.auto_decoder = False
    api.settings = {}
    api._window = _Window()
    api._run_restore()

    assert any("дешифратор" in call for call in api._window.calls)
    assert not (tmp_path / "cleaned_восстановлено.docx").exists()


def test_unbound_nested_document_rejects_root_decoder_automatically(tmp_path):
    from backend_api import BackendApi

    generated = tmp_path / "Обезличенные документы" / "анализ"
    generated.mkdir(parents=True)
    document = generated / "новый_иск.docx"
    doc = Document()
    doc.add_paragraph("Требования заявлены в интересах [ФИО_1].")
    doc.save(document)
    (tmp_path / "дешифратор.json").write_text(
        json.dumps({"[ФИО_1]": "Иванов Иван Иванович"}, ensure_ascii=False),
        encoding="utf-8",
    )

    api = BackendApi.__new__(BackendApi)
    api.restore_doc_paths = [str(document)]
    api.restore_json_paths = []
    api.auto_decoder = True
    api.settings = {"auto_decoder": True}
    api._window = None
    api._run_restore()

    restored = generated / "новый_иск_восстановлено.docx"
    assert not restored.exists()


def test_collect_supported_files_ignores_lock_and_hidden_files(tmp_path):
    # Real documents
    (tmp_path / "valid.docx").write_text("test")
    (tmp_path / "valid.xlsx").write_text("test")
    # Lock and hidden files
    (tmp_path / "~$valid.docx").write_text("lock")
    (tmp_path / ".~temp.docx").write_text("temp")
    (tmp_path / ".DS_Store").write_text("mac")

    found = collect_supported_files(tmp_path)
    found_names = [f.name for f in found]
    assert "valid.docx" in found_names
    assert "valid.xlsx" in found_names
    assert "~$valid.docx" not in found_names
    assert ".~temp.docx" not in found_names
    assert ".DS_Store" not in found_names


def test_folder_pipeline_clean_mode_creates_only_documents_and_one_decoder(tmp_path):
    """Проверяет режим чистой папки пользователя: только обезличенные документы и 1 дешифратор."""
    # Создаем 5 документов
    for i in range(1, 6):
        doc = Document()
        doc.add_paragraph(f"Документ #{i} подписан: Иванов Иван Иванович")
        doc.save(str(tmp_path / f"договор_{i}.docx"))

    api = _backend()
    api.emit_audit_sidecars = False
    api.save_log_files = False
    api.save_audit_files = False

    pipeline = FolderAnonymizationPipeline(api, emit_sidecars=False)
    result = pipeline.process(tmp_path)

    assert result.processed_count == 5
    assert result.error_count == 0

    output_dir = tmp_path / DEFAULT_OUTPUT_DIR_NAME
    assert output_dir.is_dir()
    all_output_files = [f for f in output_dir.rglob("*") if f.is_file()]

    # Ровно 5 обезличенных документов в папке вывода без каких-либо sidecar-файлов
    assert len(all_output_files) == 5
    for f in all_output_files:
        assert f.suffix == ".docx"
        assert not f.name.endswith(".provenance.json")
        assert not f.name.endswith("_log.txt")
        assert not f.name.endswith(".audit.json")

    # И ровно 1 файл дешифратора в корне (со встроенной привязкой v2)
    decoder_file = tmp_path / DEFAULT_DECODER_FILENAME
    assert decoder_file.is_file()
    assert result.decoder_path == decoder_file
    # Манифеста в корне нет
    assert not (tmp_path / "дешифратор.manifest.json").exists()
    for document in output_dir.glob("*.docx"):
        assert find_decoder_near_document(document) == decoder_file


def test_folder_pipeline_opens_output_directory(tmp_path, monkeypatch):
    """Проверяет, что при успешном обезличивании папки вызывается открытие созданной папки вывода."""
    import app_paths
    opened = []
    monkeypatch.setattr(app_paths, "open_folder_in_file_manager", lambda p: opened.append(Path(p)) or True)

    doc = Document()
    doc.add_paragraph("Тестовый документ: Иванов Иван Иванович")
    doc.save(str(tmp_path / "test.docx"))

    api = _backend()
    api.open_output_folder = True
    pipeline = FolderAnonymizationPipeline(api, emit_sidecars=False)
    result = pipeline.process(tmp_path)

    expected_output = tmp_path / DEFAULT_OUTPUT_DIR_NAME
    assert result.output_dir == expected_output
    assert expected_output in opened
