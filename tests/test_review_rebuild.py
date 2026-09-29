from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook

from backend_api import BackendApi
from privacy_audit import ReviewFinding, ReviewQueue
from folder_pipeline import FolderAnonymizationPipeline


def _api_with_private(finding, private):
    api = BackendApi.__new__(BackendApi)
    api.review_queue = ReviewQueue([finding])
    api._review_private = {finding.finding_id: private}
    api._review_rebuilt = set()
    return api


def test_rejected_docx_finding_rebuilds_one_paragraph_atomically(tmp_path):
    path = tmp_path / "result.docx"
    doc = Document()
    doc.add_paragraph("Подписант: [ФИО_1]")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Адрес: [АДРЕС_1]"
    doc.save(path)

    finding = ReviewFinding("result.docx", "body", "PER", "[ФИО_1]", 1.0)
    api = _api_with_private(finding, {
        "source_path": str(path), "output_path": str(path),
        "coordinate": {"scope": "body", "kind": "paragraph", "index": 0},
        "placeholder": "[ФИО_1]", "original": "Иванов Иван Иванович", "occurrence": 0,
    })
    api.decide_review_finding(finding.finding_id, "rejected")
    result = api.rebuild_rejected_finding(finding.finding_id)

    assert result["success"] is True
    restored = Document(path)
    assert restored.paragraphs[0].text == "Подписант: Иванов Иван Иванович"
    assert restored.tables[0].cell(0, 0).text == "Адрес: [АДРЕС_1]"
    audit = json.loads((tmp_path / "result.docx.audit.json").read_text(encoding="utf-8"))
    assert "Иванов" not in json.dumps(audit, ensure_ascii=False)


def test_rejected_xlsx_table_cell_rebuilds_only_selected_occurrence(tmp_path):
    path = tmp_path / "result.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Данные"
    sheet["A1"] = "ФИО"
    sheet["B1"] = "[ФИО_1]"
    sheet["A2"] = "Второй"
    sheet["B2"] = "[ФИО_1]"
    workbook.save(path)

    finding = ReviewFinding("result.xlsx", "sheet", "PER", "[ФИО_1]", 1.0)
    api = _api_with_private(finding, {
        "source_path": str(path), "output_path": str(path),
        "coordinate": {"scope": "sheet", "kind": "cell", "sheet": "Данные", "cell": "B2"},
        "placeholder": "[ФИО_1]", "original": "Петров Пётр Петрович", "occurrence": 0,
    })
    api.decide_review_finding(finding.finding_id, "rejected")
    result = api.rebuild_rejected_finding(finding.finding_id)

    assert result["success"] is True
    restored = load_workbook(path, data_only=False).active
    assert restored["B1"].value == "[ФИО_1]"
    assert restored["B2"].value == "Петров Пётр Петрович"


def test_rebuild_is_idempotent_and_does_not_write_when_token_is_missing(tmp_path):
    path = tmp_path / "result.docx"
    doc = Document()
    doc.add_paragraph("Без плейсхолдера")
    doc.save(path)
    before = path.read_bytes()

    finding = ReviewFinding("result.docx", "body", "PER", "[ФИО_1]", 1.0)
    api = _api_with_private(finding, {
        "source_path": str(path), "output_path": str(path),
        "coordinate": {"scope": "body", "kind": "paragraph", "index": 0},
        "placeholder": "[ФИО_1]", "original": "Секрет", "occurrence": 0,
    })
    api.decide_review_finding(finding.finding_id, "rejected")
    result = api.rebuild_rejected_finding(finding.finding_id)

    assert result["success"] is False
    assert path.read_bytes() == before


class _Occurrence:
    def __init__(self, source, start, end):
        self.sofa = source
        self.begin_char = start
        self.end_char = end - 1


class _Referent:
    type_name = "PERSON"

    def __init__(self, source, start, end):
        self.occurrence = [_Occurrence(source, start, end)]


class _Processor:
    def process(self, source, *_args):
        value = "Иванов Иван Иванович"
        start = source.text.find(value)
        return type("Result", (), {
            "entities": [_Referent(source, start, start + len(value))] if start >= 0 else []
        })()


def test_pipeline_records_coordinates_for_docx_and_xlsx_table_rebuild(tmp_path):
    source_docx = tmp_path / "source.docx"
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Иванов Иван Иванович"
    doc.save(source_docx)
    source_xlsx = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active["B2"] = "Иванов Иван Иванович"
    workbook.save(source_xlsx)

    api = BackendApi.__new__(BackendApi)
    api._pullenti_processor = _Processor()
    api.current_placeholders = {"PER": "[ФИО]"}
    api.enabled_placeholders = {"PER"}
    api.qwen_settings = type("Qwen", (), {"enabled": False})()
    api._qwen_postprocessor = type("QwenRuntime", (), {"last_status": "disabled"})()
    api.save_original = True
    api.save_pdf = False
    api.save_decoder = False
    api._window = None

    out = tmp_path / "out"
    for source in (source_docx, source_xlsx):
        api.process_single_file(str(source), set(), set(), output_dir=out,
                                batch_mapping={}, batch_entity_seen={})
    findings = api.get_review_findings()["items"]
    assert any("Таб.0 Р.0 Я.0" in item["location"] for item in findings)
    assert any(item["location"].startswith("Excel:") for item in findings)

    for item in findings:
        api.decide_review_finding(item["finding_id"], "rejected")
    assert Document(out / "source_cleaned.docx").tables[0].cell(0, 0).text == "Иванов Иван Иванович"
    assert load_workbook(out / "source_cleaned.xlsx").active["B2"].value == "Иванов Иван Иванович"
