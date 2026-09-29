import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document

from document_restorer import DocumentRestorer


def _make_docm(path: Path, source: Path) -> None:
    with zipfile.ZipFile(source) as source_zip:
        parts = {name: source_zip.read(name) for name in source_zip.namelist()}
    parts["word/vbaProject.bin"] = b"REAL_VBA_PAYLOAD"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ET.register_namespace("", rel_ns)
    rels = ET.fromstring(parts["word/_rels/document.xml.rels"])
    ET.SubElement(rels, f"{{{rel_ns}}}Relationship", {
        "Id": "rIdVba",
        "Type": "http://schemas.microsoft.com/office/2006/relationships/vbaProject",
        "Target": "vbaProject.bin",
    })
    parts["word/_rels/document.xml.rels"] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
    ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    ET.register_namespace("", ct_ns)
    types = ET.fromstring(parts["[Content_Types].xml"])
    ET.SubElement(types, f"{{{ct_ns}}}Default", {
        "Extension": "bin", "ContentType": "application/vnd.ms-office.vbaProject",
    })
    for node in types:
        if node.attrib.get("PartName") == "/word/document.xml":
            node.set("ContentType", "application/vnd.ms-word.document.macroEnabled.main+xml")
    parts["[Content_Types].xml"] = ET.tostring(types, encoding="utf-8", xml_declaration=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as output:
        for name, data in parts.items():
            output.writestr(name, data)


def test_docm_vba_preserved_for_anonymize_and_restore(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    Document().save(source)
    docm = tmp_path / "source.docm"
    _make_docm(docm, source)

    edited = Document(str(source))
    edited.add_paragraph("[ФИО_1]")
    from document_restorer import save_docm_preserving_vba
    anonymized = tmp_path / "anonymized.docm"
    save_docm_preserving_vba(docm, edited, anonymized)

    restorer = DocumentRestorer({"[ФИО_1]": "Иванов Иван"})
    monkeypatch.setattr(restorer, "_verify_document", lambda _path: None)
    restored = tmp_path / "restored.docm"
    ok, _message = restorer.restore_docx(anonymized, restored)
    assert ok

    for output in (anonymized, restored):
        with zipfile.ZipFile(output) as package:
            assert package.read("word/vbaProject.bin") == b"REAL_VBA_PAYLOAD"
            assert "macroEnabled" in package.read("[Content_Types].xml").decode("utf-8")
            assert "vbaProject" in package.read("word/_rels/document.xml.rels").decode("utf-8")
