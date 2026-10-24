# -*- coding: utf-8 -*-
"""Комплексные тесты анонимизации PPTX, сырых текстовых форматов (.txt, .log, .csv, .md, .json)
и матрицы мультиформатного пересохранения DOCXдодыр v3.0."""

from __future__ import annotations

import json
from pathlib import Path
import re
from types import SimpleNamespace
import zipfile
from xml.etree import ElementTree as ET

import pytest
from docx import Document

from backend_api import BackendApi
from constants import FORMAT_REGISTRY
from document_restorer import (
    DocumentRestorer,
    find_decoder_near_document,
    restore_document,
)
from folder_pipeline import (
    DEFAULT_DECODER_FILENAME,
    DEFAULT_OUTPUT_DIR_NAME,
    FolderAnonymizationPipeline,
    SUPPORTED_SUFFIXES,
)
from markdown_export import save_as_markdown, save_as_text
from presentation_anonymizer import (
    anonymize_pptx,
    extract_pptx_slides,
    pptx_to_markdown,
    pptx_to_text,
    restore_pptx,
)


class _MockPullenti:
    """Детерминированный процессор для обнаружения ПДн в тестах."""

    def __init__(self, entities_map=None):
        self.entities_map = entities_map or {
            "Иванов Иван Иванович": "PERSON",
            "Иванова Ивана Ивановича": "PERSON",
            "Петров Петр Петрович": "PERSON",
            "Сидоров Сидор Сидорович": "PERSON",
            "ООО «Вектор»": "ORGANIZATION",
            "ООО «Ромашка»": "ORGANIZATION",
            "ООО Ромашка": "ORGANIZATION",
            "40702810302410001062": "RU_ACCOUNT",
            "7701234567": "INN",
        }

    def process(self, source, *_args):
        text = getattr(source, "text", "") or str(source)
        entities = []
        for name, type_name in self.entities_map.items():
            start = 0
            while True:
                idx = text.find(name, start)
                if idx < 0:
                    break
                end = idx + len(name)
                occ = SimpleNamespace(sofa=source, begin_char=idx, end_char=end - 1)
                referent = SimpleNamespace(type_name=type_name, occurrence=[occ])
                entities.append(referent)
                start = end
        return SimpleNamespace(entities=entities)


def _build_test_backend(monkeypatch, custom_entities=None):
    """Инициализирует тестовый экземпляр BackendApi с синтетическим детектором."""
    api = BackendApi.__new__(BackendApi)
    api._pullenti_processor = _MockPullenti(custom_entities)
    api.current_placeholders = {
        "PER": "[ФИО]",
        "ORG": "[Организация]",
        "RU_ACCOUNT": "[Р/с]",
        "INN": "[ИНН]",
        "PHONE": "[Телефон]",
        "EMAIL": "[Email]",
        "PASSPORT": "[Паспорт]",
    }
    api.qwen_settings = SimpleNamespace(enabled=False)
    api._qwen_postprocessor = SimpleNamespace(last_status="disabled")
    api.user_exclusions = set()
    api.custom_replacements = set()
    api.save_original = False
    api.save_docx = False
    api.save_pdf = False
    api.save_markdown = False
    api.save_text = False
    api.save_decoder = True
    api.ocr_pdf = False
    api._window = None
    api.emit_audit_sidecars = False
    api.save_log_files = False
    api.save_audit_files = False
    api.open_output_folder = False
    api._last_reconcile_temp = None
    api._last_generated_outputs = []
    api.anonymize_cache = {}
    return api


def _create_sample_pptx(path: Path, title: str = "Презентация", text: str = "Иванов Иван Иванович") -> Path:
    """Создает валидный архив PPTX с текстом слайда, заметками докладчика и комментариями."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>
  <Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>
  <Override PartName="/ppt/notesSlides/notesSlide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.notesSlide+xml"/>
  <Override PartName="/ppt/comments/comment1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.comments+xml"/>
</Types>"""

    root_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>
</Relationships>"""

    presentation_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:presentation xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <p:sldIdLst>
    <p:sldId id="256" r:id="rId1"/>
  </p:sldIdLst>
</p:presentation>"""

    pres_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/>
</Relationships>"""

    slide1_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld>
    <p:spTree>
      <p:sp>
        <p:txBody>
          <a:p>
            <a:r>
              <a:t>{title}</a:t>
            </a:r>
          </a:p>
          <a:p>
            <a:r>
              <a:t>Докладчик: {text}</a:t>
            </a:r>
          </a:p>
        </p:txBody>
      </p:sp>
    </p:spTree>
  </p:cSld>
</p:sld>"""

    notes_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:notes xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld>
    <p:spTree>
      <p:sp>
        <p:txBody>
          <a:p>
            <a:r>
              <a:t>Заметка к слайду: согласовано с {text}</a:t>
            </a:r>
          </a:p>
        </p:txBody>
      </p:sp>
    </p:spTree>
  </p:cSld>
</p:notes>"""

    comment_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:cmLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cm authorId="0" dt="2026-10-23T10:00:00.000">
    <p:text>Обратить внимание: {text}</p:text>
  </p:cm>
</p:cmLst>"""

    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", root_rels)
        zf.writestr("ppt/presentation.xml", presentation_xml)
        zf.writestr("ppt/_rels/presentation.xml.rels", pres_rels)
        zf.writestr("ppt/slides/slide1.xml", slide1_xml)
        zf.writestr("ppt/notesSlides/notesSlide1.xml", notes_xml)
        zf.writestr("ppt/comments/comment1.xml", comment_xml)

    return path


# =========================================================================
# ЦИКЛ 1: Движок презентаций (presentation_anonymizer)
# =========================================================================

def test_cycle1_pptx_extraction_and_conversions(tmp_path):
    """Проверка извлечения слайдов, заметок и конвертации презентации в текст и Markdown."""
    pptx_path = tmp_path / "presentation.pptx"
    _create_sample_pptx(pptx_path, title="Годовой отчет", text="Иванов Иван Иванович")

    slides = extract_pptx_slides(pptx_path)
    assert len(slides) == 1
    assert slides[0]["slide_number"] == 1
    assert slides[0]["title"] == "Годовой отчет"
    assert any("Иванов Иван Иванович" in p for p in slides[0]["paragraphs"])
    assert any("Иванов Иван Иванович" in n for n in slides[0]["notes"])

    # Проверка экспорта в текст
    txt_content = pptx_to_text(slides)
    assert "=== Слайд 1: Годовой отчет ===" in txt_content
    assert "Докладчик: Иванов Иван Иванович" in txt_content
    assert "[Заметки:]" in txt_content

    # Проверка экспорта в Markdown
    md_content = pptx_to_markdown(slides)
    assert "## Слайд 1: Годовой отчет" in md_content
    assert "Докладчик: Иванов Иван Иванович" in md_content
    assert "> **Заметки докладчика:**" in md_content


def test_cycle1_pptx_anonymization_and_direct_restoration(tmp_path):
    """Проверка прямой анонимизации PPTX и восстановления через restore_pptx."""
    src_pptx = tmp_path / "source.pptx"
    out_pptx = tmp_path / "source_cleaned.pptx"
    restored_pptx = tmp_path / "source_restored.pptx"

    _create_sample_pptx(src_pptx, title="План 2026", text="Иванов Иван Иванович")

    def mock_anon(text):
        if "Иванов Иван Иванович" in text:
            return text.replace("Иванов Иван Иванович", "[ФИО_1]"), 1, ["Заменено ФИО"]
        return text, 0, []

    repls, logs, full_text = anonymize_pptx(src_pptx, out_pptx, mock_anon)
    assert repls >= 3  # слайд, заметки, комментарий
    assert out_pptx.is_file()

    # Проверяем, что исходный ПДн исчез из обезличенного PPTX
    cleaned_slides = extract_pptx_slides(out_pptx)
    assert "[ФИО_1]" in cleaned_slides[0]["paragraphs"][1]
    assert "Иванов Иван Иванович" not in cleaned_slides[0]["paragraphs"][1]

    # Восстанавливаем через restore_pptx
    mapping = {"[ФИО_1]": "Иванов Иван Иванович"}
    ok, msg = restore_pptx(out_pptx, restored_pptx, mapping)
    assert ok is True
    assert restored_pptx.is_file()

    restored_slides = extract_pptx_slides(restored_pptx)
    assert "Иванов Иван Иванович" in restored_slides[0]["paragraphs"][1]
    assert "[ФИО_1]" not in restored_slides[0]["paragraphs"][1]


# =========================================================================
# ЦИКЛ 2: Сырые текстовые форматы (.txt, .log, .csv, .md, .json)
# =========================================================================

@pytest.mark.parametrize("ext,content,expected_token", [
    (".txt", "Сотрудник: Иванов Иван Иванович, тел: +7 (495) 123-45-67", "[ФИО"),
    (".log", "2026-10-23 10:15:00 [INFO] User login: Иванов Иван Иванович from IP 127.0.0.1", "[ФИО"),
    (".csv", "id,name,role\n1,Иванов Иван Иванович,Директор\n2,Петров Петр Петрович,Бухгалтер", "[ФИО"),
    (".md", "# Заголовок\n\nАвтор: **Иванов Иван Иванович**\n\n- Организация: ООО «Ромашка»", "[ФИО"),
    (".json", json.dumps({"user": "Иванов Иван Иванович", "org": "ООО «Ромашка»"}, ensure_ascii=False), "[ФИО"),
])
def test_cycle2_raw_text_formats_anonymization(tmp_path, monkeypatch, ext, content, expected_token):
    """Проверка анонимизации всех текстовых форматов через BackendApi."""
    api = _build_test_backend(monkeypatch)
    src_file = tmp_path / f"test_data{ext}"
    src_file.write_text(content, encoding="utf-8")

    out_dir = tmp_path / f"out_{ext.strip('.')}"
    out_dir.mkdir()

    res = api.process_single_file(str(src_file), set(), set(), output_dir=out_dir)
    assert res > 0

    # Ищем обезличенный файл (исключая файл дешифратора)
    cleaned_files = [f for f in out_dir.glob(f"*{ext}") if "дешифратор" not in f.name.lower()]
    assert len(cleaned_files) == 1
    cleaned = cleaned_files[0]
    cleaned_text = cleaned.read_text(encoding="utf-8")

    assert "Иванов Иван Иванович" not in cleaned_text
    assert expected_token in cleaned_text


def test_cycle2_text_cp1251_encoding_support(tmp_path, monkeypatch):
    """Проверка обработки файлов в кодировке Windows-1251."""
    api = _build_test_backend(monkeypatch)
    src_file = tmp_path / "legacy_win1251.txt"
    src_file.write_bytes("Клиент: Иванов Иван Иванович".encode("cp1251"))

    out_dir = tmp_path / "out_cp1251"
    out_dir.mkdir()

    res = api.process_single_file(str(src_file), set(), set(), output_dir=out_dir)
    assert res > 0

    cleaned = list(out_dir.glob("*.txt"))[0]
    cleaned_content = cleaned.read_text(encoding="utf-8")
    assert "Иванов Иван Иванович" not in cleaned_content
    assert "[ФИО" in cleaned_content


# =========================================================================
# ЦИКЛ 3: Мультиформатное пересохранение (Матрица форматов в .txt и .md)
# =========================================================================

def test_cycle3_docx_to_only_txt(tmp_path, monkeypatch):
    """Пользователь выбрал DOCX и включил только сохранение в .txt (save_text=True, save_docx=False)."""
    api = _build_test_backend(monkeypatch)
    api.save_text = True
    api.save_docx = False
    api.save_pdf = False
    api.save_markdown = False

    doc = Document()
    doc.add_heading("Трудовой договор", level=1)
    doc.add_paragraph("Работник: Иванов Иван Иванович")
    doc.add_paragraph("Организация: ООО «Ромашка»")
    src = tmp_path / "contract.docx"
    doc.save(str(src))

    out_dir = tmp_path / "out_docx_to_txt"
    out_dir.mkdir()

    res = api.process_single_file(str(src), set(), set(), output_dir=out_dir)
    assert res > 0

    # Проверяем, что создан ТОЛЬКО .txt, а .docx НЕ создан
    txt_files = list(out_dir.glob("*.txt"))
    output_txts = [f for f in txt_files if not f.name.endswith("_log.txt") and not f.name.endswith(".manifest.json")]
    docx_files = list(out_dir.glob("*.docx"))

    assert len(docx_files) == 0, "Файл DOCX не должен был создаваться"
    assert len(output_txts) == 1, "Должен быть создан ровно один файл .txt"

    txt_body = output_txts[0].read_text(encoding="utf-8")
    assert "Иванов Иван Иванович" not in txt_body
    assert "[ФИО" in txt_body
    assert "Трудовой договор" in txt_body


def test_cycle3_pptx_to_only_markdown(tmp_path, monkeypatch):
    """Пользователь выбрал PPTX и включил только сохранение в .md (save_markdown=True)."""
    api = _build_test_backend(monkeypatch)
    api.save_markdown = True
    api.save_text = False
    api.save_docx = False
    api.save_pdf = False

    src = tmp_path / "slides.pptx"
    _create_sample_pptx(src, title="Стратегия развития", text="Иванов Иван Иванович")

    out_dir = tmp_path / "out_pptx_to_md"
    out_dir.mkdir()

    res = api.process_single_file(str(src), set(), set(), output_dir=out_dir)
    assert res > 0

    md_files = list(out_dir.glob("*.md"))
    pptx_files = list(out_dir.glob("*.pptx"))

    assert len(pptx_files) == 0, "Оригинальный формат PPTX не должен был сохраняться при явном выборе другого формата"
    assert len(md_files) == 1

    md_body = md_files[0].read_text(encoding="utf-8")
    assert "## Слайд 1: Стратегия развития" in md_body
    assert "Иванов Иван Иванович" not in md_body
    assert "[ФИО" in md_body


def test_cycle3_pptx_multi_export_matrix(tmp_path, monkeypatch):
    """Пользователь выбрал сохранение и в .txt, и в .md одновременно."""
    api = _build_test_backend(monkeypatch)
    api.save_text = True
    api.save_markdown = True
    api.save_docx = False
    api.save_pdf = False

    src = tmp_path / "multi.pptx"
    _create_sample_pptx(src, title="Маркетинг", text="Иванов Иван Иванович")

    out_dir = tmp_path / "out_multi"
    out_dir.mkdir()

    res = api.process_single_file(str(src), set(), set(), output_dir=out_dir)
    assert res > 0

    md_files = list(out_dir.glob("*.md"))
    txt_files = [f for f in out_dir.glob("*.txt") if not f.name.endswith("_log.txt")]

    assert len(md_files) == 1
    assert len(txt_files) == 1

    assert "[ФИО" in md_files[0].read_text(encoding="utf-8")
    assert "[ФИО" in txt_files[0].read_text(encoding="utf-8")


# =========================================================================
# ЦИКЛ 4: Пакетный пайплайн папок и криптографическое восстановление
# =========================================================================

def test_cycle4_folder_pipeline_mixed_batch_to_txt(tmp_path, monkeypatch):
    """Пакетная обработка смешанной папки (DOCX, PPTX, LOG, CSV) с пересохранением всех в .txt."""
    api = _build_test_backend(monkeypatch)
    api.save_text = True
    api.save_docx = False
    api.save_pdf = False
    api.save_markdown = False
    api.save_decoder = True

    in_dir = tmp_path / "input_folder"
    in_dir.mkdir()

    # 1. DOCX
    doc = Document()
    doc.add_paragraph("Сотрудник: Иванов Иван Иванович")
    doc.save(str(in_dir / "file1.docx"))

    # 2. PPTX
    _create_sample_pptx(in_dir / "file2.pptx", title="Слайды", text="Петров Петр Петрович")

    # 3. LOG
    (in_dir / "file3.log").write_text("User Сидоров Сидор Сидорович connected", encoding="utf-8")

    # 4. CSV
    (in_dir / "file4.csv").write_text("id,name\n10,Иванов Иван Иванович", encoding="utf-8")

    pipeline = FolderAnonymizationPipeline(backend=api)
    result = pipeline.process(in_dir)

    assert result.error_count == 0
    assert result.processed_count == 4

    out_dir = result.output_dir
    assert out_dir.is_dir()

    # Все документы в результирующей папке должны быть .txt
    result_txts = [
        p for p in out_dir.glob("*.txt")
        if not p.name.endswith("_log.txt") and not p.name.startswith(".")
    ]
    assert len(result_txts) == 4

    # Проверяем дешифратор
    assert result.decoder_path is not None
    assert result.decoder_path.is_file()

    # Проверяем восстановление каждого из сгенерированных .txt файлов
    for txt_doc in result_txts:
        restored = restore_document(txt_doc)
        assert restored.is_file()
        restored_text = restored.read_text(encoding="utf-8")
        assert any(name in restored_text for name in ["Иванов Иван Иванович", "Петров Петр Петрович", "Сидоров Сидор Сидорович"])


def test_cycle4_restore_pptx_with_bound_decoder(tmp_path, monkeypatch):
    """Восстановление обезличенной презентации PPTX по криптографически привязанному дешифратору."""
    api = _build_test_backend(monkeypatch)
    # Сохраняем в родном формате PPTX с дешифратором
    api.save_text = False
    api.save_markdown = False
    api.save_docx = False
    api.save_pdf = False
    api.save_decoder = True

    src = tmp_path / "report.pptx"
    _create_sample_pptx(src, title="Финансы", text="Иванов Иван Иванович")

    out_dir = tmp_path / "out_pptx_restore"
    out_dir.mkdir()

    res = api.process_single_file(str(src), set(), set(), output_dir=out_dir)
    assert res > 0

    cleaned_pptx = list(out_dir.glob("report_cleaned*.pptx"))[0]
    assert cleaned_pptx.is_file()

    # Выполняем восстановление
    restored_pptx = restore_document(cleaned_pptx)
    assert restored_pptx.is_file()

    slides = extract_pptx_slides(restored_pptx)
    assert "Иванов Иван Иванович" in slides[0]["paragraphs"][1]


# =========================================================================
# ЦИКЛ 5: UI, Pywebview Bridge, Диалоги и Симуляция кликов пользователя
# =========================================================================

def test_cycle5_ui_dom_and_bindings():
    """Проверка наличия тумблера save_text в index.html и привязки в script.js."""
    html_path = Path(__file__).resolve().parent.parent / "web" / "index.html"
    js_path = Path(__file__).resolve().parent.parent / "web" / "script.js"

    html = html_path.read_text(encoding="utf-8")
    js = js_path.read_text(encoding="utf-8")

    # Чекбокс должен существовать в HTML рядом с Word, PDF, MD
    assert 'id="cb-save-txt"' in html
    assert "Текстовый документ" in html

    # В JS должно быть чтение и сохранение настроек cb-save-txt
    assert "cb-save-txt" in js
    assert "save_text" in js


def test_cycle5_backend_settings_contract():
    """Проверка контракта get_settings / update_settings для save_text."""
    api = BackendApi()
    settings = api.get_settings()
    assert "save_text" in settings
    assert isinstance(settings["save_text"], bool)

    # Имитация изменения настройки пользователем через интерфейс
    api.update_settings({"save_text": True})
    assert api.save_text is True
    assert api.get_settings()["save_text"] is True

    api.update_settings({"save_text": False})
    assert api.save_text is False


def test_cycle5_format_registry_and_dialog_filters():
    """Проверка реестра поддерживаемых расширений и фильтров диалога файлов."""
    # Реестр форматов
    anon_extensions = FORMAT_REGISTRY["anonymize"]["extensions"]
    restore_extensions = FORMAT_REGISTRY["restore"]["extensions"]

    for ext in [".pptx", ".pptm", ".txt", ".log", ".csv", ".md", ".json"]:
        assert ext in anon_extensions, f"Расширение {ext} отсутствует в FORMAT_REGISTRY['anonymize']['extensions']"
        assert ext in SUPPORTED_SUFFIXES, f"Расширение {ext} отсутствует в SUPPORTED_SUFFIXES"
        assert ext in restore_extensions, f"Расширение {ext} отсутствует в FORMAT_REGISTRY['restore']['extensions']"

    # Проверка фильтров в main.py
    main_py_path = Path(__file__).resolve().parent.parent / "main.py"
    main_code = main_py_path.read_text(encoding="utf-8")

    assert "*.pptx" in main_code
    assert "*.txt" in main_code
    assert "*.log" in main_code
    assert "*.csv" in main_code
    assert "*.json" in main_code
