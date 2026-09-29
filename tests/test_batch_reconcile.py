from docx import Document
import pytest
import zipfile

from batch_reconcile import reconcile_docx, reconcile_text
from hidden_data import HiddenDataError


@pytest.mark.parametrize("unsafe_name", ["../outside.xml", "/absolute.xml"])
def test_reconcile_docx_rejects_unsafe_archive_members_without_republishing(tmp_path, unsafe_name):
    source = tmp_path / "unsafe.docx"
    with zipfile.ZipFile(source, "w") as package:
        package.writestr("word/document.xml", "<root/>")
        package.writestr(unsafe_name, "<private/>")
    original_bytes = source.read_bytes()

    with pytest.raises(HiddenDataError):
        reconcile_docx(source, {"[ИНН_1]": "770123456789"})

    assert source.read_bytes() == original_bytes
    assert not list(tmp_path.glob(".unsafe.*.docx"))


def test_reconcile_text_uses_complete_decoder_and_protects_existing_tokens():
    mapping = {
        "[Телефон_1]": {"type": "PHONE_NUMBER", "original": "+7 999 123-45-67"},
        "[ИНН_1]": "770123456789",
    }
    source = "Телефон [Телефон_1], ИНН 770123456789; +7 999 123-45-67."

    result, count = reconcile_text(source, mapping)

    assert result == "Телефон [Телефон_1], ИНН [ИНН_1]; [Телефон_1]."
    assert count == 2


def test_reconcile_text_does_not_treat_arbitrary_brackets_as_placeholders():
    mapping = {"[ИНН_1]": "7701234567"}
    source = "Реквизиты (ИНН 7701234567) и маршрут /ИНН 7701234567/"

    result, count = reconcile_text(source, mapping)

    assert result == "Реквизиты (ИНН [ИНН_1]) и маршрут /ИНН [ИНН_1]/"
    assert count == 2


def test_reconcile_text_rejects_nonlexical_and_placeholder_tail_values():
    mapping = {
        "[ФИО_1]": {"original": ")"},
        "[Должность_1]": {"original": "_1]"},
    }

    result, count = reconcile_text("Текст (значение) и [ФИО_1]", mapping)

    assert result == "Текст (значение) и [ФИО_1]"
    assert count == 0


def test_reconcile_text_uses_lexical_boundaries():
    mapping = {"[Телефон_1]": {"type": "PHONE_NUMBER", "original": "12345"}}

    result, count = reconcile_text("Номер 12345, но 123456 — другой", mapping)

    assert result == "Номер [Телефон_1], но 123456 — другой"
    assert count == 1


def test_reconcile_text_does_not_globalize_contextual_entities():
    mapping = {
        "[Адрес_1]": {"type": "ADDRESS", "original": "города Москвы"},
        "[Наименование_1]": {"type": "ORG", "original": "Роскомнадзор"},
        "[ФИО_1]": {"type": "PER", "original": "Света"},
    }
    source = "команды города Москвы; Роскомнадзор; света достаточно"

    result, count = reconcile_text(source, mapping)

    assert result == source
    assert count == 0


def test_reconcile_docx_replaces_value_split_across_runs_without_flattening(tmp_path):
    path = tmp_path / "early-output.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    first = paragraph.add_run("name@")
    first.bold = True
    second = paragraph.add_run("example.com")
    second.italic = True
    doc.save(path)

    count = reconcile_docx(
        path,
        {"[Email_7]": {"type": "EMAIL", "original": "name@example.com"}},
    )

    restored = Document(path)
    assert restored.paragraphs[0].text == "[Email_7]"
    assert restored.paragraphs[0].runs[0].bold is True
    assert len(restored.paragraphs[0].runs) == 2
    assert count == 1
