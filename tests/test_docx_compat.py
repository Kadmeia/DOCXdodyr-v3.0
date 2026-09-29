# -*- coding: utf-8 -*-
"""Тесты модуля docx_compat: устойчивость к отсутствию промежуточных директорий в PyInstaller."""

import os
import pytest
from pathlib import Path


def test_docx_compat_template_loading():
    """Проверяет, что патчи docx_compat позволяют загружать шаблоны даже при битом __file__."""
    import docx_compat
    import docx.parts.hdrftr as hdrftr
    import docx.parts.styles as styles
    import docx.parts.settings as settings
    import docx.parts.comments as comments
    import docx

    # Проверяем штатную загрузку
    hdr_xml = hdrftr.HeaderPart._default_header_xml()
    assert len(hdr_xml) > 0
    assert b"<w:hdr" in hdr_xml

    ftr_xml = hdrftr.FooterPart._default_footer_xml()
    assert len(ftr_xml) > 0
    assert b"<w:ftr" in ftr_xml

    styles_xml = styles.StylesPart._default_styles_xml()
    assert len(styles_xml) > 0
    assert b"<w:styles" in styles_xml

    settings_xml = settings.SettingsPart._default_settings_xml()
    assert len(settings_xml) > 0
    assert b"<w:settings" in settings_xml

    comments_xml = comments.CommentsPart._default_comments_xml()
    assert len(comments_xml) > 0
    assert b"<w:comments" in comments_xml


def test_docx_compat_simulated_broken_posix_traversal():
    """Симулирует ситуацию PyInstaller onedir, когда директория docx/parts отсутствует."""
    import docx_compat
    import docx.parts.hdrftr as hdrftr

    # Подменяем __file__ на заведомо несуществующий промежуточный путь
    saved_file = hdrftr.__file__
    try:
        hdrftr.__file__ = "/tmp/fake_nonexistent_dir_12345/parts/hdrftr.py"
        # Благодаря patch_docx_templates, функция не упадет с Errno 2
        hdr_xml = hdrftr.HeaderPart._default_header_xml()
        assert len(hdr_xml) > 0
        assert b"<w:hdr" in hdr_xml
    finally:
        hdrftr.__file__ = saved_file


def test_docx_compat_create_empty_document():
    """Проверяет, что создание пустого Document() использует корректный шаблон."""
    import docx_compat
    from docx import Document

    doc = Document()
    assert len(doc.paragraphs) == 0
    assert len(doc.sections) == 1
