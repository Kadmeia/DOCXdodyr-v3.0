from types import SimpleNamespace

import pytest
from docx import Document

from backend_api import BackendApi
from qwen_postprocessor import QwenPlaceholderPostprocessor, protect_placeholders, restore_and_validate, PlaceholderValidationError


def test_qwen_actual_writeback_all_structures_preserves_runs():
    document = Document()
    body = document.add_paragraph()
    body.add_run("Hello ").bold = True
    body.add_run("[NAME_").italic = True
    body.add_run("1] ,test 15.05.2025 must remain.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Hello [NAME_1] ,test 15.05.2025 must remain."
    table.cell(0, 0).add_paragraph("Hello [NAME_2] ,test 15.05.2025 must remain.")
    document.add_paragraph("Hello [NAME_2] ,test 15.05.2025 must remain.")
    document.sections[0].header.paragraphs[0].text = table.cell(0, 0).paragraphs[0].text
    document.sections[0].footer.paragraphs[0].text = table.cell(0, 0).paragraphs[0].text
    api = BackendApi.__new__(BackendApi)
    api.qwen_settings = SimpleNamespace(enabled=True)
    class Model:
        def generate(self, prompt, **kwargs):
            return prompt.split("ДОКУМЕНТ:\n", 1)[1].replace(" ,test", ", test")
    api._qwen_postprocessor = QwenPlaceholderPostprocessor({"enabled": True}, model_loader=lambda _: Model())
    api._write_back_qwen_document(document, [])
    paragraphs = [*document.paragraphs, *table.cell(0, 0).paragraphs,
                  document.sections[0].header.paragraphs[0], document.sections[0].footer.paragraphs[0]]
    for paragraph in paragraphs:
        assert ", test 15.05.2025 must remain." in paragraph.text
        assert " ,test" not in paragraph.text
        assert "[NAME_" in paragraph.text
    assert body.runs[0].bold and body.runs[1].italic


@pytest.mark.parametrize("replacement", ["may", "must not", "cannot"])
def test_modality_changes_are_rejected(replacement):
    protected = protect_placeholders("[NAME_1] must pay 123 on 15.05.2025.")
    with pytest.raises(PlaceholderValidationError):
        restore_and_validate(protected.text.replace("must", replacement), protected)
