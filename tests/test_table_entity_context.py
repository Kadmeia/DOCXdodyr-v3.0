# -*- coding: utf-8 -*-
from types import SimpleNamespace

from docx import Document

from backend_api import BackendApi
from table_entity_context import (
    build_table_contexts,
    contexts_from_extracted_table,
    project_span_to_value,
)


def test_horizontal_table_maps_label_to_right_value_cell():
    contexts = build_table_contexts([["ФИО", "Иванов Иван Иванович"]])
    value = next(item for item in contexts if item.column == 1)

    assert value.label == "ФИО"
    assert value.analysis_text.endswith("Иванов Иван Иванович")
    assert value.value_text not in value.analysis_text[:value.value_offset]


def test_vertical_table_maps_label_to_cell_below():
    contexts = build_table_contexts([["Кадастровый номер"], ["77:01:0000000:1"]])
    value = next(item for item in contexts if item.row == 1)

    assert value.label == "Кадастровый номер"
    assert value.label_row == 0
    assert value.label_column == 0


def test_multiline_value_keeps_cell_coordinates_and_rejects_prefix_spans():
    contexts = build_table_contexts([["Адрес"], ["г. Москва\nул. Тверская, д. 1"]])
    value = next(item for item in contexts if item.row == 1)

    assert project_span_to_value(0, 4, value) is None
    start = value.value_offset + value.value_text.index("Москва")
    end = start + len("Москва")
    assert project_span_to_value(start, end, value) == (value.value_text.index("Москва"), value.value_text.index("Москва") + len("Москва"))


def test_pdf_extracted_rows_use_the_same_context_builder():
    contexts = contexts_from_extracted_table((("ИНН", "7701234567"),))
    value = next(item for item in contexts if item.column == 1)
    assert value.label == "ИНН"


def test_real_legal_cartridge_uses_neighbour_label_but_returns_value_only():
    from legal_pullenti import initialize_ner, iter_pullenti_spans
    from pullenti.ner.ProcessorService import ProcessorService
    from pullenti.ner.SourceOfAnalysis import SourceOfAnalysis

    context = next(
        item for item in build_table_contexts([["Кадастровый номер", "77:01:0004012:123"]])
        if item.column == 1
    )
    initialize_ner()
    result = ProcessorService.create_processor().process(
        SourceOfAnalysis(context.analysis_text), None, None
    )
    span = next(item for item in iter_pullenti_spans(result, context.analysis_text) if item.label == "CADASTRAL_NUMBER")

    assert project_span_to_value(span.start, span.end, context) == (0, len(context.value_text))
    assert span.text == context.value_text


class _Occurrence:
    def __init__(self, text, start, end):
        self.begin_char = start
        self.end_char = end

    def get_text(self):
        return self.sofa.text[self.begin_char:self.end_char + 1]


class _Referent:
    type_name = "PERSON"

    def __init__(self, occurrence):
        self.occurrence = [occurrence]


class _TableProcessor:
    def process(self, source, *_):
        text = source.text
        name = "Иванов Иван Иванович"
        if name not in text:
            return SimpleNamespace(entities=[])
        start = text.index(name)
        occurrence = _Occurrence(text, start, start + len(name) - 1)
        occurrence.sofa = source
        return SimpleNamespace(entities=[_Referent(occurrence)])


def _api_for_table_test():
    api = BackendApi.__new__(BackendApi)
    api._pullenti_processor = _TableProcessor()
    api.current_placeholders = {"PER": "[ФИО]"}
    return api


def test_docx_table_replaces_only_value_cell_and_not_semantic_label():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "ФИО"
    table.cell(0, 1).text = "Иванов Иван Иванович"
    table.cell(1, 0).text = "Примечание"
    table.cell(1, 1).text = "Иванов Иван Иванович\nподписал договор"

    api = _api_for_table_test()
    count, _logs = api.clean_document(doc, set(), set())

    assert count == 2
    assert table.cell(0, 0).text == "ФИО"
    assert table.cell(0, 1).text == "[ФИО]"
    assert table.cell(1, 1).text == "[ФИО]\nподписал договор"
