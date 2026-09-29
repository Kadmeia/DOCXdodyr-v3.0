# -*- coding: utf-8 -*-
"""Тесты CLI контекстного меню Проводника Windows и автономного режима (--anonymize)."""

import os
import subprocess
import sys
import tempfile
import json
import types
from pathlib import Path
import docx
import pytest

import windows_context_menu


class _FakeRegistryKey:
    def __init__(self, registry, path):
        self.registry = registry
        self.path = path

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        return False


class _FakeWinreg:
    HKEY_CURRENT_USER = object()
    REG_SZ = 1
    KEY_ALL_ACCESS = 2
    KEY_READ = 4

    def __init__(self):
        self.values = {}

    def CreateKey(self, _root, path):
        self.values.setdefault(path, {})
        return _FakeRegistryKey(self.values, path)

    def OpenKey(self, _root, path, *_args):
        if path not in self.values:
            raise FileNotFoundError(path)
        return _FakeRegistryKey(self.values, path)

    def SetValueEx(self, key, name, _reserved, _kind, value):
        self.values[key.path][name] = value

    def QueryValueEx(self, key, name):
        return self.values[key.path][name], self.REG_SZ

    def EnumKey(self, _key, _index):
        raise OSError("no subkeys")

    def DeleteKey(self, _root, path):
        self.values.pop(path, None)


def _install_webview_stub(monkeypatch):
    """Keep CLI-only tests independent from the optional desktop WebView runtime."""
    stub = types.ModuleType("webview")
    stub.OPEN_DIALOG = 1
    stub.FOLDER_DIALOG = 2
    stub.create_window = lambda *args, **kwargs: None
    stub.start = lambda *args, **kwargs: None
    monkeypatch.setitem(sys.modules, "webview", stub)


def _capture_cli_json(monkeypatch, main_module):
    """Capture the structured CLI payload without mixing in logging output."""
    printed = []
    monkeypatch.setattr(main_module, "_safe_cli_print", printed.append)
    return printed


def test_windows_context_menu_registration_roundtrip():
    """Проверяет регистрацию, проверку статуса и отмену регистрации контекстного меню (Обезличить и Восстановить)."""
    if sys.platform != "win32":
        pytest.skip("Контекстное меню Windows поддерживается только на win32")

    # 1. Регистрация (3 для обезличивания + 3 для восстановления = 6)
    reg_result = windows_context_menu.register_windows_context_menu()
    assert reg_result["status"] == "ok"
    assert len(reg_result["registered"]) == 6
    assert windows_context_menu.is_windows_context_menu_registered() is True

    # 2. Повторный вызов безопасен (идемпотентность)
    reg_result2 = windows_context_menu.register_windows_context_menu()
    assert reg_result2["status"] == "ok"

    # 3. Отмена регистрации
    unreg_result = windows_context_menu.unregister_windows_context_menu()
    assert unreg_result["status"] == "ok"
    assert len(unreg_result["removed"]) == 6
    assert windows_context_menu.is_windows_context_menu_registered() is False

    # 4. Восстановление регистрации для рабочей системы
    restore_result = windows_context_menu.register_windows_context_menu()
    assert restore_result["status"] == "ok"
    assert windows_context_menu.is_windows_context_menu_registered() is True


def test_context_menu_commands_are_headless_and_status_checks_all_entries(monkeypatch):
    """All six Explorer commands run without GUI and status detects partial installs."""
    fake_winreg = _FakeWinreg()
    monkeypatch.setattr(windows_context_menu, "_is_windows", lambda: True)
    monkeypatch.setitem(sys.modules, "winreg", fake_winreg)

    result = windows_context_menu.register_windows_context_menu(command_prefix='"C:\\DOCXdodyr.exe"')
    assert result["status"] == "ok"
    command_values = {
        key: fake_winreg.values[f"{key}\\command"][""]
        for key in result["registered"]
    }
    assert len(command_values) == 6
    assert all("--headless" in value and "--no-open-output" in value for value in command_values.values())
    assert '"%1"' in command_values[windows_context_menu.REG_FILE_ANON_KEY]
    assert '"%V"' in command_values[windows_context_menu.REG_FOLDER_BG_RESTORE_KEY]
    assert windows_context_menu.is_windows_context_menu_registered() is True

    fake_winreg.values.pop(f"{windows_context_menu.REG_FILE_RESTORE_KEY}\\command")
    assert windows_context_menu.is_windows_context_menu_registered() is False

    # A legacy GUI command must not be reported as a current registration.
    fake_winreg.values[f"{windows_context_menu.REG_FILE_RESTORE_KEY}\\command"] = {"": '"C:\\DOCXdodyr.exe" --restore "%1"'}
    assert windows_context_menu.is_windows_context_menu_registered() is False


def test_inno_context_menu_commands_match_headless_contract():
    """The installer must preserve the same background CLI contract as runtime registration."""
    installer = Path(__file__).resolve().parents[1] / "installer" / "DOCXdodyr.iss"
    command_lines = [line for line in installer.read_text(encoding="utf-8").splitlines()
                     if "\\command\"; ValueType: string" in line and "DOCXdodyr_" in line]
    assert len(command_lines) == 6
    assert all("--headless" in line and "--no-open-output" in line for line in command_lines)


def test_headless_restore_file_uses_module_restore_function(monkeypatch, tmp_path, capsys):
    """The CLI restore path uses the public module function and returns JSON."""
    _install_webview_stub(monkeypatch)
    import main
    import document_restorer
    printed = _capture_cli_json(monkeypatch, main)

    document = tmp_path / "sample.docx"
    decoder = tmp_path / "sample_дешифратор.json"
    document.write_bytes(b"placeholder")
    decoder.write_text("{}", encoding="utf-8")
    calls = []
    monkeypatch.setattr(document_restorer, "find_decoder_near_document", lambda _path: decoder)
    monkeypatch.setattr(document_restorer, "restore_document",
                        lambda path, decoder_path=None: calls.append((Path(path), Path(decoder_path))) or path.with_name("restored.docx"))
    monkeypatch.setattr(sys, "argv", ["main.py", "--restore", str(document), "--headless"])

    assert main.handle_cli_arguments() is True
    payload = json.loads(printed[-1])
    assert payload["status"] == "ok"
    assert calls == [(document.resolve(), decoder)]
    assert payload["items"][0]["restored"][0]["output"].endswith("restored.docx")


def test_headless_restore_folder_resolves_decoder_per_document(monkeypatch, tmp_path, capsys):
    """A folder restore does not try to resolve a decoder from the directory itself."""
    _install_webview_stub(monkeypatch)
    import main
    import document_restorer
    printed = _capture_cli_json(monkeypatch, main)

    folder = tmp_path / "docs"
    folder.mkdir()
    documents = [folder / "one.docx", folder / "two.xlsx"]
    for document in documents:
        document.write_bytes(b"placeholder")
    decoders = {document: folder / f"{document.stem}_дешифратор.json" for document in documents}
    for decoder in decoders.values():
        decoder.write_text("{}", encoding="utf-8")
    calls = []
    monkeypatch.setattr(document_restorer, "find_decoder_near_document", lambda path: decoders[Path(path)])
    monkeypatch.setattr(document_restorer, "restore_document",
                        lambda path, decoder_path=None: calls.append((Path(path), Path(decoder_path))) or path)
    monkeypatch.setattr(sys, "argv", ["main.py", "--restore", str(folder), "--headless"])

    assert main.handle_cli_arguments() is True
    payload = json.loads(printed[-1])
    assert payload["status"] == "ok"
    assert calls == [(document, decoders[document]) for document in documents]


def test_headless_restore_missing_decoder_exits_nonzero(monkeypatch, tmp_path, capsys):
    """Missing decoder is an operational error, not a successful empty result."""
    _install_webview_stub(monkeypatch)
    import main
    import document_restorer
    printed = _capture_cli_json(monkeypatch, main)

    document = tmp_path / "sample.docx"
    document.write_bytes(b"placeholder")
    monkeypatch.setattr(document_restorer, "find_decoder_near_document", lambda _path: None)
    monkeypatch.setattr(sys, "argv", ["main.py", "--restore", str(document), "--headless"])

    with pytest.raises(SystemExit) as exc_info:
        main.handle_cli_arguments()
    assert exc_info.value.code == 1
    payload = json.loads(printed[-1])
    assert payload["status"] == "error"
    assert payload["items"][0]["errors"][0]["message"] == "decoder_not_found"


def test_headless_no_open_output_flag_controls_anonymize_api(monkeypatch, tmp_path, capsys):
    """The explicit no-open flag controls Explorer side effects without changing the default."""
    _install_webview_stub(monkeypatch)
    import main
    printed = _capture_cli_json(monkeypatch, main)

    document = tmp_path / "sample.docx"
    document.write_bytes(b"placeholder")
    created = []

    class FakeApi:
        def __init__(self, **kwargs):
            self._window = None
            self.open_output_folder = None
            self.save_decoder = False
            self.background_mode = kwargs.get("background_mode", False)
            created.append(self)

    class FakeWorker:
        def __init__(self, _files, _api):
            pass

        def run(self):
            return {"status": "succeeded", "processed_count": 1, "error_count": 0, "total_changes": 1}

    fake_backend = types.ModuleType("backend_api")
    fake_backend.BackendApi = FakeApi
    fake_backend.Worker = FakeWorker
    monkeypatch.setitem(sys.modules, "backend_api", fake_backend)

    monkeypatch.setattr(sys, "argv", ["main.py", "--anonymize", str(document), "--headless"])
    assert main.handle_cli_arguments() is True
    assert created[-1].open_output_folder is True
    assert json.loads(printed[-1])["status"] == "ok"

    monkeypatch.setattr(sys, "argv", ["main.py", "--anonymize", str(document), "--headless", "--no-open-output"])
    assert main.handle_cli_arguments() is True
    assert created[-1].open_output_folder is False
    assert json.loads(printed[-1])["status"] == "ok"
    assert created[-1].background_mode is True


def test_headless_anonymize_aggregates_partial_status(monkeypatch, tmp_path, capsys):
    """A successful item plus a partial item must produce an overall partial result."""
    _install_webview_stub(monkeypatch)
    import main
    printed = _capture_cli_json(monkeypatch, main)

    documents = [tmp_path / "ok.docx", tmp_path / "partial.docx"]
    for document in documents:
        document.write_bytes(b"placeholder")

    class FakeApi:
        def __init__(self):
            self._window = None

    class FakeWorker:
        def __init__(self, files, _api):
            self.document = Path(files[0])

        def run(self):
            status = "partial" if self.document.stem == "partial" else "succeeded"
            return {"status": status, "processed_count": 1, "error_count": int(status == "partial")}

    fake_backend = types.ModuleType("backend_api")
    fake_backend.BackendApi = FakeApi
    fake_backend.Worker = FakeWorker
    monkeypatch.setitem(sys.modules, "backend_api", fake_backend)
    monkeypatch.setattr(sys, "argv", ["main.py", "--anonymize", *(str(path) for path in documents),
                                       "--headless", "--no-open-output"])

    with pytest.raises(SystemExit) as exc_info:
        main.handle_cli_arguments()
    assert exc_info.value.code == 1
    payload = json.loads(printed[-1])
    assert payload["status"] == "partial"


def test_headless_restore_aggregates_partial_status(monkeypatch, tmp_path, capsys):
    """A folder with one restored and one failed document must report partial."""
    _install_webview_stub(monkeypatch)
    import main
    import document_restorer
    printed = _capture_cli_json(monkeypatch, main)

    folder = tmp_path / "docs"
    folder.mkdir()
    documents = [folder / "ok.docx", folder / "broken.docx"]
    decoder = folder / "Дешифратор.json"
    decoder.write_text("{}", encoding="utf-8")
    for document in documents:
        document.write_bytes(b"placeholder")

    monkeypatch.setattr(document_restorer, "find_decoder_near_document", lambda _path: decoder)
    def fake_restore(path, decoder_path=None):
        if Path(path).stem == "broken":
            raise ValueError("synthetic restoration failure")
        return Path(path).with_name("ok_restored.docx")
    monkeypatch.setattr(document_restorer, "restore_document", fake_restore)
    monkeypatch.setattr(sys, "argv", ["main.py", "--restore", str(folder), "--headless", "--no-open-output"])

    with pytest.raises(SystemExit) as exc_info:
        main.handle_cli_arguments()
    assert exc_info.value.code == 1
    payload = json.loads(printed[-1])
    assert payload["status"] == "partial"
    assert payload["items"][0]["status"] == "partial"


def test_worker_empty_input_returns_failure_result(monkeypatch):
    """Headless callers can rely on Worker result metadata instead of GUI callbacks."""
    # backend_api imports pywebview even though this test exercises only Worker.
    stub = types.ModuleType("webview")
    stub.OPEN_DIALOG = 1
    stub.FOLDER_DIALOG = 2
    monkeypatch.setitem(sys.modules, "webview", stub)
    from backend_api import Worker

    cleaner = types.SimpleNamespace(_window=None)
    result = Worker([], cleaner).run()
    assert result["status"] == "failed"
    assert result["error_count"] == 1


def test_background_worker_skips_crash_recovery_and_saves_decoder(monkeypatch, tmp_path):
    """Quick Actions must not reach the recovery vault/keychain."""
    _install_webview_stub(monkeypatch)
    import backend_api
    import crash_recovery
    import decoder_binding

    def fail_if_called(*_args, **_kwargs):
        raise AssertionError("background worker must not use crash recovery")

    for name in ("start_batch_checkpoint", "update_batch_checkpoint", "complete_batch_checkpoint"):
        monkeypatch.setattr(crash_recovery, name, fail_if_called)
    monkeypatch.setattr(decoder_binding, "publish_binding", lambda *_args, **_kwargs: None)

    source = tmp_path / "sample.docx"
    output = tmp_path / "sample_cleaned.docx"
    source.write_bytes(b"source")
    output.write_bytes(b"cleaned")

    class BackgroundCleaner:
        _window = None
        enable_crash_recovery = False
        save_decoder = True
        emit_audit_sidecars = False
        open_output_folder = False
        user_exclusions = set()
        custom_replacements = set()

        def process_single_file(self, _path, _exclusions, _replacements, *, batch_mapping,
                                batch_entity_seen, defer_decoder):
            assert defer_decoder is True
            batch_mapping["[PER_1]"] = "Alice"
            self._last_generated_outputs = [str(output)]
            return 1

        def update_progress(self, *_args):
            pass

        def is_cancelled(self):
            return False

        def finish_operation(self, _status):
            pass

    result = backend_api.Worker([str(source)], BackgroundCleaner()).run()
    decoder = tmp_path / "sample_Дешифратор.json"
    assert result["status"] == "succeeded"
    assert decoder.is_file()
    assert json.loads(decoder.read_text(encoding="utf-8"))["[PER_1]"] == "Alice"


def test_background_backend_does_not_touch_review_vault(monkeypatch):
    """Headless construction must not invoke the OS keyring or persist review state."""
    _install_webview_stub(monkeypatch)
    import backend_api
    import review_context

    calls = []
    migration_calls = []
    monkeypatch.setattr(
        backend_api.state_migration,
        "run_state_migration",
        lambda **kwargs: migration_calls.append(kwargs) or {},
    )

    class ExplodingVault:
        def __init__(self, *args, **kwargs):
            calls.append((args, kwargs))
            raise AssertionError("background mode must not construct SecureContextVault")

    monkeypatch.setattr(review_context, "SecureContextVault", ExplodingVault)
    api = backend_api.BackendApi(background_mode=True, lazy_pullenti=True)
    assert api.background_mode is True
    assert api._review_vault is None
    assert api._review_private == {}
    assert migration_calls == [{"migrate_review_state": False}]
    api._persist_review_state()
    api.clear_review_findings()
    assert calls == []
    assert not api._review_queue_path.exists()
    assert not api._review_context_path.exists()


def test_gui_backend_still_constructs_review_vault(monkeypatch):
    """The regular GUI constructor retains encrypted review-vault behaviour."""
    _install_webview_stub(monkeypatch)
    import backend_api
    import review_context

    created = []

    class FakeVault:
        def __init__(self, path, **kwargs):
            created.append((path, kwargs))

        def load(self):
            return {}

        def save(self, _values):
            return None

    monkeypatch.setattr(review_context, "SecureContextVault", FakeVault)
    api = backend_api.BackendApi(background_mode=False, lazy_pullenti=True)
    assert api.background_mode is False
    assert isinstance(api._review_vault, FakeVault)
    assert created and created[0][1]["account"] == "local"


def test_cli_anonymize_single_file_headless(tmp_path):
    """Проверяет запуск --anonymize --headless для одного файла."""
    doc_path = tmp_path / "договор_купли_продажи.docx"
    doc = docx.Document()
    doc.add_paragraph("Генеральный директор Иванов Иван Иванович заключил договор № 123-45 с ООО «Ромашка».")
    doc.save(str(doc_path))

    cmd = [sys.executable, "main.py", "--anonymize", str(doc_path), "--headless"]
    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "utf-8"
    res = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", env=env)
    assert res.returncode == 0

    cleaned_files = list(tmp_path.glob("*_cleaned*"))
    decoders = [path for path in tmp_path.iterdir()
                if path.is_file() and "дешифратор" in path.name.casefold() and path.suffix.casefold() == ".json"]
    assert len(cleaned_files) == 1
    assert len(decoders) == 1


def test_cli_anonymize_folder_headless(tmp_path):
    """Проверяет запуск --anonymize --headless для папки с документами."""
    folder_to_test = tmp_path / "документы"
    folder_to_test.mkdir()
    doc_path = folder_to_test / "акт_выполненных_работ.docx"
    doc = docx.Document()
    doc.add_paragraph("Исполняющий обязанности директора Петров Петр Петрович подписал акт № 99 от 10.02.2025 г.")
    doc.save(str(doc_path))

    cmd = [sys.executable, "main.py", "--anonymize", str(folder_to_test), "--headless"]
    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "utf-8"
    res = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", env=env)
    assert res.returncode == 0

    output_dir = folder_to_test / "Обезличенные документы"
    assert output_dir.exists()
    out_files = list(output_dir.glob("*.docx"))
    out_decoders = [path for path in folder_to_test.iterdir()
                    if path.is_file() and path.name.casefold().startswith("дешифратор")
                    and path.suffix.casefold() == ".json"]
    assert len(out_files) == 1
    assert len(out_decoders) == 1


def test_extract_cli_paths(tmp_path):
    """Проверяет извлечение и очистку путей из аргументов командной строки."""
    from main import extract_cli_paths

    file_1 = tmp_path / "test file.docx"
    file_1.touch()
    dir_1 = tmp_path / "test dir"
    dir_1.mkdir()

    # 1. Простой путь в кавычках
    args = ["--anonymize", f'"{file_1}"']
    extracted = extract_cli_paths(args)
    assert extracted == [str(file_1.resolve())]

    # 2. Путь к папке с экранированной кавычкой из-за обратного слэша Windows Explorer
    simulated_escaped_dir = f'{dir_1}"'
    args = ["--anonymize", simulated_escaped_dir]
    extracted = extract_cli_paths(args)
    assert extracted == [str(dir_1.resolve())]

    # 3. Несуществующие пути отсекаются
    args = ["--anonymize", str(tmp_path / "non_existent.docx")]
    assert extract_cli_paths(args) == []


def test_startup_html_path_defined_for_normal_and_quick_mode():
    """Обычный запуск не должен падать с NameError: html_path is not defined."""
    from main import startup_html_path

    normal = startup_html_path(False)
    quick = startup_html_path(True)
    assert normal.endswith("index.html")
    assert Path(normal).exists()
    assert quick.endswith("quick_progress.html")
    assert Path(quick).exists()


def test_quick_progress_html_exists_and_valid():
    """Проверяет наличие и валидность структуры quick_progress.html."""
    quick_html = Path("web") / "quick_progress.html"
    assert quick_html.exists()
    content = quick_html.read_text(encoding="utf-8")
    assert "setGlobalProgress" in content
    assert "showMissingDecoder" in content
    assert "progress-fill" in content
    assert "btn-cancel" in content
    assert "btn-pick-decoder" in content


def test_api_wrapper_quick_mode_flags():
    """Проверяет инициализацию ApiWrapper в quick_mode."""
    from main import ApiWrapper

    wrapper = ApiWrapper(startup_files=["dummy.docx"], auto_start=True, restore_mode=False)
    assert wrapper.auto_start is True
    assert wrapper.restore_mode is False
    assert wrapper.startup_files == ["dummy.docx"]
    assert wrapper._api.background_mode is True

    wrapper_restore = ApiWrapper(startup_files=["dummy.docx"], auto_start=False, restore_mode=True)
    assert wrapper_restore.auto_start is False
    assert wrapper_restore.restore_mode is True
    assert wrapper_restore._api.background_mode is True
