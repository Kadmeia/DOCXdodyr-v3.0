"""Synthetic GUI smoke for the frozen runtime; never reads user documents."""

import json
import tempfile
import threading
import time
from pathlib import Path


def run():
    import webview
    import app_paths
    from docx import Document
    from main import ApiWrapper

    import os
    previous_env = {k: os.environ.get(f"DOCXDODYR_{k}_DIR") for k in ("DATA", "USER_DATA", "CONFIG", "CACHE", "LOG")}
    with tempfile.TemporaryDirectory(prefix="docxdodyr-gui-smoke-") as raw:
        root = Path(raw)
        for kind in ("DATA", "USER_DATA", "CONFIG", "CACHE", "LOG"):
            os.environ[f"DOCXDODYR_{kind}_DIR"] = str(root / kind.lower())
        input_dir = root / "input"
        input_dir.mkdir()
        source = input_dir / "synthetic.docx"
        doc = Document()
        doc.add_paragraph("fixture@example.invalid")
        doc.save(source)
        wrapper = ApiWrapper()
        window = webview.create_window(
            "DOCXdodyr frozen GUI smoke",
            url=str(app_paths.get_web_dir() / "index.html"),
            js_api=wrapper,
            width=800,
            height=600,
        )
        wrapper.set_window(window)
        result = {"status": "error", "dom_click": False, "processed": False, "closed": False}
        closed = threading.Event()
        loaded = threading.Event()

        def on_loaded():
            loaded.set()

        window.events.loaded += on_loaded

        def on_closed():
            wrapper.shutdown()
            result["closed"] = True
            closed.set()

        window.events.closed += on_closed

        def actions():
            try:
                if not loaded.wait(timeout=30):
                    raise RuntimeError("Frozen GUI did not load")
                window.evaluate_js("document.getElementById('btn-tab-anonymize').click();")
                result["dom_click"] = window.evaluate_js("document.getElementById('tab-anonymize').style.display === 'flex';") is True
                window.evaluate_js("window.pywebview.api.process_folder_dialog(%s, false);" % json.dumps(str(input_dir)))
                deadline = time.time() + 90
                while time.time() < deadline:
                    status = window.evaluate_js("document.getElementById('folder-status').textContent || '';" ) or ""
                    output_dir = input_dir / "Обезличенные документы"
                    outputs = list(output_dir.glob("*.docx")) if output_dir.exists() else []
                    if ("завершена" in status.lower() or "успешно" in status.lower()) and outputs:
                        from document_restorer import DocumentRestorer, find_decoder_near_document
                        from docx import Document
                        if any("fixture@example.invalid" in "\n".join(p.text for p in Document(item).paragraphs) for item in outputs):
                            raise RuntimeError("Frozen GUI did not anonymize synthetic DOCX")
                        decoder = find_decoder_near_document(outputs[0])
                        if decoder is None:
                            raise RuntimeError("Frozen GUI did not produce decoder")
                        restored = root / "restored.docx"
                        mapping = json.loads(decoder.read_text(encoding="utf-8"))
                        ok, _ = DocumentRestorer(mapping).restore_docx(outputs[0], restored)
                        if not ok or "fixture@example.invalid" not in "\n".join(p.text for p in Document(restored).paragraphs):
                            raise RuntimeError("Frozen GUI DOCX restoration failed")
                        result["processed"] = True
                        break
                    if "ошибка" in status.lower():
                        raise RuntimeError("Frozen GUI reported processing error")
                    time.sleep(0.5)
            except Exception:
                pass
            finally:
                try:
                    window.destroy()
                except Exception:
                    closed.set()

        threading.Thread(target=actions, daemon=True).start()
        webview.start(debug=False)
        closed.wait(timeout=5)
        wrapper.shutdown()
        for kind, value in previous_env.items():
            key = f"DOCXDODYR_{kind}_DIR"
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value
        if not (result["dom_click"] and result["processed"] and result["closed"]):
            raise RuntimeError("Frozen GUI smoke failed")
        result["status"] = "ok"
        return result
