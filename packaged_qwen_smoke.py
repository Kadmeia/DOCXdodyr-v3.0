"""Opt-in synthetic inference inside the frozen executable; local weights only."""
import os
from pathlib import Path
import tempfile


def run():
    from docx import Document
    from backend_api import BackendApi
    from qwen_postprocessor import (QwenPlaceholderPostprocessor, extract_placeholders,
                                    _FACT_NUMBER_RE, _MODALITY_RE)
    model = os.environ.get('QWEN_MODEL_DIR', '').strip()
    if not model or not Path(model).is_dir():
        raise ValueError('QWEN_MODEL_DIR must name existing local weights')
    processor = QwenPlaceholderPostprocessor({'enabled': True, 'model_path': model,
        'local_files_only': True, 'allow_network_download': False, 'max_new_tokens': 128,
        'temperature': 0.0})
    base = 'Уважаемый [ФИО_1] ,направляем Вам проект договора № [НомерДоговора_1] от 15.05.2025.'
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('Уважаемый [ФИО_1] ')
    paragraph.add_run(',направляем Вам проект договора № [НомерДоговора_1] от 15.05.2025.')
    document.add_paragraph(base)
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = base
    cell.add_paragraph(base)
    document.sections[0].header.paragraphs[0].text = base
    document.sections[0].footer.paragraphs[0].text = base
    def texts(doc):
        return [p.text for p in [*doc.paragraphs, *doc.tables[0].cell(0, 0).paragraphs,
            doc.sections[0].header.paragraphs[0], doc.sections[0].footer.paragraphs[0]]]
    def facts(text):
        return (extract_placeholders(text), _FACT_NUMBER_RE.findall(text),
                _MODALITY_RE.findall(text.casefold()))
    before = texts(document)
    try:
        api = BackendApi.__new__(BackendApi)
        api.qwen_settings = processor.settings
        api._qwen_postprocessor = processor
        api._write_back_qwen_document(document, [])
        with tempfile.TemporaryDirectory(prefix='docxdodyr-qwen-smoke-') as folder:
            path = Path(folder) / 'synthetic.docx'
            document.save(path)
            after = texts(Document(path))
        if len(before) != 6 or len(after) != 6 or processor.last_status != 'ok':
            raise RuntimeError('Frozen Qwen did not complete structural processing')
        for original, corrected in zip(before, after):
            if original == corrected or ' ,направляем' in corrected or facts(original) != facts(corrected):
                raise RuntimeError('Frozen Qwen structural invariants failed')
        return {'status': 'ok', 'local_inference': True, 'saved_reopened_structures': 6,
                'factual_invariants_preserved': True}
    finally:
        processor.reset()
